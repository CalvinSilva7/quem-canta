"""Primeira etapa do escritório: quem gravou cada obra do titular.

A busca usa a Deezer e a Apple Music, que mostram o compositor de cada gravação sem precisar de navegador:
uma gravação entra quando o crédito traz um autor do relatório, ou quando o intérprete já está ligado ao
titular. O resultado é só a lista, obra por obra, numa planilha simples: ela vai para o compositor conferir
e corrigir antes de qualquer coleta de prints. Por isso não leva status nem link.

A única marca é a cor: em preto, o intérprete de que o app tem certeza; em vermelho, o que ele só supõe.
Há certeza quando alguma gravação daquele intérprete traz o próprio titular no crédito, ou traz outro autor
da obra sem ressalva nenhuma. Quando a ligação é só presumida, ou pede revisão (crédito de um coautor que
pode ser homônimo, medley, título parecido), o nome sai em vermelho.
"""

import io

from cantor.matching import normalizar

from . import classificador as c
from . import pipeline

PLATAFORMAS = ("deezer", "apple")
VERMELHO = "C00000"
# O que não é intérprete confirmado de obra nenhuma: homônimo de terceiro, leitura que falhou, faixa fora do relatório.
_FORA = (c.HOMONIMA, c.ERRO_TECNICO, c.INDISPONIVEL, c.FORA_DO_REPERTORIO)


def _chave(nome: str) -> str:
    return " ".join(normalizar(str(nome or "").replace("&", " e ")).split())


def _da_certeza(g) -> bool:
    """Esta gravação prova que o intérprete gravou a obra do titular?"""
    k = g.classificacao
    if k.status == c.OK:  # o próprio titular está no crédito
        return True
    return k.status in c.NEGATIVOS and not k.revisar and bool(k.autores_reconhecidos)  # outro autor da obra, sem ressalva


def por_obra(relatorio, coletas) -> list[tuple[str, list[str], set[str]]]:
    """[(título, [intérpretes], {intérpretes sem certeza})], na ordem do relatório, um título por linha.

    Obra sem gravação vem com a lista vazia. Os intérpretes saem do mais gravado para o menos; o mesmo nome escrito
    com "e" e com "&" conta uma vez. O terceiro item diz quais nomes o app só supõe (ver `_da_certeza`).
    """
    achados = {}  # título -> {chave do nome: [nome como aparece, quantas gravações, há certeza?]}
    for coleta in coletas:
        for g in coleta.gravacoes:
            if g.classificacao.status in _FORA or not str(g.interprete or "").strip():
                continue
            for titulo in g.obra.split("; "):
                item = achados.setdefault(titulo, {}).setdefault(_chave(g.interprete), [g.interprete.strip(), 0, False])
                item[1] += 1
                item[2] = item[2] or (_da_certeza(g) and "; " not in g.obra)  # medley não dá certeza sobre nenhuma das obras
    saida = []
    for titulo in dict.fromkeys(o.titulo for o in relatorio.obras):
        nomes = sorted(achados.get(titulo, {}).values(), key=lambda item: (-item[1], normalizar(item[0])))
        saida.append((titulo, [nome for nome, _, _ in nomes], {nome for nome, _, certo in nomes if not certo}))
    return saida


def buscar(relatorio, config, pasta, nomes_dos_coautores=(), ao_avancar=None, limite=None) -> list[tuple[str, list[str], set[str]]]:
    """Roda a busca (Deezer e Apple Music, sem navegador e sem prints) e devolve a lista de `por_obra`."""
    coletas = pipeline.executar(relatorio, config, pasta, nomes_dos_coautores, plataformas=PLATAFORMAS, prints=False,
                                ao_avancar=ao_avancar, limite_youtube=limite)
    return por_obra(relatorio, coletas)


def planilha(relatorio, lista) -> bytes:
    """A planilha de intérpretes: o nome do titular, o cabeçalho e uma linha por (obra, intérprete).

    Obra sem intérprete encontrado sai numa linha só, com a segunda coluna em branco, para o compositor preencher.
    Intérprete de que o app não tem certeza sai em vermelho; os demais, em preto. `lista` aceita itens de dois
    campos (obra, intérpretes), todos em preto, ou de três, com o conjunto dos nomes sem certeza.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    livro = Workbook()
    ws = livro.active
    ws.title = "Intérpretes"
    titulo = Font(name="Bookman Old Style", bold=True, color="FFFFFF", size=10.5)
    corpo = Font(name="Verdana", size=9.5)
    em_duvida = Font(name="Verdana", size=9.5, color=VERMELHO)
    ws.append([(relatorio.pseudonimo_titular or relatorio.nome_titular or "").upper()])
    ws.append(["Obras", "Intérpretes"])
    for linha in ws.iter_rows(min_row=1, max_row=2, max_col=2):
        for celula in linha:
            celula.font, celula.fill = titulo, PatternFill("solid", fgColor="1B3A6B")
            celula.alignment = Alignment(horizontal="center")
    for obra, nomes, *resto in lista:
        duvidosos = resto[0] if resto else set()
        for nome in nomes or [""]:
            ws.append([obra, nome.upper() if nome else None])
            ws.cell(ws.max_row, 1).font = corpo
            ws.cell(ws.max_row, 2).font = em_duvida if nome in duvidosos else corpo
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 46, 40
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _colunas(linhas) -> tuple[int, int, int]:
    """Em que colunas estão a obra e o intérprete, e em que linha começam os dados: (obra, intérprete, primeira linha).

    Vale o cabeçalho, se houver ("Obras" / "Título da Música" / "Composição" e "Intérpretes"), em qualquer coluna:
    o escritório costuma pôr uma coluna de número antes. Sem cabeçalho reconhecível, valem as duas primeiras
    colunas, pulando a primeira se ela for só de números.
    """
    for i, linha in enumerate(linhas[:6]):
        nomes = [normalizar(str(c or "")) for c in linha]
        do_interprete = next((j for j, n in enumerate(nomes) if n.startswith("interprete")), None)
        da_obra = next((j for j, n in enumerate(nomes) if j != do_interprete and any(p in n.split() for p in ("obra", "obras", "titulo", "musica", "musicas", "composicao"))), None)
        if do_interprete is not None and da_obra is not None:
            return da_obra, do_interprete, i + 1
    primeira = [str(l[0] or "").strip() for l in linhas if l and str(l[0] or "").strip()]
    so_numeros = bool(primeira) and sum(v.rstrip(".").isdigit() or normalizar(v) in ("n", "no", "num", "numero") for v in primeira) >= 0.8 * len(primeira)
    return (1, 2, 0) if so_numeros else (0, 1, 0)


def ler_planilha(conteudo: bytes, relatorio) -> tuple[dict, list[str]]:
    """Lê a planilha de intérpretes (a nossa, ou a que o compositor devolveu corrigida).

    Devolve ({título do relatório: [intérpretes]}, [obras da planilha que não estão no relatório]). As colunas da obra
    e do intérprete são achadas pelo cabeçalho (ver `_colunas`); linhas de título e de cabeçalho são ignoradas.
    Vários intérpretes na mesma célula, separados por barra ("Fulano / Beltrano") ou quebra de linha, contam um a um.
    """
    import re

    from openpyxl import load_workbook

    por_titulo = {normalizar(o.titulo): o.titulo for o in relatorio.obras}
    linhas = [list(l) for l in load_workbook(io.BytesIO(conteudo), read_only=True).worksheets[0].iter_rows(min_col=1, max_col=8, values_only=True)]
    da_obra, do_interprete, comeco = _colunas(linhas)
    pares, fora = {}, []
    for linha in linhas[comeco:]:
        linha = linha + [None] * (max(da_obra, do_interprete) + 1 - len(linha))
        obra, celula = str(linha[da_obra] or "").strip(), str(linha[do_interprete] or "").strip()
        if not obra or not celula or normalizar(obra) in ("obras", "obra"):
            continue
        titulo = por_titulo.get(normalizar(obra))
        if titulo is None:
            fora.append(obra)
            continue
        for interprete in [n.strip() for n in re.split(r"\s*/{1,2}\s*|\n|;", celula) if n.strip()]:
            if _chave(interprete) not in {_chave(n) for n in pares.setdefault(titulo, [])}:
                pares[titulo].append(interprete)
    return pares, list(dict.fromkeys(fora))
