"""Juntar planilhas de obras de coletas feitas em separado (ex.: YouTube e Spotify num dia, Deezer e Palco MP3 em outro).

Copiar e colar uma coluna de uma planilha na outra põe os links nas obras erradas: cada coleta monta as linhas com o
que ela achou, então as linhas das duas não batem (uma obra tem 5 intérpretes numa e 1 na outra). Aqui nada é juntado
pela posição: as gravações de cada plataforma são lidas da aba "Gravações" de cada planilha (com a obra e o
intérprete de cada uma) e a aba "Obras" é montada de novo, do mesmo jeito que a coleta monta.

Por isso, o que alguém colou ou digitou à mão na aba "Obras" das planilhas enviadas não entra: vale o que o app coletou.
Se a mesma plataforma vier em mais de uma planilha, fica a coleta mais recente dela, e o app avisa.
"""

import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook, load_workbook

from cantor.matching import normalizar

from . import saida
from .classificador import Classificacao
from .modelo import Obra, Relatorio
from .pipeline import Coleta, Gravacao

ABAS = ("Obras", "Gravações", "Pendências - homônimos", "Resumo")
_COLUNAS = ["Plataforma", "Obra (relatório)", "Título exibido", "Intérprete", "Álbum", "Lançamento", "Link", "ISRC",
            "Fornecedor do metadado", "Crédito exibido", "Status", "Revisar", "Fundamento", "Coautores omitidos",
            "Vínculo com o titular", "Como foi achada", "Coletado em", "Capturas de tela"]


class PlanilhaInvalida(ValueError):
    """Mensagem pronta para mostrar na tela."""


@dataclass
class Lida:
    nome: str
    titular: str
    titulos: list[str]
    gravacoes: dict[str, list[Gravacao]]  # plataforma -> gravações (das abas Gravações e Pendências)
    resumo: dict[str, list[tuple[str, object]]]  # plataforma -> linhas da seção dela no Resumo
    coletadas: list[str] = field(default_factory=list)

    def coletada_em(self, plataforma: str) -> str:
        return max((g.coletado_em or "" for g in self.gravacoes.get(plataforma, [])), default="")


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _lista(valor, separador) -> list[str]:
    return [p.strip() for p in _texto(valor).split(separador) if p.strip()]


def _gravacao(linha) -> Gravacao:
    v = dict(zip(_COLUNAS, linha))
    creditos = _lista(v["Crédito exibido"], ",")
    return Gravacao(
        plataforma=_texto(v["Plataforma"]), obra=_texto(v["Obra (relatório)"]), titulo=_texto(v["Título exibido"]),
        interprete=_texto(v["Intérprete"]), album=_texto(v["Álbum"]), lancamento=_texto(v["Lançamento"]),
        link=_texto(v["Link"]), isrc=_texto(v["ISRC"]), fornecedor=_texto(v["Fornecedor do metadado"]),
        creditos=[] if creditos == ["(nenhum)"] else creditos, vinculo=_texto(v["Vínculo com o titular"]),
        origem=_texto(v["Como foi achada"]), coletado_em=_texto(v["Coletado em"]), provas=_lista(v["Capturas de tela"], ";"),
        classificacao=Classificacao(status=_texto(v["Status"]), fundamento=_texto(v["Fundamento"]),
                                    revisar=_texto(v["Revisar"]).lower() == "sim",
                                    coautores_omitidos=_lista(v["Coautores omitidos"], ",")),
    )  # fmt: skip


def ler(conteudo: bytes, nome: str = "") -> Lida:
    """Uma planilha de obras gerada pelo app (com as abas Obras, Gravações, Pendências e Resumo)."""
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        raise PlanilhaInvalida(f"“{nome}” não abriu como planilha do Excel.") from erro
    if any(aba not in wb.sheetnames for aba in ABAS):
        raise PlanilhaInvalida(f"“{nome}” não é uma planilha de obras gerada pelo app (faltam as abas {', '.join(ABAS)}).")
    obras = list(wb["Obras"].iter_rows(values_only=True))
    if len(obras) < 2 or [_texto(x) for x in obras[1][:2]] != ["Obras", "Intérpretes"]:
        raise PlanilhaInvalida(f"“{nome}”: a aba Obras não está no formato do app.")
    titular = _texto(obras[0][0])
    # Os títulos do relatório (uma linha "(nenhuma gravação encontrada)" também conta); "A; B" é um medley de duas obras
    titulos = [t for linha in obras[2:] if _texto(linha[1]) for t in _texto(linha[0]).split("; ") if t]

    gravacoes: dict[str, list[Gravacao]] = {}
    for aba in ("Gravações", "Pendências - homônimos"):
        linhas = list(wb[aba].iter_rows(values_only=True))
        if not linhas or [_texto(x) for x in linhas[0][: len(_COLUNAS)]] != _COLUNAS:
            raise PlanilhaInvalida(f"“{nome}”: a aba {aba} não está no formato do app.")
        for linha in linhas[1:]:
            linha = list(linha[: len(_COLUNAS)]) + [None] * (len(_COLUNAS) - len(linha))
            if _texto(linha[0]) in saida.PLATAFORMAS:
                g = _gravacao(linha)
                gravacoes.setdefault(g.plataforma, []).append(g)

    # Resumo: cada plataforma coletada tem a seção dela (mesmo sem gravação nenhuma, ela foi verificada)
    resumo: dict[str, list[tuple[str, object]]] = {}
    atual = None
    for rotulo, valor, *_ in (list(r) + [None, None] for r in wb["Resumo"].iter_rows(min_row=2, values_only=True)):
        rotulo = _texto(rotulo)
        if rotulo in saida.PLATAFORMAS and valor in (None, ""):
            atual = resumo.setdefault(rotulo, [])
        elif not rotulo and valor in (None, ""):
            atual = None
        elif atual is not None:
            atual.append((rotulo, valor if valor is not None else ""))
    coletadas = [p for p in saida.PLATAFORMAS if p in resumo or p in gravacoes]
    wb.close()
    return Lida(nome=nome, titular=titular, titulos=titulos, gravacoes=gravacoes, resumo=resumo, coletadas=coletadas)


@dataclass
class Juntada:
    planilha: bytes
    de_onde: dict[str, str]  # plataforma -> nome da planilha de onde veio
    avisos: list[str]


def juntar(lidas: list[Lida], relatorio: Relatorio | None = None) -> Juntada:
    if len(lidas) < 2:
        raise PlanilhaInvalida("Envie pelo menos duas planilhas para juntar.")
    titulares = {normalizar(l.titular) for l in lidas}
    if len(titulares) > 1:
        raise PlanilhaInvalida("As planilhas são de compositores diferentes: " + "; ".join(f"“{l.nome}”: {l.titular}" for l in lidas))
    avisos = []
    if relatorio is not None:
        do_relatorio = " - ".join(x for x in (relatorio.nome_titular, relatorio.pseudonimo_titular) if x)
        if normalizar(do_relatorio) not in titulares:
            avisos.append(f"As planilhas são de {lidas[0].titular}, e o relatório aberto no app é de {do_relatorio}.")
            relatorio = None

    fontes_escolhidas, coletas = {}, []
    for plataforma in saida.PLATAFORMAS:
        fontes = [l for l in lidas if plataforma in l.coletadas]
        if not fontes:
            continue
        # A coleta mais recente da plataforma (empate ou sem data: a planilha enviada por último)
        escolhida = max(enumerate(fontes), key=lambda par: (par[1].coletada_em(plataforma), par[0]))[1]
        if len(fontes) > 1:
            avisos.append(f"{plataforma} veio em {len(fontes)} planilhas: ficou a coleta mais recente, de “{escolhida.nome}”.")
        fontes_escolhidas[plataforma] = escolhida
        coletas.append(Coleta(plataforma, gravacoes=list(escolhida.gravacoes.get(plataforma, []))))

    if relatorio is None:
        nome, _, pseudonimo = lidas[0].titular.partition(" - ")
        titulos = dict.fromkeys(t for l in lidas for t in l.titulos)
        relatorio = Relatorio(nome_titular=nome, pseudonimo_titular=pseudonimo, obras=[Obra(titulo=t) for t in titulos])

    wb = Workbook()
    saida._aba_obras(wb.active, relatorio, coletas)
    saida._aba_gravacoes(wb.create_sheet("Gravações"), coletas, lambda g: g.classificacao.status != saida.c.HOMONIMA)
    saida._aba_gravacoes(wb.create_sheet("Pendências - homônimos"), coletas, lambda g: g.classificacao.status == saida.c.HOMONIMA)
    _aba_resumo(wb.create_sheet("Resumo"), lidas[0].titular, relatorio, coletas, fontes_escolhidas)
    arquivo = io.BytesIO()
    wb.save(arquivo)
    return Juntada(arquivo.getvalue(), {p: l.nome for p, l in fontes_escolhidas.items()}, avisos)


def _aba_resumo(ws, titular, relatorio, coletas, fontes):
    """O resumo de cada plataforma vem da planilha de onde ela veio, como estava."""
    saida._cabecalho(ws, 1, ["Resumo da execução", ""])
    linhas = [
        ("Titular", titular),
        ("Obras no relatório", len({normalizar(o.titulo) for o in relatorio.obras})),
        ("Planilha juntada pelo app em", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    linhas += [(f"   {p}", f"de “{fonte.nome}”") for p, fonte in fontes.items()]
    for coleta in coletas:
        linhas += [("", ""), (coleta.plataforma, "")] + fontes[coleta.plataforma].resumo.get(coleta.plataforma, [])
    for n, (rotulo, valor) in enumerate(linhas, start=2):
        ws.cell(n, 1, rotulo).font = saida._NEGRITO if valor == "" else saida._CORPO
        ws.cell(n, 2, valor).font = saida._CORPO
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 62, 90
