"""Saídas: a planilha no formato do escritório e o resumo da execução.

Aba "Obras": uma linha por (obra, intérprete) e uma coluna por plataforma, como a
planilha que o escritório preenche à mão. Em cada célula: os links das gravações
sem crédito ou com crédito errado (a prova), separados por " -- "; "tem créditos";
"N/A" quando a obra não está lá; ou "verificar manualmente".
Aba "Gravações": uma linha por gravação, com status, fundamento e origem.
Aba "Resumo": os denominadores reais (verificadas / coletadas) e o que falhou.
"""

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cantor.matching import normalizar

from . import classificador as c

# Ordem das colunas da planilha do escritório.
# A Apple Music não está na planilha do escritório: entra no fim, como plataforma de controle.
PLATAFORMAS = ["YOUTUBE", "SPOTIFY", "TIDAL", "DEEZER", "CLARO", "BRISA MUSIC", "PALCO MP3", "VAGALUME", "NAPSTER", "AMAZON",
               "AMAZON - SITE",
               "APPLE MUSIC"]
ENCERRADAS = {"BRISA MUSIC", "NAPSTER"}
# "AMAZON" é o aplicativo de desktop da Amazon Music, o único lugar em que ela mostra o compositor ("Créditos").
# Ele não é um site e exige a conta logada: o app não coleta ali. O site é a coluna "AMAZON - SITE".
SO_NO_APLICATIVO = {"AMAZON"}
SEPARADOR = " -- "
TEM_CREDITOS, NAO_ESTA, VERIFICAR = "tem créditos", "N/A", "verificar manualmente"
NAO_COLETADO = "não coletado nesta versão"
A_CONFIRMAR = "A CONFIRMAR (pode ser outra música): "

NAVY = "1B3A6B"
CORES = {  # identidade visual do escritório, por status
    c.OK: "D9EAD3", c.VIOLACAO: "F4CCCC", c.SEM_CREDITOS: "EFEFEF", c.LINK_DIVERGENTE: "FCE5CD",
    c.ERRO_TECNICO: "FFF2CC", c.INDISPONIVEL: "FFF2CC", c.OK_VARIANTE: "EEF6E8", c.GRAFIA: "E6DDF2",
    c.MEDLEY: "DDEBF7", c.HOMONIMA: "DDEBF7", c.TITULO_APROXIMADO: "DDEBF7", A_CONFIRMAR: "FFF2CC",
}
_CABECALHO = Font(name="Bookman Old Style", bold=True, color="FFFFFF", size=10.5)
_CORPO = Font(name="Verdana", size=9.5)
_NEGRITO = Font(name="Verdana", size=9.5, bold=True)


def _fundo(cor):
    return PatternFill("solid", fgColor=cor)


def linhas_por_obra(relatorio, coletas) -> list[dict]:
    """[{"obra", "interprete", "celulas": {plataforma: (texto, status que dá a cor)}}], em ordem de título."""
    pares = {}  # (título da obra, intérprete normalizado) -> {"interprete", plataforma: [gravações]}
    for coleta in coletas:
        for g in coleta.gravacoes:
            if g.classificacao.status == c.HOMONIMA:
                continue  # homônimo provável nunca entra direto: fica na aba de pendências até alguém confirmar
            par = pares.setdefault((g.obra, normalizar(g.interprete)), {"interprete": g.interprete})
            par.setdefault(coleta.plataforma, []).append(g)
    coletadas = {coleta.plataforma for coleta in coletas}
    linhas = []
    for (obra, _), par in pares.items():
        celulas = {p: _celula(par.get(p, [])) if p in coletadas else _nao_coletada(p) for p in PLATAFORMAS}
        linhas.append({"obra": obra, "interprete": par["interprete"], "celulas": celulas})
    com_gravacao = {normalizar(t) for linha in linhas for t in linha["obra"].split("; ")}
    for titulo in dict.fromkeys(o.titulo for o in relatorio.obras):
        if normalizar(titulo) not in com_gravacao:
            vazias = {p: (NAO_ESTA, "") if p in coletadas else _nao_coletada(p) for p in PLATAFORMAS}
            linhas.append({"obra": titulo, "interprete": "(nenhuma gravação encontrada)", "celulas": vazias})
    return sorted(linhas, key=lambda l: (normalizar(l["obra"]), normalizar(l["interprete"])))


def _nao_coletada(plataforma):
    if plataforma in ENCERRADAS:
        return "plataforma encerrada", ""
    if plataforma in SO_NO_APLICATIVO:
        return "não coletado: aplicativo de desktop (print manual)", ""
    return NAO_COLETADO, ""


def _celula(gravacoes) -> tuple[str, str]:
    """Texto da célula e o status que define a cor."""
    if not gravacoes:
        return NAO_ESTA, ""
    negativas = [g for g in gravacoes if g.classificacao.status in c.NEGATIVOS]
    if negativas:
        # Negativa marcada para revisão não é prova firme: sai separada, com aviso, e não entra na contagem do rodapé.
        firmes = [g for g in negativas if not g.classificacao.revisar]
        duvidosas = [g for g in negativas if g.classificacao.revisar]
        texto = SEPARADOR.join(g.link for g in firmes)
        if duvidosas:
            texto += ("\n" if firmes else "") + A_CONFIRMAR + SEPARADOR.join(g.link for g in duvidosas)
        if not firmes:
            return texto, A_CONFIRMAR
        status = c.VIOLACAO if any(g.classificacao.status == c.VIOLACAO for g in firmes) else c.SEM_CREDITOS
        return texto, status
    pendentes = [g for g in gravacoes if g.classificacao.status != c.OK]
    if any(g.classificacao.status == c.OK for g in gravacoes):
        return TEM_CREDITOS + (f" (+{len(pendentes)} a verificar)" if pendentes else ""), c.OK
    return VERIFICAR, pendentes[0].classificacao.status


def links_de_prova(linhas, plataforma) -> int:
    return sum(l["celulas"][plataforma][0].split(A_CONFIRMAR)[0].count("http") for l in linhas)


def planilha(relatorio, coletas) -> bytes:
    wb = Workbook()
    _aba_obras(wb.active, relatorio, coletas)
    _aba_gravacoes(wb.create_sheet("Gravações"), coletas, lambda g: g.classificacao.status != c.HOMONIMA)
    _aba_gravacoes(wb.create_sheet("Pendências - homônimos"), coletas, lambda g: g.classificacao.status == c.HOMONIMA)
    _aba_resumo(wb.create_sheet("Resumo"), relatorio, coletas)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _cabecalho(ws, linha, textos):
    for coluna, texto in enumerate(textos, start=1):
        celula = ws.cell(linha, coluna, texto)
        celula.font, celula.fill = _CABECALHO, _fundo(NAVY)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _aba_obras(ws, relatorio, coletas):
    ws.title = "Obras"
    titulo = " - ".join(x for x in (relatorio.nome_titular, relatorio.pseudonimo_titular) if x)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    _cabecalho(ws, 1, [titulo])
    _cabecalho(ws, 2, ["Obras", "Intérpretes"] + PLATAFORMAS)
    linhas = linhas_por_obra(relatorio, coletas)
    for n, linha in enumerate(linhas, start=3):
        ws.cell(n, 1, linha["obra"]).font = _CORPO
        ws.cell(n, 2, linha["interprete"]).font = _CORPO
        for coluna, plataforma in enumerate(PLATAFORMAS, start=3):
            texto, status = linha["celulas"][plataforma]
            celula = ws.cell(n, coluna, texto)
            celula.font, celula.alignment = _CORPO, Alignment(wrap_text=True, vertical="top")
            if status in CORES:
                celula.fill = _fundo(CORES[status])
    rodape = len(linhas) + 3
    coletadas = {coleta.plataforma for coleta in coletas}
    for coluna, plataforma in enumerate(PLATAFORMAS, start=3):
        if plataforma in coletadas:
            total = links_de_prova(linhas, plataforma)
            # "obras" é como o escritório chama cada link listado na petição.
            ws.cell(rodape, coluna, f"{total} obras" if total else "NENHUMA OBRA").font = _NEGRITO
    ws.freeze_panes = "C3"
    for letra, largura in zip("AB", (38, 28)):
        ws.column_dimensions[letra].width = largura
    for coluna in range(3, 3 + len(PLATAFORMAS)):
        ws.column_dimensions[get_column_letter(coluna)].width = 30


def _aba_gravacoes(ws, coletas, entra):
    colunas = ["Plataforma", "Obra (relatório)", "Título exibido", "Intérprete", "Álbum", "Lançamento", "Link", "ISRC",
               "Fornecedor do metadado", "Crédito exibido", "Status", "Revisar", "Fundamento", "Coautores omitidos",
               "Vínculo com o titular", "Como foi achada", "Coletado em", "Capturas de tela"]
    _cabecalho(ws, 1, colunas)
    n = 1
    for coleta in coletas:
        for g in sorted(filter(entra, coleta.gravacoes), key=lambda g: (normalizar(g.obra), normalizar(g.interprete), g.link)):
            n += 1
            k = g.classificacao
            valores = [g.plataforma, g.obra, g.titulo, g.interprete, g.album, g.lancamento, g.link, g.isrc, g.fornecedor,
                       ", ".join(g.creditos) or "(nenhum)", k.status, "sim" if k.revisar else "", k.fundamento,
                       ", ".join(k.coautores_omitidos), g.vinculo, g.origem, g.coletado_em, "; ".join(g.provas)]
            cor = CORES.get(A_CONFIRMAR if k.status in c.NEGATIVOS and k.revisar else k.status)
            for coluna, valor in enumerate(valores, start=1):
                celula = ws.cell(n, coluna, valor)
                celula.font = _CORPO
                if cor:
                    celula.fill = _fundo(cor)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(n, 1)}"
    for coluna, largura in enumerate([11, 30, 30, 24, 28, 12, 40, 15, 24, 34, 34, 8, 80, 30, 34, 30, 20, 50], start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura


def contagens(relatorio, coletas) -> list[tuple[str, object]]:
    """O resumo da execução, sempre com o denominador real."""
    linhas = [("Titular", " - ".join(x for x in (relatorio.nome_titular, relatorio.pseudonimo_titular) if x)),
              ("Obras no relatório", len(relatorio.obras))]
    for coleta in coletas:
        status = Counter(g.classificacao.status for g in coleta.gravacoes)
        verificadas = sum(n for s, n in status.items() if s not in (c.ERRO_TECNICO, c.INDISPONIVEL))
        com_gravacao = {t for g in coleta.gravacoes if g.classificacao.status != c.HOMONIMA for t in g.obra.split("; ")}
        linhas += [
            ("", ""), (f"{coleta.plataforma}", ""),
            ("Nomes cuja discografia foi percorrida", ", ".join(coleta.sementes)),
            ("Álbuns (ou faixas) lidos / com erro", f"{coleta.albuns_lidos} / {len(coleta.albuns_com_erro)}"),
            ("Resultados de busca com o título, de artista sem ligação com o titular (não abertos)", coleta.sem_vinculo_nao_abertos),
            ("Faixas lidas que não são do relatório", coleta.fora_do_repertorio),
            ("Gravações do repertório: verificadas / coletadas", f"{verificadas} / {len(coleta.gravacoes)}"),
            ("Obras (títulos) com gravação ligada ao titular / no relatório",
             f"{len(com_gravacao)} / {len({o.titulo for o in relatorio.obras})}"),
            ("Gravações que pedem revisão humana", sum(g.classificacao.revisar for g in coleta.gravacoes)),
        ]
        linhas += [(f"   {s}", n) for s, n in status.most_common()]
        linhas += [("Intérpretes com vínculo inferido (confirmar)",
                    "; ".join(f"{i['nome']} ({i['obras']} títulos, {i['com_autor']} com autor no crédito)" for i in coleta.interpretes_inferidos) or "nenhum")]
        negativas = Counter(g.fornecedor or "(sem fornecedor)" for g in coleta.gravacoes if g.classificacao.status in c.NEGATIVOS)
        linhas += [(f"   sem crédito ou crédito errado, fornecedor {f}", n) for f, n in negativas.most_common()]
        linhas += [("Requisições feitas / bloqueadas pelo filtro de dados sensíveis", f"{coleta.requisicoes} / {coleta.bloqueadas}"),
                   ("Tempo de coleta", f"{coleta.segundos:.0f} s")]
        linhas += [("ERRO (não verificado)", f"{link}: {erro}") for link, erro in coleta.albuns_com_erro]
        linhas += [("Aviso", a) for a in coleta.avisos]
    return linhas


def _aba_resumo(ws, relatorio, coletas):
    _cabecalho(ws, 1, ["Resumo da execução", ""])
    for n, (rotulo, valor) in enumerate(contagens(relatorio, coletas), start=2):
        ws.cell(n, 1, rotulo).font = _NEGRITO if valor == "" else _CORPO
        ws.cell(n, 2, valor).font = _CORPO
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 62, 90


def resumo_em_texto(relatorio, coletas) -> str:
    return "\n".join(f"{rotulo}: {valor}" if valor != "" else rotulo for rotulo, valor in contagens(relatorio, coletas))
