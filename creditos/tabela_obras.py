"""Tabela de obras: o Word com uma linha por obra e os recortes do relatório analítico, do ISWC e do Ecadnet.

Entra o que o escritório já tem do caso: o relatório analítico do ECAD (PDF), os PDFs do ISWC (um só, vários ou o .zip
de um por obra) e os prints do Ecadnet (Word ou PDF com os prints de tela). As linhas da tabela são as obras do ISWC,
na ordem do relatório, com o título do relatório.

- Relatório e ISWC são PDFs com texto: a obra é achada pelo ISWC e o recorte sai do próprio PDF, em alta resolução.
- Ecadnet são prints de tela (imagem): o leitor de texto que vem no Windows lê o endereço (que traz o código da obra,
  o mesmo "CÓD. OBRA" do relatório), o ISWC e o título; o recorte vai da barra azul com o título até o fim da última
  tabela. Print sem a barra azul é a continuação (página rolada) da obra do print anterior.
Nada é instalado: o leitor é o do Windows (Windows.Media.Ocr), chamado pelo PowerShell.
"""

import io
import json
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

ESCALA_DO_RELATORIO = 4  # pontos do PDF -> pixels (288 dpi): o recorte fica nítido na largura da coluna
ESCALA_DO_ISWC = 3
LARGURAS_DXA = (2262, 3355, 4017, 5989)  # as colunas do modelo do escritório (em 1/20 de ponto)
MARGEM_DA_CELULA_DXA = 108
CABECALHO = [
    ("OBRA MUSICAL", None),
    ("RELATÓRIO ANALÍTICO", "(os repasses de arrecadação e distribuição de direitos ocorrem através desse registro)"),
    ("ISWC", "(código exclusivo atribuído a uma obra ou composição específica escrita pelo compositor)"),
    ("ECADNET", "(site que disponibiliza os detalhes de cada obra musical, como a indicação do compositor)"),
]
FONTE = "Bookman Old Style"
_ISWC = re.compile(r"T-\d{3}\.\d{3}\.\d{3}-\d")


def digitos(iswc: str) -> str:
    return re.sub(r"\D", "", iswc or "")


def titulo_da_tabela(titulo: str) -> str:
    """ "24 HORAS DE SOLTEIRO" -> "24 Horas De Solteiro", como o escritório escreve na tabela."""
    return " ".join(p[:1].upper() + p[1:].lower() for p in (titulo or "").split())


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto or "") if not unicodedata.combining(c))


def _chave_do_titulo(texto: str) -> str:
    return re.sub(r"[^a-z0-9]", "", _sem_acento(texto).lower())


# ---------------------------------------------------------------- relatório analítico


@dataclass
class Trecho:
    pagina: int
    x0: float
    top: float
    x1: float
    bottom: float


@dataclass
class ObraDoRelatorio:
    codigo: str
    iswc: str
    titulo: str
    trechos: list[Trecho] = field(default_factory=list)
    ate_o_cliente: bool = False  # já chegou na linha do titular do relatório: os nomes de depois não entram


def _linhas(palavras, tolerancia=3):
    linhas = []
    for p in sorted(palavras, key=lambda p: (round(p["top"]), p["x0"])):
        if linhas and abs(linhas[-1][0]["top"] - p["top"]) <= tolerancia:
            linhas[-1].append(p)
        else:
            linhas.append([p])
    return [sorted(linha, key=lambda p: p["x0"]) for linha in linhas]


def obras_do_relatorio(conteudo: bytes) -> list[ObraDoRelatorio]:
    """As obras do relatório analítico do ECAD com o lugar de cada uma: a linha da obra (código, ISWC e título, em
    negrito) e os nomes dos titulares (em itálico) que vêm embaixo, até o do próprio cliente (o titular do relatório,
    pelo código dele), como o escritório recorta; se os nomes passam para a página seguinte, ela entra também."""
    import pdfplumber

    obras: list[ObraDoRelatorio] = []
    colunas, cliente = {}, ""
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        for n, pagina in enumerate(pdf.pages):
            palavras = pagina.extract_words(x_tolerance=1.5, extra_attrs=["fontname"])  # ISWC e título não se juntam
            for p in palavras:
                if not p["fontname"].startswith("Courier") and p["text"] in ("ISWC", "TÍTULO", "ASS.", "PSEUDÔNIMO", "NOME"):
                    colunas.setdefault(p["text"], p["x0"])
            if not cliente:  # "TITULAR: 368368 PAULO …" no cabeçalho
                rotulo = next((p for p in palavras if p["text"] == "TITULAR:"), None)
                codigo = rotulo and next((p for p in palavras if abs(p["top"] - rotulo["top"]) < 3 and p["x0"] > rotulo["x1"]), None)
                if codigo and codigo["text"].isdigit():
                    cliente = codigo["text"]
            if "ISWC" not in colunas:
                continue
            x_iswc, x_titulo = colunas["ISWC"], colunas.get("TÍTULO", colunas["ISWC"] + 60)
            x_ass = colunas.get("ASS.", 1e9)
            for linha in _linhas(p for p in palavras if p["fontname"].startswith("Courier")):
                negrito = [p for p in linha if "Bold" in p["fontname"]]
                italico = [p for p in linha if "Oblique" in p["fontname"]]
                if negrito and re.fullmatch(r"\d+", negrito[0]["text"]) and negrito[0]["x0"] < x_iswc - 10:
                    iswc = next((p for p in negrito if _ISWC.fullmatch(p["text"])), None)
                    inicio = (iswc["x1"] + 1) if iswc else x_titulo - 6
                    titulo = [p for p in negrito if p["x0"] >= inicio and p["x1"] < x_ass - 1]
                    partes = ([iswc] if iswc else []) + titulo
                    obra = ObraDoRelatorio(negrito[0]["text"], iswc["text"] if iswc else "", " ".join(p["text"] for p in titulo))
                    if partes:
                        obra.trechos.append(Trecho(n, x_iswc - 2, min(p["top"] for p in partes) - 2,
                                                   max(p["x1"] for p in partes) + 2, max(p["bottom"] for p in partes) + 2))
                    obras.append(obra)
                elif italico and not negrito and obras and not obras[-1].ate_o_cliente:
                    obra = obras[-1]
                    trecho = obra.trechos[-1] if obra.trechos and obra.trechos[-1].pagina == n else None
                    if trecho is None:  # os nomes continuam no alto da página seguinte
                        trecho = Trecho(n, x_iswc - 2, min(p["top"] for p in italico) - 2, x_iswc, 0)
                        obra.trechos.append(trecho)
                    trecho.x1 = max(trecho.x1, max(p["x1"] for p in italico) + 2)
                    trecho.bottom = max(trecho.bottom, max(p["bottom"] for p in italico) + 2)
                    if cliente and linha[0]["text"] == cliente and linha[0]["x0"] < x_iswc - 10:
                        obra.ate_o_cliente = True
    return obras


def _recortar_pdf(documento, numero, x0, top, x1, bottom, escala):
    pagina = documento[numero]
    largura, altura = pagina.get_size()
    imagem = pagina.render(scale=escala, crop=(max(0, x0), max(0, altura - bottom), max(0, largura - x1), max(0, top)))
    return imagem.to_pil().convert("RGB")


def recortes_do_relatorio(conteudo: bytes, obra: ObraDoRelatorio) -> list:
    """Os recortes da obra no relatório (um por página em que ela aparece), como imagens."""
    import pypdfium2 as pdfium

    documento = pdfium.PdfDocument(conteudo)
    try:
        return [_recortar_pdf(documento, t.pagina, t.x0, t.top, t.x1, t.bottom, ESCALA_DO_RELATORIO) for t in obra.trechos]
    finally:
        documento.close()


# ---------------------------------------------------------------- ISWC


def recortes_do_iswc(conteudo: bytes) -> dict:
    """{dígitos do ISWC: imagem} de um PDF do ISWC Network (uma obra por página): o quadro do resultado, de "ISWC:"
    até "Creator Name(s)" (sem o "IP Name Number" e o "View More"), como o escritório recorta."""
    import pdfplumber
    import pypdfium2 as pdfium

    achados = {}
    documento = pdfium.PdfDocument(conteudo)
    try:
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            for n, pagina in enumerate(pdf.pages):
                palavras = pagina.extract_words()
                for rotulo in (p for p in palavras if p["text"] == "ISWC:"):
                    valor = next((p for p in palavras if abs(p["top"] - rotulo["top"]) < 4 and p["x0"] > rotulo["x1"]
                                  and _ISWC.fullmatch(p["text"])), None)  # fmt: skip
                    if valor is None:
                        continue  # o "ISWC:" do campo de busca: o valor fica na caixa embaixo, não na mesma linha
                    caixas = [r for r in pagina.rects if r["x0"] <= rotulo["x0"] and r["x1"] >= valor["x1"]
                              and r["top"] <= rotulo["top"] and r["bottom"] >= rotulo["bottom"] + 20]  # fmt: skip
                    if not caixas:
                        continue
                    caixa = min(caixas, key=lambda r: (r["bottom"] - r["top"]) * (r["x1"] - r["x0"]))
                    ip = next((p for p in palavras if p["text"] == "IP" and caixa["top"] < p["top"] < caixa["bottom"]
                               and abs(p["x0"] - rotulo["x0"]) < 5), None)  # fmt: skip
                    fim = caixa["bottom"]
                    if ip is not None:
                        divisorias = [r["top"] for r in pagina.rects if r["bottom"] - r["top"] < 2
                                      and r["x1"] - r["x0"] > 0.8 * (caixa["x1"] - caixa["x0"]) and rotulo["top"] < r["top"] < ip["top"]]
                        fim = max(divisorias) if divisorias else ip["top"] - 8
                    achados.setdefault(digitos(valor["text"]), _recortar_pdf(
                        documento, n, caixa["x0"] - 1, caixa["top"] - 1, caixa["x1"] + 1, fim + 1, ESCALA_DO_ISWC))
    finally:
        documento.close()
    return achados


# ---------------------------------------------------------------- arquivos enviados


def pdfs_de(nome: str, conteudo: bytes) -> list[bytes]:
    """Os PDFs de um arquivo enviado: o próprio PDF, ou os de dentro de um .zip."""
    if nome.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
            return [pacote.read(i) for i in sorted(pacote.namelist()) if i.lower().endswith(".pdf")]
    return [conteudo]


def _pixels(objeto) -> int:
    """Largura x altura (em pixels) de uma imagem dentro do PDF. O nome do método mudou entre as versões do
    pypdfium2 (get_size, depois get_px_size), e a versão é a que veio com o pdfplumber de cada computador."""
    largura, altura = (getattr(objeto, "get_px_size", None) or objeto.get_size)()
    return largura * altura


def imagens_de(nome: str, conteudo: bytes) -> list[bytes]:
    """Os prints de um arquivo enviado, na ordem, cada um como arquivo de imagem (PNG ou JPG): as imagens de um Word,
    as de um PDF (uma por página), uma imagem solta, ou tudo isso dentro de um .zip. Ficam comprimidos: são abertos
    um lote por vez na leitura."""
    minusculo = nome.lower()
    if minusculo.endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
            return [im for i in sorted(pacote.namelist()) for im in imagens_de(i, pacote.read(i))]
    if minusculo.endswith(".docx"):
        from docx import Document
        from docx.oxml.ns import qn

        documento = Document(io.BytesIO(conteudo))
        partes = documento.part.related_parts
        return [partes[b.get(qn("r:embed"))].blob for b in documento.element.body.iter(qn("a:blip")) if b.get(qn("r:embed")) in partes]
    if minusculo.endswith(".pdf"):
        import pypdfium2 as pdfium
        import pypdfium2.raw as pdfium_c

        documento = pdfium.PdfDocument(conteudo)
        imagens = []
        try:
            for pagina in documento:
                objetos = list(pagina.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE]))
                maior = max(objetos, key=_pixels) if objetos else None
                if maior is not None and _pixels(maior) >= 500_000:  # um print de tela; um ícone não conta
                    imagens.append(png(maior.get_bitmap(render=False).to_pil().convert("RGB")))
                else:
                    imagens.append(png(pagina.render(scale=2).to_pil().convert("RGB")))
        finally:
            documento.close()
        return imagens
    if minusculo.endswith((".png", ".jpg", ".jpeg")):
        return [conteudo]
    return []


# ---------------------------------------------------------------- Ecadnet

_LEITOR = r"""
param([string]$Lista)
$ErrorActionPreference = "Stop"
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType = WindowsRuntime]
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType = WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics, ContentType = WindowsRuntime]
$asTask = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Esperar($op, [Type]$tipo) { $t = $asTask.MakeGenericMethod($tipo).Invoke($null, @($op)); $t.Wait(-1) | Out-Null; $t.Result }
$motor = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
$saida = foreach ($imagem in (Get-Content -LiteralPath $Lista -Encoding UTF8)) {
  $arquivo = Esperar ([Windows.Storage.StorageFile]::GetFileFromPathAsync($imagem)) ([Windows.Storage.StorageFile])
  $fluxo = Esperar ($arquivo.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
  $decoder = Esperar ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($fluxo)) ([Windows.Graphics.Imaging.BitmapDecoder])
  $bitmap = Esperar ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
  $resultado = Esperar ($motor.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
  $fluxo.Dispose()
  ,@(foreach ($l in $resultado.Lines) {
    $r = @($l.Words | ForEach-Object { $_.BoundingRect })
    [pscustomobject]@{ texto = $l.Text; x0 = ($r | ForEach-Object { $_.X } | Measure-Object -Minimum).Minimum
      y0 = ($r | ForEach-Object { $_.Y } | Measure-Object -Minimum).Minimum; y1 = ($r | ForEach-Object { $_.Y + $_.Height } | Measure-Object -Maximum).Maximum }
  })
}
[Console]::OutputEncoding = [Text.Encoding]::UTF8
ConvertTo-Json -InputObject @($saida) -Depth 5 -Compress
"""


class LeitorIndisponivel(Exception):
    """O leitor de texto do Windows não respondeu (Windows sem o recurso, ou sem o idioma)."""


def ler_textos(imagens) -> list[list[dict]]:
    """O texto de cada imagem, linha a linha com a posição ({texto, x0, y0, y1}), pelo leitor do próprio Windows."""
    if not imagens:
        return []
    with tempfile.TemporaryDirectory() as pasta:
        pasta = Path(pasta)
        caminhos = []
        for i, imagem in enumerate(imagens):
            caminho = pasta / f"{i:04d}.png"
            imagem.save(caminho)
            caminhos.append(str(caminho))
        (pasta / "lista.txt").write_text("\n".join(caminhos), encoding="utf-8")
        (pasta / "ler.ps1").write_text(_LEITOR, encoding="utf-8-sig")
        try:
            resposta = subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(pasta / "ler.ps1"),
                 "-Lista", str(pasta / "lista.txt")], capture_output=True, timeout=60 + 20 * len(imagens),
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))  # fmt: skip
            dados = json.loads(resposta.stdout.decode("utf-8-sig") or "null")
        except (OSError, subprocess.TimeoutExpired, ValueError) as e:
            raise LeitorIndisponivel(f"{type(e).__name__}: {e}") from e
    if len(imagens) == 1 and (isinstance(dados, dict) or (isinstance(dados, list) and all(isinstance(d, dict) for d in dados))):
        dados = [dados]  # com uma imagem só, o PowerShell desembrulha a lista de fora
    if not isinstance(dados, list) or len(dados) != len(imagens):
        raise LeitorIndisponivel(resposta.stderr.decode("utf-8", "replace")[-300:] or "resposta vazia")
    # imagem com uma linha só também vem desembrulhada (o objeto, sem a lista)
    return [[l for l in ([linhas] if isinstance(linhas, dict) else linhas or []) if isinstance(l, dict)] for linhas in dados]


@dataclass
class PrintDoEcadnet:
    imagem: object
    barra: tuple | None = None  # (x0, topo, x1, base) da barra azul com o título da obra
    codigo: str = ""  # do endereço .../obras-musicais/<código>
    iswc: str = ""  # só os dígitos
    titulo: str = ""
    topo_da_pagina: int = 0  # onde a página começa, abaixo da barra de endereço do navegador
    obra: str = ""  # dígitos do ISWC da obra (da tabela) a que o print foi ligado
    como: str = ""  # "código", "ISWC", "título", "continuação", "fora da tabela" ou "pessoa"
    recorte: bytes | None = None  # o quadro da obra (PNG); a imagem inteira é solta depois do recorte


def barra_azul(imagem) -> tuple | None:
    """A barra azul do Ecadnet com o título da obra (x0, topo, x1, base), ou None (print da página rolada)."""
    import numpy as np

    a = np.asarray(imagem.convert("RGB")).astype(int)
    azul = (a[:, :, 2] > 220) & (a[:, :, 0] < 40) & (a[:, :, 1] > 60) & (a[:, :, 1] < 130)
    linhas = np.where(azul.sum(1) > 0.2 * a.shape[1])[0]
    if not len(linhas):
        return None
    grupo = linhas[: np.argmax(np.diff(np.append(linhas, linhas[-1] + 99)) > 1) + 1]  # a primeira faixa azul contínua
    colunas = np.where(azul[grupo[0]: grupo[-1] + 1].sum(0) > 0.5 * len(grupo))[0]
    return int(colunas[0]), int(grupo[0]), int(colunas[-1]), int(grupo[-1])


_TROCAS = str.maketrans({"o": "0", "O": "0", "l": "1", "I": "1", "|": "1", "z": "4", "g": "9", "S": "5", "B": "8", ",": "."})


def interpretar_print(p: PrintDoEcadnet, linhas: list[dict], escala: int):
    """Código (do endereço), ISWC e título lidos no print. `linhas` vêm da imagem ampliada `escala` vezes."""
    for linha in linhas:
        texto = linha.get("texto", "")
        m = re.search(r"obras-musicais/(\d{3,})", texto.replace(" ", ""))
        if m and not p.codigo:
            p.codigo = m.group(1)
            p.topo_da_pagina = int(linha.get("y1", 0) / escala) + 16
        m = re.search(r"T\s*-\s*([\dozgSBlI|.,\s]{9,20})-\s*([\dozgSBlI|])", texto)
        if m and p.barra and linha.get("y0", 0) / escala > p.barra[3] and not p.iswc:
            numero = digitos((m.group(1) + m.group(2)).translate(_TROCAS))
            if len(numero) == 10:
                p.iswc = numero
        if p.barra and p.barra[1] - 8 <= linha.get("y0", 0) / escala <= p.barra[3] + 8 and not p.titulo:
            p.titulo = re.sub(r"^[^A-Za-zÀ-ú0-9]+", "", re.sub(r"^u[.,•\s]*•+\s*", "", texto)).strip()


def ler_prints(arquivos: list[bytes], ao_avancar=None, por_vez=4) -> list[PrintDoEcadnet]:
    """Para cada print: a barra azul, o que está escrito no alto (endereço, título e ISWC, lido com a imagem ampliada
    3 vezes) e o recorte do quadro da obra. Abre `por_vez` prints de cada vez e guarda só o recorte."""
    from PIL import Image

    prints, colunas, escala = [], None, 3
    for inicio in range(0, len(arquivos), por_vez):
        lote = []
        for conteudo in arquivos[inicio: inicio + por_vez]:
            imagem = Image.open(io.BytesIO(conteudo)).convert("RGB")
            lote.append(PrintDoEcadnet(imagem=imagem, barra=barra_azul(imagem)))
        alto = []
        for p in lote:  # só o alto do print: o endereço e, abaixo da barra azul, a linha do ISWC
            altura = min(p.imagem.height, p.barra[3] + 5 * (p.barra[3] - p.barra[1])) if p.barra else min(p.imagem.height, 160)
            pedaco = p.imagem.crop((0, 0, p.imagem.width, altura))
            alto.append(pedaco.resize((pedaco.width * escala, pedaco.height * escala)))
        for p, linhas in zip(lote, ler_textos(alto)):
            interpretar_print(p, linhas, escala)
            if not p.topo_da_pagina:
                p.topo_da_pagina = round(87 * p.imagem.height / 1079)
            if p.barra:  # a continuação (sem a barra azul) usa a largura do quadro do print anterior
                colunas = (p.barra[0], p.barra[2])
            if colunas:
                p.recorte = png(recorte_do_ecadnet(p, colunas))
            p.imagem = None  # a tela inteira não é mais usada: não fica ocupando memória
        prints.extend(lote)
        if ao_avancar:
            ao_avancar(len(prints), len(arquivos))
    return prints


def ligar_prints(prints: list[PrintDoEcadnet], obras: list[ObraDoRelatorio], da_tabela: set[str]):
    """Liga cada print a uma obra da tabela: pelo código do endereço, pelo ISWC ou pelo título. Print sem a barra azul
    é a continuação do anterior (a mesma página, rolada). Print de obra fora da tabela fica sem obra."""
    por_codigo = {o.codigo: digitos(o.iswc) for o in obras}
    titulos = {}
    for o in obras:
        if digitos(o.iswc) in da_tabela:
            titulos.setdefault(_chave_do_titulo(o.titulo), set()).add(digitos(o.iswc))

    def pelo_iswc(numero):
        if numero in da_tabela:
            return numero
        perto = [i for i in da_tabela if len(i) == len(numero) and sum(a != b for a, b in zip(i, numero)) == 1]
        return perto[0] if len(perto) == 1 else ""

    def pelo_titulo(titulo):
        chave = _chave_do_titulo(titulo)
        if not chave:
            return ""
        notas = sorted(((SequenceMatcher(None, chave, k).ratio(), k) for k in titulos), reverse=True)
        if notas and notas[0][0] >= 0.85 and (len(notas) == 1 or notas[1][0] < notas[0][0] - 0.05) and len(titulos[notas[0][1]]) == 1:
            return next(iter(titulos[notas[0][1]]))
        return ""

    anterior = None
    for p in prints:
        if p.barra is None:
            if anterior is not None and (not p.codigo or p.codigo == anterior.codigo):
                p.obra, p.como = anterior.obra, ("continuação" if anterior.obra else anterior.como)
                if not p.codigo:
                    p.codigo = anterior.codigo
        elif p.codigo in por_codigo:
            p.obra, p.como = (por_codigo[p.codigo], "código") if por_codigo[p.codigo] in da_tabela else ("", "fora da tabela")
        elif p.iswc and pelo_iswc(p.iswc):
            p.obra, p.como = pelo_iswc(p.iswc), "ISWC"
        elif pelo_titulo(p.titulo):
            p.obra, p.como = pelo_titulo(p.titulo), "título"
        anterior = p if p.barra is not None or p.obra else anterior
    return prints


def _faixa_cinza(a, x0, x1):
    """As linhas (y) em que a faixa [x0, x1] é quase toda cinza claro: as divisórias e os cabeçalhos das tabelas."""
    import numpy as np

    faixa = a[:, x0: x1 + 1]
    cinza = (faixa.min(2) >= 190) & (faixa.max(2) <= 245) & (faixa.max(2) - faixa.min(2) <= 10)
    return np.where(cinza.mean(1) >= 0.9)[0]


def recorte_do_ecadnet(p: PrintDoEcadnet, colunas: tuple[int, int]):
    """O quadro da obra no print: da barra azul (ou do alto da página, na continuação) até o fim da última tabela."""
    import numpy as np

    a = np.asarray(p.imagem.convert("RGB")).astype(int)
    x0, x1 = colunas
    topo = (p.barra[1] - 2) if p.barra else p.topo_da_pagina
    # Divisória e cabeçalho de tabela têm o fundo branco da página logo embaixo. O cinza da barra de tarefas do Windows
    # e do fundo do site, embaixo do quadro, não: sem isso, eles passariam por linha de tabela.
    branco = (a[:, x0: x1 + 1].min(2) >= 250).mean(1)
    cinzas = [y for y in _faixa_cinza(a, x0, x1) if y > topo + 10 and y + 2 < len(branco) and branco[y + 1: y + 3].max() >= 0.8]
    base = (max(cinzas) + 2) if cinzas else p.imagem.height
    return p.imagem.crop((max(0, x0 - 1), max(0, topo), min(p.imagem.width, x1 + 2), min(p.imagem.height, base)))


# ---------------------------------------------------------------- a tabela


@dataclass
class LinhaDaTabela:
    titulo: str
    iswc: str
    relatorio: list[bytes] = field(default_factory=list)  # recortes em PNG
    iswc_imagem: bytes | None = None
    ecadnet: list[bytes] = field(default_factory=list)


def png(imagem) -> bytes:
    saida = io.BytesIO()
    imagem.save(saida, "PNG", optimize=True)
    return saida.getvalue()


def preparar(relatorio_pdf: bytes, iswc_pdfs: list[bytes]) -> tuple[list[LinhaDaTabela], list[ObraDoRelatorio], list[str]]:
    """As linhas da tabela (as obras do ISWC, na ordem do relatório, com o título do relatório), as obras do relatório
    e os avisos do que não foi achado."""
    obras = obras_do_relatorio(relatorio_pdf)
    if not obras:
        raise ValueError("não achei obras no relatório: ele precisa ser o relatório analítico do ECAD em PDF")
    do_iswc = {}
    for conteudo in iswc_pdfs:
        for numero, imagem in recortes_do_iswc(conteudo).items():
            do_iswc.setdefault(numero, imagem)
    vistas, linhas = set(), []
    for obra in obras:
        numero = digitos(obra.iswc)
        if numero not in do_iswc or numero in vistas:
            continue
        vistas.add(numero)
        linhas.append(LinhaDaTabela(titulo_da_tabela(obra.titulo), obra.iswc,
                                    [png(im) for im in recortes_do_relatorio(relatorio_pdf, obra)], png(do_iswc[numero])))
    avisos = [f"O ISWC T-{n[:3]}.{n[3:6]}.{n[6:9]}-{n[9:]} está no PDF do ISWC mas não no relatório: ficou fora da tabela."
              for n in sorted(do_iswc.keys() - vistas)]
    return linhas, obras, avisos


def preencher_ecadnet(linhas: list[LinhaDaTabela], prints: list[PrintDoEcadnet]):
    """Põe em cada linha os recortes dos prints ligados a ela, na ordem dos prints."""
    por_iswc = {digitos(l.iswc): l for l in linhas}
    for linha in linhas:
        linha.ecadnet = []
    for p in prints:
        if p.obra in por_iswc and p.recorte:
            por_iswc[p.obra].ecadnet.append(p.recorte)


def montar_linhas(relatorio_pdf: bytes, iswc_pdfs: list[bytes], prints: list[PrintDoEcadnet]) -> tuple[list[LinhaDaTabela], list[str]]:
    linhas, obras, avisos = preparar(relatorio_pdf, iswc_pdfs)
    ligar_prints(prints, obras, {digitos(l.iswc) for l in linhas})
    preencher_ecadnet(linhas, prints)
    return linhas, avisos


def montar_docx(linhas: list[LinhaDaTabela]) -> bytes:
    """O Word no formato do modelo do escritório: A4 deitado, margens de 1,27 cm, tabela de 4 colunas com borda,
    Bookman Old Style 10, cada recorte na largura da coluna."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Pt, Twips

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(8), 1.08
    secao = doc.sections[0]
    secao.orientation = WD_ORIENT.LANDSCAPE
    secao.page_width, secao.page_height = Twips(16838), Twips(11906)
    for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(secao, lado, Twips(720))

    tabela = doc.add_table(rows=1, cols=4)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabela.autofit = False
    propriedades = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borda = OxmlElement(f"w:{lado}")
        for chave, valor in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            borda.set(qn(f"w:{chave}"), valor)
        bordas.append(borda)
    propriedades.append(bordas)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    propriedades.append(layout)
    for coluna, largura in zip(tabela.columns, LARGURAS_DXA):
        coluna.width = Twips(largura)

    def texto(paragrafo, conteudo, negrito=False, italico=False):
        run = paragrafo.add_run(conteudo)
        run.font.name, run.font.size, run.bold, run.italic = FONTE, Pt(10), negrito, italico
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        return run

    for celula, largura, (titulo, explicacao) in zip(tabela.rows[0].cells, LARGURAS_DXA, CABECALHO):
        celula.width = Twips(largura)
        p = celula.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(p, titulo, negrito=True)
        p = celula.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if explicacao and explicacao.startswith("(site"):
            texto(p, "(")
            texto(p, "site", italico=True)
            texto(p, explicacao[5:])
        elif explicacao:
            texto(p, explicacao)

    for linha in linhas:
        celulas = tabela.add_row().cells
        for celula, largura in zip(celulas, LARGURAS_DXA):
            celula.width = Twips(largura)
        texto(celulas[0].paragraphs[0], linha.titulo)
        colunas = ((1, linha.relatorio), (2, [linha.iswc_imagem] if linha.iswc_imagem is not None else []), (3, linha.ecadnet))
        for indice, imagens in colunas:
            util = Emu((LARGURAS_DXA[indice] - 2 * MARGEM_DA_CELULA_DXA) * 635)
            for posicao, imagem in enumerate(imagens):
                p = celulas[indice].paragraphs[0] if posicao == 0 else celulas[indice].add_paragraph()
                if indice == 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.add_run().add_picture(io.BytesIO(imagem), width=util)
    doc.add_paragraph()
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
