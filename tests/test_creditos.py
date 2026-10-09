"""Testes sem rede do fluxo de créditos: importação do relatório, filtro de saída e classificador.

Todos os nomes, títulos e códigos daqui são inventados.
"""

import collections
import sys
import pandas as pd
import pytest

from creditos import classificador as c
from creditos.classificador import Config, Item, classificar
from creditos.ecad import RelatorioIlegivel, interpretar
from creditos.filtro import DadoSensivel, FiltroDeSaida
from creditos.modelo import Obra, Relatorio, Titular
from creditos.ubc import ler_planilha

# --- relatório do ECAD montado como o PDF entrega: palavras com a posição horizontal ---

X_OBRA = {"codigo": 30, "iswc": 77, "titulo": 142, "ass": 305, "situacao": 367, "tipo": 425, "nacional": 489, "inclusao": 525}
X_TITULAR = {"codigo": 30, "nome": 80, "pseudonimo": 226, "cae": 308, "ass": 369, "cat": 430, "pct": 445, "contrato": 468, "link": 560}


def _palavras(x, texto):
    """Espalha as palavras de um campo a partir de x, como no PDF (cerca de 6 pontos por letra)."""
    saida = []
    for palavra in str(texto).split():
        saida.append((x, palavra))
        x += 5 * len(palavra) + 4
    return saida


def linha(posicoes, **campos):
    return sorted(p for campo, texto in campos.items() if texto != "" for p in _palavras(posicoes[campo], texto))


CABECALHO = [
    [(130, "RELATÓRIO"), (204, "ANALÍTICO")],
    _palavras(43, "TITULAR: 1000 JOSE CARLOS SOUZA LIMA") + _palavras(324, "PSEUDÔNIMO: ZECA LIMA") + _palavras(471, "CATEGORIA: TODAS"),
    [(30, "CÓD."), (51, "OBRA"), (77, "ISWC"), (139, "TÍTULO"), (171, "PRINCIPAL"), (305, "ASS."), (325, "RESPONS."),
     (367, "SITUAÇÃO*"), (418, "TIPO"), (441, "OBRA"), (474, "NACIONAL"), (525, "INCLUSÃO")],
    [(473, "CONTRATO")],
    [(30, "CÓDIGO"), (80, "NOME"), (104, "DO"), (118, "TITULAR"), (226, "PSEUDÔNIMO"), (308, "CAE"), (369, "ASSOCIAÇÃO"),
     (425, "CAT"), (445, "(%)"), (469, "INÍCIO"), (501, "/"), (509, "FIM"), (544, "LINK")],
]
DONO = dict(codigo="1000", nome="JOSE CARLOS SOUZA LIMA", pseudonimo="ZECA LIMA", cae="00111.22.33.44", ass="AAA", cat="CA", link="1")
PARCEIRA = dict(codigo="2000", nome="MARIA APARECIDA DIAS", pseudonimo="CIDA DIAS", cae="00555.66.77.88", ass="BBB", cat="CA", link="2")
EDITORA = dict(codigo="3000", nome="EDITORA INVENTADA LTDA", pseudonimo="INVENTADA", cae="00999.00.11.22", ass="CCC", cat="E",
               pct="10,00", contrato="05/03/21", link="3")


def obra(codigo, titulo, iswc="T-000.000.001-0", situacao="LB", inclusao="10/06/2005"):
    return linha(X_OBRA, codigo=codigo, iswc=iswc, titulo=titulo, ass="AAA", situacao=situacao, tipo="ORIGINAL", nacional="SIM", inclusao=inclusao)


def rodape(total):
    return [_palavras(357, f"TOTAL DE OBRAS DO TITULAR: {total}"), _palavras(28, "19/02/2026 - 14:04") + [(540, "Página"), (563, "1")]]


def relatorio_de_teste() -> Relatorio:
    linhas = CABECALHO + [
        obra("11", "LIGUE O RADIO"),
        linha(X_TITULAR, **DONO, pct="100,"),
        obra("12", "COISA FEITA", iswc="- . . -"),
        linha(X_TITULAR, **DONO, pct="50,00"),
        linha(X_TITULAR, **PARCEIRA, pct="40,00"),
        linha(X_TITULAR, **EDITORA),
        obra("13", "BAILE NA SERRA", situacao="BL/DU"),
        linha(X_TITULAR, **DONO, pct="100,"),
        obra("14", "BAILE NA SERRA", situacao="BL/DU"),
        linha(X_TITULAR, **{**DONO, "pseudonimo": "ZECA DO BAILE"}, pct="50,00"),
        linha(X_TITULAR, **{**PARCEIRA, "pseudonimo": ""}, pct="50,00"),
        obra("15", "A DANCA DA PANELA"),
        linha(X_TITULAR, **DONO, pct="100,"),
        obra("16", "DANCA DA PANELLA"),
        linha(X_TITULAR, **DONO, pct="100,"),
    ] + rodape(6)
    return interpretar(linhas)


@pytest.fixture
def rel():
    return relatorio_de_teste()


# --- importação -------------------------------------------------------------


def test_le_obras_titulares_e_separa_nome_de_pseudonimo_pela_coluna(rel):
    assert (rel.codigo_titular, rel.nome_titular, rel.pseudonimo_titular) == ("1000", "JOSE CARLOS SOUZA LIMA", "ZECA LIMA")
    assert (len(rel.obras), rel.total_declarado, rel.emitido_em, rel.avisos) == (6, 6, "19/02/2026 14:04", [])
    feita = rel.obras[1]
    assert (feita.codigo, feita.titulo, feita.iswc, feita.situacao, feita.inclusao) == ("12", "COISA FEITA", "", "LB", "10/06/2005")
    assert [(t.nome, t.pseudonimo, t.categoria, t.percentual) for t in feita.titulares] == [
        ("JOSE CARLOS SOUZA LIMA", "ZECA LIMA", "CA", 50.0),
        ("MARIA APARECIDA DIAS", "CIDA DIAS", "CA", 40.0),
        ("EDITORA INVENTADA LTDA", "INVENTADA", "E", 10.0),
    ]
    assert feita.titulares[2].contrato == "05/03/21" and [a.nome for a in feita.autores] == ["JOSE CARLOS SOUZA LIMA", "MARIA APARECIDA DIAS"]
    assert rel.obras[0].titulares[0].percentual == 100.0  # o PDF corta "100,00" em "100,"
    assert rel.obras[3].titulares[1].pseudonimo == ""  # pseudônimo em branco não puxa o CAE para o lugar dele
    assert rel.obras[2].situacoes == ["BL", "DU"]


def test_situacao_impressa_fora_da_linha_da_obra_e_reaproveitada():
    sem = [p for p in obra("11", "LIGUE O RADIO") if p[1] != "LB"]
    antes = interpretar(CABECALHO + [[(367, "LB")], sem, linha(X_TITULAR, **DONO, pct="100,")] + rodape(1))
    depois = interpretar(CABECALHO + [sem, [(367, "DU")], linha(X_TITULAR, **DONO, pct="100,")] + rodape(1))
    assert (antes.obras[0].situacao, depois.obras[0].situacao) == ("LB", "DU") and not antes.avisos


def test_contagem_que_nao_bate_vira_aviso_e_arquivo_estranho_e_recusado():
    faltando = interpretar(CABECALHO + [obra("11", "LIGUE O RADIO"), linha(X_TITULAR, **DONO, pct="100,")] + rodape(2))
    assert any(a.startswith("CONTAGEM_DIVERGENTE") for a in faltando.avisos)
    sem_titular = interpretar(CABECALHO + [obra("11", "LIGUE O RADIO")] + rodape(1))
    assert any("sem nenhum titular" in a for a in sem_titular.avisos)
    with pytest.raises(RelatorioIlegivel):
        interpretar([[(30, "qualquer"), (80, "coisa")]])


def test_pseudonimos_do_titular_e_dos_coautores_viram_nomes_artisticos(rel):
    nomes = rel.nomes_artisticos()
    assert [(n["nome"], n["titular"], n["obras"]) for n in nomes] == [
        ("ZECA LIMA", True, 5), ("ZECA DO BAILE", True, 1), ("CIDA DIAS", False, 1),
    ]
    assert all(n["nome"] != "INVENTADA" for n in nomes)  # pseudônimo de editora não é nome artístico
    assert rel.pseudonimos()["1000"] == ["ZECA LIMA", "ZECA DO BAILE"]


def test_sinaliza_titulos_iguais_ou_parecidos_e_situacao_de_duplicidade(rel):
    grupos = [(g["motivo"], [o.codigo for o in g["obras"]]) for g in rel.duplicidades()]
    assert grupos == [("titulo", ["13", "14"]), ("titulo", ["15", "16"])]
    sozinha = interpretar(CABECALHO + [obra("11", "LIGUE O RADIO", situacao="HO"), linha(X_TITULAR, **DONO, pct="100,")] + rodape(1))
    assert [(g["motivo"], g["obras"][0].codigo) for g in sozinha.duplicidades()] == [("situacao", "11")]


def test_relatorio_vai_e_volta_de_json(rel):
    assert Relatorio.de_dict(rel.para_dict()) == rel


def test_planilha_no_formato_da_ubc_vira_o_mesmo_modelo():
    df = pd.DataFrame({
        "titulo": ["Ligue o Rádio", "Coisa Feita", ""],
        "compositor": ["José Carlos Souza Lima"] * 3,
        "coautores": ["", "Maria Aparecida Dias, Outro Parceiro", ""],
        "iswc": ["T-000.000.001-0", "", ""],
        "cod_ecad": ["11", "12", ""],
        "data_cadastro": ["2005-06-10", "2006-01-02", ""],
        "situacao": ["LIBERADA", "EM DUPLICIDADE", ""],
        "editores": ["", "Editora Inventada Ltda", ""],
    })
    rel = ler_planilha(df)
    assert (rel.origem, rel.nome_titular, len(rel.obras)) == ("UBC", "José Carlos Souza Lima", 2)
    feita = rel.obras[1]
    assert (feita.codigo, feita.situacao, feita.inclusao) == ("12", "DU", "2006-01-02")
    assert [a.nome for a in feita.autores] == ["José Carlos Souza Lima", "Maria Aparecida Dias", "Outro Parceiro"]
    assert [t.nome for t in feita.titulares if not t.e_autor] == ["Editora Inventada Ltda"]
    with pytest.raises(ValueError):
        ler_planilha(pd.DataFrame({"compositor": ["x"]}))


# --- filtro de saída --------------------------------------------------------


def test_filtro_bloqueia_dado_sensivel_e_libera_titulo_e_autores(rel):
    liberados = [o.titulo for o in rel.obras] + [x for o in rel.obras for a in o.autores for x in (a.nome, a.pseudonimo)]
    filtro = FiltroDeSaida(rel.termos_sensiveis(), liberados)
    for texto in ["https://api.exemplo/search?q=LIGUE O RADIO ZECA LIMA", "q=coisa feita Maria Aparecida Dias", "isrc:BRXXX2100001"]:
        assert filtro.motivo(texto) == ""
    for texto in ["q=00111.22.33.44", "obra 12 codigo=3000", "q=Editora Inventada Ltda", "partes 50,00", "25 %", "cpf 123.456.789-09",
                  "inclusao=10/06/2005", "contrato 05/03/21"]:
        assert filtro.motivo(texto), texto
    with pytest.raises(DadoSensivel):
        filtro.conferir("https://api.exemplo/search", {"q": "LIGUE O RADIO"}, "cae=00111.22.33.44")
    assert filtro.bloqueadas == 1


def test_filtro_nao_confunde_codigo_curto_nem_numero_dentro_de_outro(rel):
    filtro = FiltroDeSaida(rel.termos_sensiveis())
    assert filtro.motivo("track/11") == "" and filtro.motivo("album/130001") == ""  # "11", "12"... são curtos demais para vigiar


# --- classificador: estado da coleta ---------------------------------------


def test_erro_tecnico_nunca_vira_sem_creditos(rel):
    assert classificar(Item(titulo="Ligue o Rádio", coleta="erro", erro="timeout"), rel).status == c.ERRO_TECNICO
    assert classificar(Item(titulo="Ligue o Rádio", coleta="indisponivel"), rel).status == c.INDISPONIVEL
    assert classificar(Item(titulo="Ligue o Rádio", coleta="qualquer coisa"), rel).status == c.ERRO_TECNICO


def test_link_que_abre_outro_interprete_e_link_divergente(rel):
    item = Item(titulo="Ligue o Rádio", interprete="Banda Estranha", interprete_esperado="Zeca Lima", creditos=[])
    r = classificar(item, rel)
    assert r.status == c.LINK_DIVERGENTE and r.revisar


# --- classificador: título --------------------------------------------------


def test_titulo_exato_ignora_caixa_acento_parenteses_e_sufixo_de_versao(rel):
    for titulo in ["Ligue o Rádio", "LIGUE O RADIO (Ao Vivo)", "Ligue o Rádio - Ao Vivo", "Ligue o Rádio - Remasterizado 2010"]:
        r = classificar(Item(titulo=titulo, interprete="Zeca Lima", creditos=["Zeca Lima"]), rel)
        assert (r.status, r.titulo, r.obras) == (c.OK, "exato", ["11"]), titulo


def test_titulo_fora_do_relatorio_nao_gera_acusacao(rel):
    r = classificar(Item(titulo="Outra Música Qualquer", interprete="Zeca Lima"), rel)
    assert r.status == c.FORA_DO_REPERTORIO and not r.revisar


def test_titulo_aproximado_nunca_sustenta_resultado_negativo(rel):
    # "Largue o Rádio" casa com "LIGUE O RADIO" por semelhança: pode ser outra música.
    sem = classificar(Item(titulo="Largue o Rádio", interprete="Zeca Lima"), rel)
    assert (sem.status, sem.sugestao, sem.revisar, sem.titulo) == (c.TITULO_APROXIMADO, c.SEM_CREDITOS, True, "aproximado")
    errado = classificar(Item(titulo="Coisa Feia", interprete="Zeca Lima", creditos=["Zeca Lima Banda"]), rel, Config(interpretes=["Zeca Lima Banda"]))
    assert (errado.status, errado.sugestao) == (c.TITULO_APROXIMADO, c.VIOLACAO)
    # Com o titular creditado, o título aproximado não atrapalha: o resultado é positivo.
    ok = classificar(Item(titulo="Largue o Rádio", interprete="Zeca Lima", creditos=["Zeca Lima"]), rel)
    assert ok.status == c.OK and "aproximado" in ok.fundamento


def test_medley_e_dividido_e_vai_para_revisao(rel):
    for titulo in ["Ligue o Rádio / Outra Música", "Pot-Pourri: Ligue o Rádio - Coisa Feita", "Medley Ligue o Rádio"]:
        r = classificar(Item(titulo=titulo, interprete="Zeca Lima", creditos=["Fulano"]), rel)
        assert (r.status, r.revisar) == (c.MEDLEY, True), titulo
        assert "LIGUE O RADIO" in r.fundamento
    assert classificar(Item(titulo="Uma / Outra", interprete="Zeca Lima"), rel).status == c.FORA_DO_REPERTORIO


# --- classificador: crédito exibido ----------------------------------------


@pytest.mark.parametrize("credito", ["Zeca Lima", "ZECA LIMA", "José Carlos Souza Lima", "José Lima", "Jose Souza Lima", "Zeca do Baile"])
def test_titular_creditado_pelo_nome_pseudonimo_ou_nome_abreviado_e_ok(rel, credito):
    r = classificar(Item(titulo="Ligue o Rádio", interprete="Banda Qualquer", creditos=[credito]), rel)
    assert (r.status, r.revisar) == (c.OK, False)


def test_sem_creditos_exige_tela_lida_e_vinculo_com_o_titular(rel):
    r = classificar(Item(titulo="Coisa Feita", interprete="Zeca Lima", creditos=[]), rel)
    assert (r.status, r.revisar, r.coautores_omitidos) == (c.SEM_CREDITOS, False, ["MARIA APARECIDA DIAS"])
    # Mesmo título, sem crédito, com um intérprete que ninguém ligou ao titular: pode ser música de outra pessoa.
    estranho = classificar(Item(titulo="Coisa Feita", interprete="Banda Estranha", creditos=[]), rel)
    assert (estranho.status, estranho.revisar) == (c.HOMONIMA, True)
    for item, config in [
        (Item(titulo="Coisa Feita", interprete="Banda Estranha", vinculo=True), Config()),
        (Item(titulo="Coisa Feita", interprete="Banda Estranha"), Config(interpretes=["Banda Estranha"])),
        (Item(titulo="Coisa Feita", interprete="Cida Dias"), Config()),  # a coautora cantando
    ]:
        assert classificar(item, rel, config).status == c.SEM_CREDITOS


def test_banda_ou_interprete_no_campo_de_autor_e_violacao_sem_pedir_revisao(rel):
    r = classificar(Item(titulo="Coisa Feita", interprete="Banda do Baile", creditos=["Banda do Baile"]), rel,
                    Config(interpretes=["Banda do Baile"]))
    assert (r.status, r.revisar) == (c.VIOLACAO, False)
    assert "pessoa física" in r.fundamento and r.coautores_omitidos == ["JOSE CARLOS SOUZA LIMA", "MARIA APARECIDA DIAS"]


def test_interprete_sem_vinculo_que_credita_a_si_mesmo_e_homonimo_e_nunca_violacao(rel):
    # Caso real da primeira rodada: um cantor famoso com uma música de mesmo título, composta por ele.
    r = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Famoso", creditos=["Cantor Famoso"]), rel)
    assert (r.status, r.revisar, r.vinculo) == (c.HOMONIMA, True, "")
    varios = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Famoso", creditos=["Cantor Famoso, Outro Autor"]), rel)
    assert varios.status == c.HOMONIMA


def test_vinculo_pelo_interprete_exige_o_nome_inteiro(rel):
    # "Zeca" não é "Zeca Lima", e "Maria Dias" (outra pessoa) não é a coautora "Maria Aparecida Dias".
    for interprete in ["Zeca", "Lima", "Cida"]:
        assert classificar(Item(titulo="Coisa Feita", interprete=interprete), rel).status == c.HOMONIMA, interprete
    r = classificar(Item(titulo="Coisa Feita", interprete="ZECA LIMA"), rel)
    assert r.status == c.SEM_CREDITOS and "Zeca Lima".upper() in r.vinculo.upper()


def test_nome_parecido_que_e_de_outro_autor_do_relatorio_nao_e_grafia():
    # Dois irmãos no relatório: a obra é só de um, e a plataforma credita o outro.
    irmao = Titular(codigo="1001", nome="JOAO CARLOS SOUZA LIMA", pseudonimo="JOAO LIMA", categoria="CA")
    dono = Titular(codigo="1000", nome="JOSE CARLOS SOUZA LIMA", pseudonimo="ZECA LIMA", categoria="CA")
    rel = Relatorio(codigo_titular="1000", nome_titular=dono.nome, pseudonimo_titular=dono.pseudonimo, obras=[
        Obra(codigo="1", titulo="SO DELE", titulares=[dono]),
        Obra(codigo="2", titulo="DOS DOIS", titulares=[dono, irmao]),
    ])
    r = classificar(Item(titulo="Só Dele", interprete="Banda do Baile", creditos=["João Carlos Souza Lima"]), rel)
    assert (r.status, r.revisar) == (c.VIOLACAO, True) and "outra pessoa" in r.fundamento
    assert r.coautores_omitidos == ["JOSE CARLOS SOUZA LIMA"] and r.vinculo == "autor do relatório no crédito exibido"
    # A grafia errada do próprio titular continua sendo divergência.
    assert classificar(Item(titulo="Só Dele", interprete="Zeca Lima", creditos=["José Carlos Sousa Lima"]), rel).status == c.GRAFIA


def test_nome_de_pessoa_desconhecido_e_violacao_mas_sempre_com_revisao(rel):
    r = classificar(Item(titulo="Ligue o Rádio", interprete="Zeca Lima", creditos=["Santino Prado"]), rel)
    assert (r.status, r.revisar) == (c.VIOLACAO, True) and "outro nome artístico" in r.fundamento
    # Sem vínculo nenhum com o titular, nem violação é: pode ser obra homônima.
    r = classificar(Item(titulo="Ligue o Rádio", interprete="Banda Estranha", creditos=["Santino Prado"]), rel)
    assert (r.status, r.revisar) == (c.HOMONIMA, True)


def test_so_o_coautor_creditado_e_violacao_contra_o_titular(rel):
    r = classificar(Item(titulo="Coisa Feita", interprete="Banda Estranha", creditos=["Cida Dias"]), rel)
    assert (r.status, r.revisar, r.coautores_omitidos) == (c.VIOLACAO, False, ["JOSE CARLOS SOUZA LIMA"])


def test_coautor_omitido_sai_na_nota_mesmo_quando_o_titular_esta_creditado(rel):
    r = classificar(Item(titulo="Coisa Feita", interprete="Zeca Lima", creditos=["Zeca Lima"]), rel)
    assert r.status == c.OK and r.coautores_omitidos == ["MARIA APARECIDA DIAS"] and "MARIA APARECIDA DIAS" in r.fundamento
    completo = classificar(Item(titulo="Coisa Feita", interprete="Zeca Lima", creditos=["Zeca Lima, Cida Dias"]), rel)
    assert completo.status == c.OK and completo.coautores_omitidos == []
    extra = classificar(Item(titulo="Coisa Feita", interprete="Zeca Lima", creditos=["Zeca Lima", "Cida Dias", "Santino Prado"]), rel)
    assert (extra.status, extra.revisar) == (c.OK, True) and "Santino Prado" in extra.fundamento


@pytest.mark.parametrize("credito", ["Zecca Lima", "Zeca Lyma", "José Carlos Sousa Lima"])
def test_grafia_errada_do_nome_ou_pseudonimo_e_divergencia(rel, credito):
    r = classificar(Item(titulo="Ligue o Rádio", interprete="Zeca Lima", creditos=[credito]), rel)
    assert (r.status, r.revisar) == (c.GRAFIA, True)


def test_nome_artistico_so_vale_como_ok_depois_de_confirmado(rel):
    item = Item(titulo="Ligue o Rádio", interprete="Zeca Lima", creditos=["Carlinhos do Acordeon"])
    assert classificar(item, rel).status == c.VIOLACAO  # desconhecido: violação com revisão
    a_confirmar = classificar(item, rel, Config(nomes_a_confirmar=["Carlinhos do Acordeon"]))
    assert (a_confirmar.status, a_confirmar.revisar) == (c.OK_VARIANTE, True)
    assert classificar(item, rel, Config(nomes_confirmados=["Carlinhos do Acordeon"])).status == c.OK
    # Combinação de pseudônimo com sobrenome ("Zeca Souza") identifica o autor, mas pede confirmação.
    combinado = classificar(Item(titulo="Ligue o Rádio", interprete="Zeca Lima", creditos=["Zeca Souza"]), rel)
    assert (combinado.status, combinado.revisar) == (c.OK_VARIANTE, True)


def test_registros_em_duplicidade_com_autorias_diferentes_sao_todos_considerados(rel):
    # "BAILE NA SERRA" tem dois registros: um só do titular, outro com a coautora.
    r = classificar(Item(titulo="Baile na Serra", interprete="Zeca Lima", creditos=["Maria Aparecida Dias"]), rel)
    assert (r.status, r.obras) == (c.VIOLACAO, ["13", "14"]) and "2 registros" in r.fundamento
    ok = classificar(Item(titulo="Baile na Serra", interprete="Zeca Lima", creditos=["Zeca do Baile"]), rel)
    assert ok.status == c.OK and ok.coautores_omitidos == ["MARIA APARECIDA DIAS"]


def test_titulo_do_relatorio_sem_o_titular_entre_os_autores_nao_acusa_ninguem():
    rel = Relatorio(codigo_titular="1000", nome_titular="JOSE CARLOS SOUZA LIMA", obras=[
        Obra(codigo="9", titulo="MUSICA ALHEIA", titulares=[Titular(codigo="2000", nome="MARIA APARECIDA DIAS", categoria="CA")]),
    ])
    assert classificar(Item(titulo="Música Alheia", interprete="Zeca Lima"), rel).status == c.HOMONIMA


# --- Deezer e pipeline, com a plataforma trocada por respostas fixas --------

import json as _json

from creditos import pipeline, saida
from creditos.deezer import Album, Deezer, ErroDeColeta, Faixa, interpretar_pagina


def pagina_de_album(album_id, faixas, declaradas=None, selo="Selo Inventado"):
    estado = {
        "DATA": {"ALB_ID": str(album_id), "ALB_TITLE": f"Álbum {album_id}", "ART_NAME": "Banda do Baile", "LABEL_NAME": selo,
                 "UPC": "1", "NUMBER_TRACK": str(len(faixas) if declaradas is None else declaradas), "ORIGINAL_RELEASE_DATE": "2010-05-01"},
        "SONGS": {"data": [
            {"SNG_ID": str(i), "SNG_TITLE": t, "VERSION": v, "ART_NAME": "Banda do Baile", "ISRC": f"BRXXX10000{i:02d}",
             "SNG_CONTRIBUTORS": papeis, "DURATION": "180"}
            for i, (t, v, papeis) in enumerate(faixas, start=1)
        ]},
    }
    return f"<html><script>window.__DZR_APP_STATE__ = {_json.dumps(estado)}</script><p>resto</p></html>"


def test_le_creditos_isrc_e_fornecedor_do_estado_da_pagina():
    album = interpretar_pagina(pagina_de_album(7, [
        ("Ligue o Rádio", "", {"composer": ["Zeca Lima"], "author": ["Zeca Lima", "Cida Dias"], "main_artist": ["Banda do Baile"]}),
        ("Coisa Feita", "(Ao Vivo)", []),  # sem contribuidores, a página manda uma lista vazia
    ]), "7")
    assert (album.fornecedor, album.lancamento, album.faixas_declaradas, album.aviso) == ("Selo Inventado", "2010-05-01", 2, "")
    um, dois = album.faixas
    assert (um.creditos, um.isrc, um.link) == (["Zeca Lima", "Cida Dias"], "BRXXX1000001", "https://www.deezer.com/track/1")
    assert (dois.titulo, dois.creditos) == ("Coisa Feita (Ao Vivo)", [])


def test_guardas_da_pagina_erro_nunca_vira_album_sem_creditos():
    with pytest.raises(ErroDeColeta, match="em vez do 8"):
        interpretar_pagina(pagina_de_album(7, []), "8")  # a página devolveu outro álbum
    with pytest.raises(ErroDeColeta, match="estado"):
        interpretar_pagina("<html>faça login</html>", "7")
    faltando = interpretar_pagina(pagina_de_album(7, [("Ligue o Rádio", "", {})], declaradas=2), "7")
    assert faltando.aviso.startswith("CONTAGEM_DIVERGENTE")


class DeezerFalso(Deezer):
    """Mesma classe, sem rede: discografias, buscas e álbuns vêm de dicionários."""

    def __init__(self, discografias=None, buscas=None, albuns=None, filtro=None):
        super().__init__(filtro=filtro)
        self.discografias, self.buscas, self.lidos = discografias or {}, buscas or {}, albuns or {}
        self.pedidos = []

    def artistas(self, nome, maximo=3):
        self.pedidos.append(("artista", nome))
        return [{"id": nome, "nome": nome, "albuns": 1, "fas": 1}] if nome in self.discografias else []

    def albuns_do_artista(self, artista_id):
        return self.discografias[artista_id]

    def buscar_faixas(self, consulta, maximo=300):
        self.pedidos.append(("busca", consulta))
        return self.buscas.get(consulta, [])

    def album(self, album_id):
        return self.lidos.get(album_id) or Album(id=album_id, erro="HTTP 503 após 3 tentativas")


def _album(album_id, *faixas, selo="Selo Inventado", lancamento="2010-05-01"):
    return Album(id=album_id, titulo=f"Álbum {album_id}", fornecedor=selo, lancamento=lancamento, faixas_declaradas=len(faixas),
                 faixas=[Faixa(id=f"{album_id}{i}", titulo=t, interprete=quem, creditos=list(cred), isrc=f"BR{album_id}{i}")
                         for i, (t, quem, cred) in enumerate(faixas)])


def _coletar(rel, deezer, **config):
    return pipeline.coletar_deezer(rel, Config(nomes_confirmados=["ZECA LIMA"], **config), deezer)


def test_pipeline_classifica_cada_gravacao_e_deduplica_por_link(rel):
    deezer = DeezerFalso(
        discografias={"ZECA LIMA": ["a1", "a2"]},
        albuns={
            "a1": _album("a1", ("Ligue o Rádio", "Zeca Lima", ["Zeca Lima"]), ("Coisa Feita", "Zeca Lima", []),
                         ("Música de Outro", "Zeca Lima", ["Outro"])),
            "a2": _album("a2", ("Ligue o Rádio", "Zeca Lima", []), selo="Agregador Inventado"),  # a mesma obra, outro cadastro
        },
    )
    coleta = _coletar(rel, deezer)
    assert [(g.obra, g.classificacao.status, g.fornecedor) for g in coleta.gravacoes] == [
        ("LIGUE O RADIO", c.OK, "Selo Inventado"), ("COISA FEITA", c.SEM_CREDITOS, "Selo Inventado"),
        ("LIGUE O RADIO", c.SEM_CREDITOS, "Agregador Inventado"),
    ]
    assert (coleta.fora_do_repertorio, coleta.albuns_lidos, coleta.albuns_com_erro) == (1, 2, [])
    assert all(g.isrc and g.link.startswith("https://www.deezer.com/track/") and g.coletado_em for g in coleta.gravacoes)


def test_album_que_nao_abre_vira_erro_no_resumo_e_nunca_gravacao_sem_credito(rel):
    coleta = _coletar(rel, DeezerFalso(discografias={"ZECA LIMA": ["a1"]}))
    assert coleta.gravacoes == [] and coleta.albuns_com_erro == [("https://www.deezer.com/br/album/a1", "HTTP 503 após 3 tentativas")]


def test_homonimo_de_terceiro_vai_para_pendencias_e_nao_entra_na_planilha(rel):
    achado = {"id": "x", "titulo": "Ligue o Rádio", "interprete": "Cantor Famoso", "album_id": "b1"}
    deezer = DeezerFalso(buscas={"LIGUE O RADIO": [achado]},
                         albuns={"b1": _album("b1", ("Ligue o Rádio", "Cantor Famoso", ["Cantor Famoso"]))})
    coleta = _coletar(rel, deezer)
    assert [(g.interprete, g.classificacao.status, g.origem) for g in coleta.gravacoes] == [
        ("Cantor Famoso", c.HOMONIMA, "busca por título (LIGUE O RADIO)")]
    linhas = saida.linhas_por_obra(rel, [coleta])
    assert all(l["interprete"] == "(nenhuma gravação encontrada)" for l in linhas)


def test_interprete_recorrente_ganha_vinculo_inferido_sempre_marcado_para_confirmar(rel):
    banda = "Banda do Baile"
    deezer = DeezerFalso(
        buscas={t: [{"id": t, "titulo": t, "interprete": banda, "album_id": "d1"}] for t in ["LIGUE O RADIO"]},
        albuns={"d1": _album("d1", ("Ligue o Rádio", banda, ["Zeca Lima"]), ("Coisa Feita", banda, ["Zeca Lima"]),
                             ("Baile na Serra", banda, []), ("A Dança da Panela", banda, []))},
    )
    coleta = _coletar(rel, deezer)
    assert coleta.interpretes_inferidos == [{"nome": banda, "obras": 4, "com_autor": 2}]
    sem = [g for g in coleta.gravacoes if g.classificacao.status == c.SEM_CREDITOS]
    assert len(sem) == 2 and all(g.classificacao.revisar and g.classificacao.presumido and "PRESUMIDO" in g.classificacao.fundamento for g in sem)
    # Com menos títulos, ou sem nenhum autor reconhecido, não há inferência: fica como possível homônimo.
    pouco = DeezerFalso(buscas=deezer.buscas, albuns={"d1": _album("d1", ("Ligue o Rádio", banda, ["Zeca Lima"]), ("Baile na Serra", banda, []))})
    assert [g.classificacao.status for g in _coletar(rel, pouco).gravacoes] == [c.OK, c.HOMONIMA]


def test_lancamento_muito_anterior_ao_cadastro_pede_revisao_sem_revelar_a_data(rel):
    deezer = DeezerFalso(discografias={"ZECA LIMA": ["a1"]},
                         albuns={"a1": _album("a1", ("Coisa Feita", "Zeca Lima", []), lancamento="1990-01-01")})
    (g,) = _coletar(rel, deezer).gravacoes
    assert g.classificacao.status == c.SEM_CREDITOS and g.classificacao.revisar
    assert "homônima" in g.classificacao.fundamento and "2005" not in g.classificacao.fundamento


def test_nenhuma_requisicao_leva_dado_sensivel(rel):
    config = Config(nomes_confirmados=["ZECA LIMA"])
    deezer = DeezerFalso(filtro=pipeline.filtro_do_relatorio(rel, config))
    pipeline.coletar_deezer(rel, config, deezer, nomes_dos_coautores=["CIDA DIAS"])
    assert ("artista", "CIDA DIAS") in deezer.pedidos and ("busca", "COISA FEITA ZECA LIMA") in deezer.pedidos
    sensiveis = rel.termos_sensiveis()
    assert not any(termo in consulta for _, consulta in deezer.pedidos for termo in sensiveis if len(termo) >= 4)
    assert all(deezer.filtro.motivo(consulta) == "" for _, consulta in deezer.pedidos)


def test_planilha_no_formato_do_escritorio(rel):
    import io
    from openpyxl import load_workbook

    deezer = DeezerFalso(discografias={"ZECA LIMA": ["a1", "a2"]}, albuns={
        "a1": _album("a1", ("Ligue o Rádio", "Zeca Lima", ["Zeca Lima"]), ("Coisa Feita", "Zeca Lima", [])),
        "a2": _album("a2", ("Coisa Feita", "Zeca Lima", [])),
    })
    coleta = _coletar(rel, deezer)
    wb = load_workbook(io.BytesIO(saida.planilha(rel, [coleta])))
    assert wb.sheetnames == ["Obras", "Gravações", "Pendências - homônimos", "Resumo"]
    ws = wb["Obras"]
    assert ws["A1"].value == "JOSE CARLOS SOUZA LIMA - ZECA LIMA" and ws["A1"].fill.fgColor.rgb.endswith("1B3A6B")
    assert [x.value for x in ws[2]] == ["Obras", "Intérpretes"] + saida.PLATAFORMAS
    deezer_col = saida.PLATAFORMAS.index("DEEZER") + 3
    por_obra = {ws.cell(r, 1).value: ws.cell(r, deezer_col).value for r in range(3, ws.max_row)}
    assert por_obra["LIGUE O RADIO"] == "tem créditos"
    assert por_obra["COISA FEITA"] == "https://www.deezer.com/track/a11 -- https://www.deezer.com/track/a20"
    assert por_obra["BAILE NA SERRA"] == "N/A"
    assert ws.cell(ws.max_row, deezer_col).value == "2 obras"
    assert ws.cell(3, saida.PLATAFORMAS.index("NAPSTER") + 3).value == "plataforma encerrada"
    assert ws.cell(3, saida.PLATAFORMAS.index("SPOTIFY") + 3).value == saida.NAO_COLETADO
    gravacoes = [[x.value for x in linha] for linha in wb["Gravações"].iter_rows(min_row=2)]
    assert len(gravacoes) == 3 and {g[10] for g in gravacoes} == {c.OK, c.SEM_CREDITOS}
    assert "Gravações do repertório: verificadas / coletadas: 3 / 3" in saida.resumo_em_texto(rel, [coleta])


# --- YouTube Music: interpretação do que a tela mostra (sem navegador) -------

from datetime import datetime
from zoneinfo import ZoneInfo

from creditos import captura, provas, ytmusic

MENU_SEM = ["Criar mix", "Tocar a seguir", "Adicionar à fila", "Salvar na playlist", "Ir para a página do artista", "Compartilhar", "Denunciar"]
MENU_COM = MENU_SEM[:5] + ["Mostrar créditos da música"] + MENU_SEM[5:]


def test_menu_sem_o_item_de_creditos_so_vale_se_o_menu_da_faixa_abriu_mesmo():
    assert ytmusic.interpretar_menu(MENU_SEM) == "sem_creditos"
    assert ytmusic.interpretar_menu(MENU_COM) == "com_creditos"
    assert ytmusic.interpretar_menu(["Ver créditos", "Compartilhar"]) == "com_creditos"
    # Menu que não abriu, ou outro menu da página (conta, configurações): nunca é "sem créditos".
    for itens in [None, [], ["Fazer login", "Configurações", "Termos de Serviço"]]:
        assert ytmusic.interpretar_menu(itens) == "invalido"


def test_janela_de_creditos_e_lida_por_secao_com_cada_nome_no_seu_elemento():
    secoes = [["rotulo", "Interpretação de"], ["nome", "Banda do Baile"], ["rotulo", "Composição de"], ["nome", "Zeca Lima"],
              ["nome", "Cida Dias"], ["rotulo", "Metadados de música fornecidos por"], ["nome", "Selo Inventado"]]
    por_rotulo = ytmusic.interpretar_janela(secoes)
    assert por_rotulo["Composição de"] == ["Zeca Lima", "Cida Dias"]
    assert ytmusic.autores_e_fornecedor(por_rotulo) == (["Zeca Lima", "Cida Dias"], "Selo Inventado")
    assert ytmusic.autores_e_fornecedor(ytmusic.interpretar_janela(secoes[:2])) == ([], "")  # janela sem "Composição de"


def test_linha_do_player_e_resultados_da_busca():
    assert ytmusic.interpretar_linha_do_player("Banda do Baile\n • \nÁlbum Tal\n • \n2012") == ("Banda do Baile", "Álbum Tal • 2012")
    itens = [
        {"titulo": "Ligue o Rádio", "colunas": ["Ligue o Rádio", "Banda do Baile • Álbum Tal • 3:09"], "link": "watch?v=AbC_123-xyz&list=RD", "artistas": ["Banda do Baile"]},
        {"titulo": "Um álbum", "colunas": ["Um álbum"], "link": "", "artistas": []},
    ]
    (cand,) = ytmusic.interpretar_resultados(itens, "Música")
    assert (cand.video, cand.titulo, cand.interprete, cand.detalhe, cand.tipo) == ("AbC_123-xyz", "Ligue o Rádio", "Banda do Baile", "Álbum Tal • 3:09", "Música")
    assert cand.link == "https://music.youtube.com/watch?v=AbC_123-xyz"


def test_captcha_login_e_consentimento_param_a_fila():
    assert ytmusic.sinal_de_bloqueio("https://www.google.com/sorry/index?continue=x")
    assert ytmusic.sinal_de_bloqueio("https://accounts.google.com/ServiceLogin")
    assert ytmusic.sinal_de_bloqueio("https://music.youtube.com/watch?v=x", "Faça login para confirmar que você não é um robô")
    assert ytmusic.sinal_de_bloqueio("https://music.youtube.com/watch?v=x", "Banda do Baile · Compartilhar") == ""


def test_nome_do_arquivo_e_carimbo_da_captura():
    quando = datetime(2026, 3, 2, 15, 56, 7, tzinfo=ZoneInfo("America/Sao_Paulo"))
    assert captura.nome_base("A Dança da Panela", "Banda do Baile", "YouTube Music", quando) == "a-danca-da-panela_banda-do-baile_youtube-music_2026-03-02_15h56"
    carimbo = captura.texto_do_carimbo("https://music.youtube.com/watch?v=x", quando, "abc", "Mon, 02 Mar 2026 18:56:07 GMT")
    assert "URL: https://music.youtube.com/watch?v=x" in carimbo and "02/03/2026 15:56:07 -0300 (America/Sao_Paulo)" in carimbo
    assert "Captura: abc" in carimbo and "HTTP Date): Mon, 02 Mar 2026 18:56:07 GMT" in carimbo


class YouTubeFalso:
    """Busca e leitura com respostas fixas, no lugar do navegador."""
    filtro, navegacoes = None, 0

    def __init__(self, buscas, leituras, bloquear_em=None):
        self.buscas, self.leituras, self.bloquear_em, self.lidos = buscas, leituras, bloquear_em, []

    def buscar(self, consulta, abas):
        return [ytmusic.Candidato(*c_) for c_ in self.buscas.get(consulta, [])]

    def ler(self, video, obra, esperado=None):
        if video == self.bloquear_em:
            raise ytmusic.Bloqueio("a plataforma desviou para www.google.com")
        self.lidos.append(video)
        return self.leituras[video]


def _leitura(video, titulo, interprete, menu, creditos=(), fornecedor=""):
    return ytmusic.Leitura(video=video, titulo=titulo, interprete=interprete, menu=menu, creditos=list(creditos),
                           tem_item_de_creditos=bool(creditos), fornecedor=fornecedor, provas=[f"prova-{video}"])


def _youtube(rel, ytm):
    return pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Banda do Baile"])


def test_youtube_so_abre_faixa_de_quem_tem_ligacao_com_o_titular(rel):
    ytm = YouTubeFalso(
        buscas={"LIGUE O RADIO Banda do Baile": [
            ("v1", "Ligue o Rádio", "Banda do Baile", "Álbum • 3:00", "Música"),
            ("v2", "Banda do Baile - Ligue o Rádio", "Canal de Fã", "5 mil visualizações", "Vídeo"),
            ("v3", "Ligue o Rádio", "Cantor Famoso", "Outro Álbum", "Música"),  # homônimo: conta, mas não abre
            ("v4", "Outra Música", "Banda do Baile", "Álbum", "Música"),
        ]},
        leituras={"v1": _leitura("v1", "Ligue o Rádio", "Banda do Baile", MENU_COM, ["Zeca Lima"], "Selo Inventado"),
                  "v2": _leitura("v2", "Banda do Baile - Ligue o Rádio", "Canal de Fã", MENU_SEM)},
    )
    coleta = _youtube(rel, ytm)
    assert ytm.lidos == ["v1", "v2"] and coleta.sem_vinculo_nao_abertos == 1
    pendencia, ok, sem = coleta.gravacoes
    # O homônimo não é aberto, mas fica nas pendências, com o canal, para alguém olhar.
    assert (pendencia.classificacao.status, pendencia.interprete, pendencia.provas) == (c.HOMONIMA, "Cantor Famoso", [])
    assert pendencia.link.endswith("v3") and "não foi aberto" in pendencia.classificacao.fundamento
    assert (ok.classificacao.status, ok.fornecedor, ok.link) == (c.OK, "Selo Inventado", "https://music.youtube.com/watch?v=v1")
    assert (sem.classificacao.status, sem.interprete, sem.provas) == (c.SEM_CREDITOS, "Banda do Baile", ["prova-v2"])
    assert 'canal "Canal de Fã"' in sem.classificacao.fundamento and "título do vídeo" in sem.vinculo


def test_youtube_erro_de_leitura_e_pagina_trocada_nunca_viram_sem_creditos(rel):
    ytm = YouTubeFalso(
        buscas={"LIGUE O RADIO Banda do Baile": [("v1", "Ligue o Rádio", "Banda do Baile", "", "Música"), ("v2", "Ligue o Rádio", "Banda do Baile", "", "Música")]},
        leituras={"v1": ytmusic.Leitura(video="v1", coleta="erro", erro="o menu de ações da faixa não abriu"),
                  "v2": _leitura("v2", "Uma Música Completamente Diferente", "Banda do Baile", MENU_SEM)},
    )
    assert [g.classificacao.status for g in _youtube(rel, ytm).gravacoes] == [c.ERRO_TECNICO, c.ERRO_TECNICO]


def test_youtube_bloqueio_para_a_fila_guarda_o_que_ja_leu_e_avisa(rel):
    ytm = YouTubeFalso(
        buscas={"LIGUE O RADIO Banda do Baile": [("v1", "Ligue o Rádio", "Banda do Baile", "", "Música")],
                "COISA FEITA Banda do Baile": [("v2", "Coisa Feita", "Banda do Baile", "", "Música")]},
        leituras={"v1": _leitura("v1", "Ligue o Rádio", "Banda do Baile", MENU_SEM)}, bloquear_em="v2",
    )
    coleta = _youtube(rel, ytm)
    assert len(coleta.gravacoes) == 1 and coleta.interrompida.startswith("Bloqueio") and "PAROU" in coleta.avisos[-1]


# --- saídas para a petição ---------------------------------------------------


def _coleta_para_peticao(tmp_path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    from PIL import Image as Pil

    def gravacao(n, obra, status, revisar=False, com_prova=True):
        base = f"prova-{n}"
        if com_prova:
            destino = io.BytesIO()
            Pil.new("RGB", (400, 220), (20 * n, 30, 60)).save(destino, "PNG")
            (tmp_path / f"{base}.png").write_bytes(destino.getvalue())
            (tmp_path / f"{base}.json").write_text(_json.dumps({
                "captura": base, "capturado_em": "2026-03-02T15:56:07-03:00", "fuso": "America/Sao_Paulo", "data_do_servidor": "",
                "arquivos": {"png": f"{base}.png", "html": f"{base}.html"}, "sha256": {"png": captura.sha256(destino.getvalue()), "html": "x"},
            }))
        return pipeline.Gravacao(plataforma="YOUTUBE", link=f"https://music.youtube.com/watch?v=v{n}", titulo=obra, interprete="Banda do Baile",
                                 obra=obra, provas=[base], classificacao=c.Classificacao(status, "motivo", revisar=revisar))
    return pipeline.Coleta("YOUTUBE", gravacoes=[
        gravacao(1, "LIGUE O RADIO", c.SEM_CREDITOS), gravacao(2, "COISA FEITA", c.VIOLACAO), gravacao(3, "BAILE NA SERRA", c.OK),
        gravacao(4, "A DANCA DA PANELA", c.SEM_CREDITOS, revisar=True), gravacao(5, "DANCA DA PANELLA", c.SEM_CREDITOS, com_prova=False),
    ])


import io


def test_lista_da_peticao_so_conta_o_que_esta_firme(rel, tmp_path):
    coleta = _coleta_para_peticao(tmp_path)
    firmes, a_confirmar = provas.itens_da_peticao([coleta], "YOUTUBE")
    assert [provas.linha_da_peticao(i) for i in firmes] == [
        "1 - Música: Coisa Feita https://music.youtube.com/watch?v=v2 Interpretada por: Banda do Baile",
        "2 - Música: Danca da Panella https://music.youtube.com/watch?v=v5 Interpretada por: Banda do Baile",
        "3 - Música: Ligue o Radio https://music.youtube.com/watch?v=v1 Interpretada por: Banda do Baile",
    ]
    assert [i["obra"] for i in a_confirmar] == ["A Danca da Panela"]
    from docx import Document
    texto = "\n".join(p.text for p in Document(io.BytesIO(provas.docx_da_lista(rel, [coleta], "YOUTUBE", "YouTube Music"))).paragraphs)
    assert "Há 3 obras intelectuais" in texto and "A confirmar antes de usar" in texto and "1 - Música: Coisa Feita" in texto


def test_pdf_de_provas_tem_indice_e_um_print_por_pagina_e_recusa_imagem_alterada(rel, tmp_path):
    import pdfplumber

    coleta = _coleta_para_peticao(tmp_path)
    pdf, avisos = provas.pdf_de_provas(rel, [coleta], "YOUTUBE", tmp_path, "YouTube Music")
    with pdfplumber.open(io.BytesIO(pdf)) as lido:
        paginas = [p.extract_text() for p in lido.pages]
    assert len(paginas) == 3 and "Coleta de provas - YouTube Music" in paginas[0] and "SEM CAPTURA" in paginas[0]
    assert "1 - Música: Coisa Feita" in paginas[1] and "3 - Música: Ligue o Radio" in paginas[2]
    assert avisos == ["2 (Danca da Panella / Banda do Baile): sem captura de tela"]
    (tmp_path / "prova-1.png").write_bytes(b"imagem trocada depois")
    _, avisos = provas.pdf_de_provas(rel, [coleta], "YOUTUBE", tmp_path)
    assert any("hash não confere" in a for a in avisos)


def test_url_codificada_nao_e_confundida_com_percentual(rel):
    from urllib.parse import quote, unquote

    filtro = FiltroDeSaida(rel.termos_sensiveis())
    url = "https://music.youtube.com/search?q=" + quote("A DANÇA DA PANELA 3 Zeca Lima")
    assert filtro.motivo(url) == "percentual"  # é por isso que o coletor confere o texto decodificado
    assert filtro.motivo(unquote(url)) == ""


# --- Apple Music, Spotify, Tidal e Vagalume: interpretação sem rede ------------

from creditos import apple, navegador, spotify, tidal, vagalume


def _pagina_da_apple(album_id, faixas, declaradas=None, endereco=None):
    secoes = [
        {"itemKind": "containerDetailHeaderLockup", "items": [{"title": "Álbum Tal", "subtitleLinks": [{"title": "Banda do Baile"}],
                                                               "trackCount": len(faixas) if declaradas is None else declaradas,
                                                               "quaternaryTitle": "Sertanejo · 2012"}]},
        {"itemKind": "trackLockup", "items": [
            {"title": t, "artistName": "Banda do Baile", "composer": comp, "trackNumber": i,
             "contentDescriptor": {"url": f"https://music.apple.com/br/album/x/{album_id}?i={i}"}}
            for i, (t, comp) in enumerate(faixas, start=1)]},
        {"itemKind": "containerDetailTracklistFooterLockup", "items": [{"description": "1 de maio de 2012 2 músicas ℗ 2012 Selo Inventado"}]},
    ]
    estado = _json.dumps({"data": [{"data": {"sections": secoes}}]})
    endereco = endereco or f"https://music.apple.com/br/album/album-tal/{album_id}"
    return f'<meta property="og:url" content="{endereco}"><script type="application/json" id="serialized-server-data">{estado}</script>'


def test_apple_le_o_compositor_de_cada_faixa_e_confere_album_e_contagem():
    album = apple.interpretar_pagina(_pagina_da_apple(55, [("Ligue o Rádio", "Zeca Lima & Cida Dias"), ("Coisa Feita", None)]), "55")
    assert (album.titulo, album.artista, album.lancamento, album.fornecedor, album.aviso) == ("Álbum Tal", "Banda do Baile", "2012", "℗ 2012 Selo Inventado", "")
    assert [(f.titulo, f.creditos, f.link) for f in album.faixas] == [
        ("Ligue o Rádio", ["Zeca Lima", "Cida Dias"], "https://music.apple.com/br/album/x/55?i=1"),
        ("Coisa Feita", [], "https://music.apple.com/br/album/x/55?i=2"),
    ]
    # A Apple redireciona para um endereço com o nome do álbum: a guarda é o final do og:url, não o endereço pedido.
    with pytest.raises(apple.ErroDeColeta, match="não é a do álbum 56"):
        apple.interpretar_pagina(_pagina_da_apple(55, [("X", "Y")]), "56")
    with pytest.raises(apple.ErroDeColeta):
        apple.interpretar_pagina('<meta property="og:url" content="https://music.apple.com/br/album/a/55">sem estado', "55")
    assert apple.interpretar_pagina(_pagina_da_apple(55, [("X", "Y")], declaradas=3), "55").aviso.startswith("CONTAGEM_DIVERGENTE")


def test_spotify_separa_secao_nome_e_papel_e_sem_secao_de_composicao_nao_ha_autor():
    partes = [["texto", "Artista"], ["nome", "Banda do Baile"], ["texto", "Artista principal"], ["texto", "Composição e letra"],
              ["nome", "Zeca Lima"], ["texto", "Autores"], ["nome", "Cida Dias"], ["texto", "Autores"], ["texto", "Fontes"], ["texto", "2000 Selo Inventado"]]
    secoes = spotify.interpretar_janela(partes)
    assert secoes == {"Artista": ["Banda do Baile"], "Composição e letra": ["Zeca Lima", "Cida Dias"], "Fontes": ["2000 Selo Inventado"]}
    assert spotify.autores_e_fornecedor(secoes) == (["Zeca Lima", "Cida Dias"], "2000 Selo Inventado")
    sem = spotify.interpretar_janela(partes[:3] + partes[8:])  # a janela abre, mas não tem a seção de composição
    assert spotify.autores_e_fornecedor(sem) == ([], "2000 Selo Inventado")
    linhas = [{"titulo": "Ligue o Rádio", "faixa": "abc123", "artistas": ["Banda do Baile", "Convidado"], "album": "Álbum Tal"}, {"titulo": "", "faixa": ""}]
    (cand,) = spotify.interpretar_linhas(linhas)
    assert (cand.link, cand.interprete) == ("https://open.spotify.com/intl-pt/track/abc123", "Banda do Baile, Convidado")


def test_tidal_autores_sao_composer_e_lyricist_e_fornecedor_e_producer_e_publisher():
    dados = {"titulo": "Álbum Tal", "artista": "Banda do Baile", "contagem": "2 MÚSICAS", "data": "19 de janeiro de 2024", "faixas": [
        {"id": "901", "numero": "1", "titulo": "Ligue o Rádio", "interprete": "Banda do Baile",
         "celulas": [["Producer", ["Selo Inventado"]], ["Composer", ["Zeca Lima", "Cida Dias"]], ["Lyricist", ["Zeca Lima"]], ["Music Publisher", ["Editora X"]]]},
        {"id": "902", "numero": "2", "titulo": "Coisa Feita", "interprete": "", "celulas": []},
    ]}
    album = tidal.interpretar_album(dados, "77")
    assert (album.lancamento, album.faixas_declaradas, album.fornecedor, album.aviso) == ("2024", 2, "Selo Inventado, Editora X", "")
    um, dois = album.faixas
    assert (um.creditos, um.link) == (["Zeca Lima", "Cida Dias"], "https://tidal.com/album/77/track/901")
    assert (dois.creditos, dois.interprete) == ([], "Banda do Baile")
    assert tidal.interpretar_album({**dados, "contagem": "3 MÚSICAS"}, "77").aviso.startswith("CONTAGEM_DIVERGENTE")
    with pytest.raises(ValueError):
        tidal.interpretar_album({"faixas": []}, "77")  # página sem faixa é erro, nunca álbum sem créditos


PAGINA_DO_VAGALUME = (
    "<h1>Banda do Baile</h1><div>letra</div><small class=styleDesc id=author><span class=tit-CA><b>Compositores:</b> "
    "José Carlos Souza Lima (Zeca Lima) (<a href=https://x target=_blank>ABC</a>), Maria Aparecida Dias (<a href=https://y>DEF</a>)</span>"
    "<span class=tit-E><b>Editores:</b> Editora X (<a href=https://x>ABC</a>)</span></small><p>rodapé Para autores</p>"
)


def test_vagalume_le_so_o_bloco_de_autoria_e_registra_a_negativa_expressa():
    letra = vagalume.interpretar_pagina(PAGINA_DO_VAGALUME, "https://www.vagalume.com.br/banda/ligue-o-radio.html")
    assert letra.creditos == ["José Carlos Souza Lima", "Maria Aparecida Dias"] and not letra.desconhecido
    assert "Compositores: José Carlos Souza Lima (Zeca Lima)" in letra.bloco and "Editores" in letra.bloco
    negada = vagalume.interpretar_pagina("<small id=author><span class=tit-CA><b>Compositor:</b> Desconhecido no ECAD</span></small>", "x")
    assert (negada.creditos, negada.desconhecido) == ([], True)
    with pytest.raises(vagalume.ErroDeColeta):  # sem o bloco, nada se conclui (o rodapé "Para autores" não conta)
        vagalume.interpretar_pagina("<p>rodapé Para autores</p>", "x")
    lista = vagalume.interpretar_lista('<a href=/banda-do-baile/ligue-o-radio.html>Ligue o Rádio</a><a href="/outra/x.html">X</a>', "banda-do-baile")
    assert lista == {"/banda-do-baile/ligue-o-radio.html": "Ligue o Rádio"}


def test_bloqueio_generico_login_captcha_e_desvio():
    for url in ["https://accounts.spotify.com/pt-BR/login", "https://account.deezer.com/login", "https://login.tidal.com/x"]:
        assert navegador.sinal_de_bloqueio(url), url
    assert navegador.sinal_de_bloqueio("https://tidal.com/album/1", "Verify you are human to continue")
    assert navegador.sinal_de_bloqueio("https://open.spotify.com/intl-pt/track/x", "Compartilhar · Ver créditos") == ""


class _VagalumeFalso:
    filtro, requisicoes = None, 0

    def letras_de(self, nome):
        return {"/banda-do-baile/coisa-feita.html": "Coisa Feita", "/banda-do-baile/outra.html": "Outra"} if nome == "Banda do Baile" else {}

    def letra(self, endereco):
        return vagalume.Letra(link="https://www.vagalume.com.br" + endereco, creditos=[], desconhecido=True, bloco="Compositor: Desconhecido no ECAD")


def test_vagalume_negativa_expressa_entra_como_sem_creditos_com_a_nota(rel):
    coleta = pipeline.coletar_vagalume(rel, Config(nomes_confirmados=["ZECA LIMA"]), _VagalumeFalso(), recorrentes=["Banda do Baile"])
    (g,) = coleta.gravacoes
    assert (g.plataforma, g.obra, g.classificacao.status) == ("VAGALUME", "COISA FEITA", c.SEM_CREDITOS)
    assert "AFIRMA que o compositor é desconhecido" in g.classificacao.fundamento


def test_pacote_de_provas_tem_uma_pasta_por_plataforma_com_gravacao_negativa(rel, tmp_path):
    import zipfile

    youtube = _coleta_para_peticao(tmp_path / "youtube")
    sem_nada = pipeline.Coleta("TIDAL")
    pacote, resumo, avisos = provas.pacote(rel, [youtube, sem_nada], tmp_path, b"planilha")
    assert zipfile.ZipFile(io.BytesIO(pacote)).namelist() == [
        "Planilha de Obras.xlsx", "YouTube Music/Lista da petição - YouTube Music.docx", "YouTube Music/Provas - YouTube Music.pdf"]
    assert resumo == [{"plataforma": "YouTube Music", "firmes": 3, "a_confirmar": 1}, {"plataforma": "Tidal", "firmes": 0, "a_confirmar": 0}]
    assert avisos == ["YouTube Music: 2 (Danca da Panella / Banda do Baile): sem captura de tela"]


def test_apple_o_print_e_da_pagina_da_musica_e_a_tela_manda_quando_diverge_dos_dados(rel, tmp_path, monkeypatch):
    assert pipeline.pagina_da_musica_na_apple("https://music.apple.com/br/album/x/55?i=777") == "https://music.apple.com/br/song/777"
    assert pipeline.autores_na_pagina_da_apple(["Composição E Letra", "Zeca Lima", "Composição", "Cida Dias", "Letra"]) == ["Zeca Lima", "Cida Dias"]

    class NavegadorFalso:
        class pagina:
            url = "https://music.apple.com/br/song/777"
            wait_for_timeout = staticmethod(lambda ms: None)
            evaluate = staticmethod(lambda *a: "")
        abertos = []
        def ir(self, url, pausa=None): self.abertos.append(url)
        def conferir(self, texto=""): pass

    capturas = []
    monkeypatch.setattr(captura, "capturar", lambda pagina, pasta, obra, interprete, plataforma, exibido, etapa="": capturas.append(exibido) or {"captura": f"print-{len(capturas)}"})
    sem = lambda n: pipeline.Gravacao(plataforma="APPLE MUSIC", link=f"https://music.apple.com/br/album/x/55?i={n}", titulo="Ligue o Rádio",
                                      interprete="Zeca Lima", obra="LIGUE O RADIO", vinculo="x", classificacao=c.Classificacao(c.SEM_CREDITOS, "sem compositor nos dados do álbum"))
    coleta = pipeline.Coleta("APPLE MUSIC", gravacoes=[sem(1), sem(2), sem(3)])
    telas = iter([[], ["Zeca Lima"], ["Banda do Baile"]])  # igual aos dados; a tela mostra o titular; a tela mostra outro nome
    reclassificar = lambda g: classificar(Item(titulo=g.titulo, interprete=g.interprete, creditos=g.creditos, vinculo=True), rel)
    nav = NavegadorFalso()
    pipeline.capturar_paginas(coleta, nav, tmp_path, lambda n: next(telas), None, pipeline.pagina_da_musica_na_apple, reclassificar)
    igual, creditada, outro = coleta.gravacoes
    assert nav.abertos == ["https://music.apple.com/br/song/1", "https://music.apple.com/br/song/2", "https://music.apple.com/br/song/3"]
    assert (igual.classificacao.status, igual.provas) == (c.SEM_CREDITOS, ["print-1"])
    assert (creditada.classificacao.status, creditada.provas) == (c.OK, []) and "vale o que a tela mostra" in creditada.classificacao.fundamento
    assert (outro.classificacao.status, outro.creditos, outro.provas) == (c.VIOLACAO, ["Banda do Baile"], ["print-2"])


# --- ligação provada x ligação presumida (os dois erros do primeiro caso real) ---


def test_nome_de_uma_palavra_nunca_casa_com_nome_maior():
    assert c.mesmo_artista("Banda do Baile", "Grupo Banda do Baile") and c.mesmo_artista("BANDA DO BAILE", "banda do baile")
    assert not c.mesmo_artista("Fulano", "Beltrano Fulano") and not c.mesmo_artista("Fulano", "Oficial Beltrano Fulano")
    assert c.mesmo_artista("Fulano", "FULANO")
    assert c.partes_do_interprete("Fulano, Beltrano & Sicrano feat. Outro") == ["Fulano", "Beltrano", "Sicrano", "Outro"]
    assert pipeline._citado("Banda do Baile", "Banda do Baile - Ligue o Rádio")
    assert pipeline._citado("Fulano", "Ligue o Rádio | Fulano (Ao Vivo)")
    assert not pipeline._citado("Fulano", "Ligue o Rádio | Beltrano Fulano - coreografia")


def test_artista_que_so_grava_outras_obras_do_titular_nao_sustenta_acusacao(rel):
    presumindo = Config(nomes_confirmados=["ZECA LIMA"], presumidos=["Cantor Popular"])
    # Mesmo título, cantado por quem grava outras obras do titular, com compositores que não são os do relatório:
    # é a marca de uma música homônima, e não de um crédito errado.
    outra = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Popular, Convidado",
                             creditos=["Cantor Popular", "Convidado", "Outro Autor Qualquer"]), rel, presumindo)
    assert (outra.status, outra.revisar) == (c.HOMONIMA, True)
    # Sem crédito nenhum: pode ser a obra, mas só com a presunção não entra em contagem nenhuma.
    sem = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Popular"), rel, presumindo)
    assert (sem.status, sem.revisar, sem.presumido) == (c.SEM_CREDITOS, True, True) and "PRESUMIDO" in sem.fundamento
    # Quando alguém confirma o intérprete, ou o par já foi provado em outra plataforma, o resultado é firme.
    confirmado = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Popular"), rel, Config(interpretes=["Cantor Popular"]))
    provado = classificar(Item(titulo="Ligue o Rádio", interprete="Cantor Popular", vinculo=True), rel, presumindo)
    assert [(k.status, k.revisar, k.presumido) for k in (confirmado, provado)] == [(c.SEM_CREDITOS, False, False)] * 2
    # E um autor do relatório no crédito prova a ligação por si só.
    creditado = classificar(Item(titulo="Coisa Feita", interprete="Cantor Popular", creditos=["Cida Dias"]), rel, presumindo)
    assert (creditado.status, creditado.revisar) == (c.VIOLACAO, False)


def test_youtube_nao_liga_pelo_sobrenome_e_presuncao_fica_a_confirmar(rel):
    ytm = YouTubeFalso(
        buscas={"LIGUE O RADIO Fulano": [
            ("v1", "Ligue o Rádio", "Fulano", "Álbum", "Música"),
            ("v2", "Beltrano Fulano - Ligue o Rádio", "Oficial Beltrano Fulano", "2 mil visualizações", "Vídeo"),  # outro artista
        ], "COISA FEITA Fulano": [("v3", "Coisa Feita", "Fulano", "Álbum", "Música")]},
        leituras={"v1": _leitura("v1", "Ligue o Rádio", "Fulano", MENU_SEM), "v3": _leitura("v3", "Coisa Feita", "Fulano", MENU_SEM)},
    )
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Fulano"],
                                      pares={"COISA FEITA": ["Fulano"]})
    assert ytm.lidos == ["v1", "v3"] and coleta.sem_vinculo_nao_abertos == 1  # o vídeo do homônimo de sobrenome nem é aberto
    pendencia, presumida, provada = coleta.gravacoes
    assert (pendencia.classificacao.status, pendencia.interprete) == (c.HOMONIMA, "Oficial Beltrano Fulano")
    assert (presumida.classificacao.status, presumida.classificacao.presumido) == (c.SEM_CREDITOS, True)
    assert (provada.classificacao.status, provada.classificacao.revisar) == (c.SEM_CREDITOS, False)  # par provado em outra plataforma
    firmes, a_confirmar = provas.itens_da_peticao([coleta], "YOUTUBE")
    assert [i["obra"] for i in firmes] == ["Coisa Feita"] and [i["obra"] for i in a_confirmar] == ["Ligue o Radio"]


def test_print_ja_tirado_e_reaproveitado_na_rodada_seguinte(tmp_path):
    (tmp_path / "p1.png").write_bytes(b"png")
    ficha = {"captura": "p1", "url": "https://music.apple.com/br/song/nome-da-musica/777", "arquivos": {"png": "p1.png"},  # com o redirecionamento
             "exibido": {"creditos_lidos": ["Banda do Baile"]}}
    (tmp_path / "p1.json").write_text(_json.dumps(ficha))
    (tmp_path / "solta.json").write_text(_json.dumps({"captura": "p2", "url": "x", "arquivos": {"png": "sumiu.png"}}))
    assert [f["captura"] for f in captura.capturas_guardadas(tmp_path)] == ["p1"]  # ficha sem a imagem não vale
    g = lambda n, creditos: pipeline.Gravacao(plataforma="APPLE MUSIC", link=f"https://music.apple.com/br/album/x/55?i={n}", creditos=creditos,
                                              classificacao=c.Classificacao(c.VIOLACAO, "motivo"))
    coleta = pipeline.Coleta("APPLE MUSIC", gravacoes=[g(777, ["Banda do Baile"]), g(777, ["Outro Nome"]), g(778, ["Banda do Baile"])])
    pipeline.reaproveitar_prints(coleta, tmp_path, pipeline.pagina_da_musica_na_apple)
    # Só a gravação do mesmo endereço e com os mesmos créditos que estavam na tela quando o print foi tirado.
    assert [x.provas for x in coleta.gravacoes] == [["p1"], [], []]


# --- estimativa de tempo e atualização do app ---------------------------------

from creditos import atualizacao


def test_estimativa_de_tempo_cresce_com_obras_e_plataformas_e_nao_finge_precisao():
    todas = pipeline.estimar_minutos(119, pipeline.TODAS)
    assert todas > pipeline.estimar_minutos(119, ("deezer", "apple")) > pipeline.estimar_minutos(10, ("deezer", "apple"))
    assert pipeline.estimar_minutos(10, ()) == pipeline.estimar_minutos(10, ("deezer",))  # a Deezer roda sempre
    assert pipeline.texto_da_estimativa(3) == "cerca de 3 minutos" and pipeline.texto_da_estimativa(1) == "cerca de 1 minuto"
    assert pipeline.texto_da_estimativa(22) == "cerca de 20 minutos"
    assert pipeline.texto_da_estimativa(171) == "cerca de 2h50" and pipeline.texto_da_estimativa(118) == "cerca de 2 horas"


def _pacote(arquivos: dict, raiz="quem-canta-9.9.9/") -> bytes:
    import zipfile
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        for nome, conteudo in arquivos.items():
            z.writestr(raiz + nome, conteudo)
    return saida.getvalue()


def test_so_avisa_quando_a_versao_publicada_e_maior(monkeypatch):
    assert atualizacao.mais_nova("0.10.0", "0.9.5") and atualizacao.mais_nova("v1.0", "0.9.9")
    assert not atualizacao.mais_nova("0.4.0", "0.4.0") and not atualizacao.mais_nova("0.3.9", "0.4.0") and not atualizacao.mais_nova("", "0.4.0")

    class Resposta:
        def __init__(self, dados): self.dados = dados
        def json(self): return self.dados
    publicada = {"versao": "99.0.0", "url": "https://exemplo/app.zip", "sha256": "A" * 64, "notas": "melhorias"}
    monkeypatch.setattr(atualizacao.requests, "get", lambda url, timeout: Resposta(publicada))
    assert atualizacao.consultar("https://exemplo/versao.json") == {**publicada, "sha256": "a" * 64}
    monkeypatch.setattr(atualizacao.requests, "get", lambda url, timeout: Resposta({"versao": "0.0.1"}))
    assert atualizacao.consultar("https://exemplo/versao.json") is None
    assert atualizacao.consultar("http://sem-https/versao.json") is None

    def fora_do_ar(url, timeout):
        raise atualizacao.requests.ConnectionError()
    monkeypatch.setattr(atualizacao.requests, "get", fora_do_ar)
    assert atualizacao.consultar("https://exemplo/versao.json") is None  # sem rede, o app abre do mesmo jeito


def test_atualizacao_troca_o_programa_e_nunca_toca_nos_dados_nem_no_ambiente(tmp_path):
    (tmp_path / "dados").mkdir()
    (tmp_path / "dados" / "caso.txt").write_text("prints do cliente")
    (tmp_path / "creditos").mkdir()
    (tmp_path / "creditos" / "versao.py").write_text("antiga")
    (tmp_path / "atualizacao.json").write_text('{"consultar_em": "https://do-escritorio"}')
    zip_novo = _pacote({"creditos/versao.py": "nova", "creditos_app.py": "tela nova", "dados/caso.txt": "APAGARIA",
                        ".venv/x.txt": "não entra", "atualizacao.json": "{}"})
    publicada = {"versao": "9.9.9", "url": "https://exemplo/app.zip", "sha256": captura.sha256(zip_novo)}
    escritos = atualizacao.instalar(publicada, tmp_path, baixar=lambda url: zip_novo)
    assert sorted(escritos) == ["creditos/versao.py", "creditos_app.py"]
    assert (tmp_path / "creditos" / "versao.py").read_text() == "nova"
    assert (tmp_path / "dados" / "caso.txt").read_text() == "prints do cliente" and not (tmp_path / ".venv").exists()
    assert "do-escritorio" in (tmp_path / "atualizacao.json").read_text()
    assert atualizacao.endereco_de_consulta(tmp_path) == "https://do-escritorio"


def test_endereco_de_consulta_tem_padrao_e_pode_ser_trocado_ou_desligado(tmp_path, monkeypatch):
    monkeypatch.delenv("QUEMCANTA_ATUALIZACAO", raising=False)
    assert atualizacao.endereco_de_consulta(tmp_path) == atualizacao.ENDERECO_PADRAO  # sem arquivo
    (tmp_path / "atualizacao.json").write_text('{"consultar_em": ""}')
    assert atualizacao.endereco_de_consulta(tmp_path) == atualizacao.ENDERECO_PADRAO  # arquivo vazio, como no pacote
    (tmp_path / "atualizacao.json").write_text("isto não é json")
    assert atualizacao.endereco_de_consulta(tmp_path) == atualizacao.ENDERECO_PADRAO
    (tmp_path / "atualizacao.json").write_text('{"consultar_em": "https://outro/versao.json"}')
    assert atualizacao.endereco_de_consulta(tmp_path) == "https://outro/versao.json"
    monkeypatch.setenv("QUEMCANTA_ATUALIZACAO", "desligada")
    assert atualizacao.endereco_de_consulta(tmp_path) == "desligada"
    monkeypatch.setattr(atualizacao.requests, "get", lambda url, timeout: pytest.fail("não podia consultar"))
    assert atualizacao.verificar(atualizacao.endereco_de_consulta(tmp_path)) == (atualizacao.SEM_RESPOSTA, None)
    assert atualizacao.ENDERECO_PADRAO.startswith("https://raw.githubusercontent.com/")
    assert atualizacao.endereco_do_pacote("1.2.3").startswith("https://github.com/") and "/v1.2.3/" in atualizacao.endereco_do_pacote("1.2.3")


def test_verificar_separa_versao_nova_em_dia_e_sem_resposta(monkeypatch):
    class Resposta:
        def __init__(self, dados): self.dados = dados
        def json(self):
            if isinstance(self.dados, Exception):
                raise self.dados
            return self.dados
    def com(dados):
        monkeypatch.setattr(atualizacao.requests, "get", lambda url, timeout: Resposta(dados))
        return atualizacao.verificar("https://exemplo/versao.json")
    estado, publicada = com({"versao": "99.0.0", "url": "https://exemplo/app.zip", "sha256": "A" * 64})
    assert estado == atualizacao.NOVA and publicada["versao"] == "99.0.0" and publicada["sha256"] == "a" * 64
    assert com({"versao": atualizacao.VERSAO}) == (atualizacao.EM_DIA, None)
    assert com({"versao": "0.0.1"}) == (atualizacao.EM_DIA, None)
    # Resposta que não é a esperada (página de erro, lista, sem versão) nunca vira "em dia".
    for ruim in ({}, [], {"versao": "abc"}, "404: Not Found", ValueError("não é json")):
        assert com(ruim) == (atualizacao.SEM_RESPOSTA, None)


def test_atualizacao_nao_regrava_arquivo_igual(tmp_path):
    (tmp_path / "requirements.txt").write_text("streamlit")
    (tmp_path / "creditos_app.py").write_text("tela velha")
    pacote = _pacote({"requirements.txt": "streamlit", "creditos_app.py": "tela nova"})
    publicada = {"url": "https://exemplo/app.zip", "sha256": captura.sha256(pacote)}
    assert atualizacao.instalar(publicada, tmp_path, baixar=lambda url: pacote) == ["creditos_app.py"]
    assert atualizacao.instalar(publicada, tmp_path, baixar=lambda url: pacote) == []  # instalar de novo não muda nada


def test_pacote_do_windows_nao_leva_dados_nem_o_arquivo_de_versao_publicada(tmp_path, monkeypatch):
    import zipfile
    import empacotar
    (tmp_path / "app").mkdir()
    for nome in ("creditos_app.py", "requirements.txt", "Instalar.cmd", "atualizacao.json"):
        (tmp_path / "app" / nome).write_text("x\n")
    (tmp_path / "app" / "dados").mkdir()
    (tmp_path / "app" / "dados" / "caso.pdf").write_text("relatório de cliente")
    (tmp_path / "app" / "versao-publicada.json").write_text("{}")
    monkeypatch.setattr(empacotar, "AQUI", tmp_path / "app")
    monkeypatch.setattr(empacotar.sys, "argv", ["empacotar.py", str(tmp_path), "o", "que", "mudou"])
    empacotar.main()
    pacote = tmp_path / f"quem-canta-{atualizacao.VERSAO}-windows.zip"
    nomes = [n.split("/", 1)[1] for n in zipfile.ZipFile(pacote).namelist()]
    assert sorted(nomes) == ["Instalar.cmd", "atualizacao.json", "creditos_app.py", "requirements.txt"]
    publicada = _json.loads((tmp_path / "app" / "versao-publicada.json").read_text())
    assert publicada == {"versao": atualizacao.VERSAO, "url": atualizacao.endereco_do_pacote(), "notas": "o que mudou",
                         "sha256": captura.sha256(pacote.read_bytes())}
    # O que o empacotar publica é exatamente o que o app instalado aceita instalar.
    (tmp_path / "instalado").mkdir()
    assert sorted(atualizacao.instalar(publicada, tmp_path / "instalado", baixar=lambda url: pacote.read_bytes())) == [
        "Instalar.cmd", "creditos_app.py", "requirements.txt"]


def test_atualizacao_recusa_pacote_sem_hash_adulterado_ou_que_sai_da_pasta(tmp_path):
    zip_bom = _pacote({"creditos_app.py": "tela"})
    for publicada, motivo in [
        ({"url": "https://x/app.zip", "sha256": ""}, "hash"),
        ({"url": "http://x/app.zip", "sha256": "a" * 64}, "https"),
        ({"url": "https://x/app.zip", "sha256": "b" * 64}, "não confere"),
    ]:
        with pytest.raises(atualizacao.FalhaNaAtualizacao, match=motivo):
            atualizacao.instalar(publicada, tmp_path, baixar=lambda url: zip_bom)
    malicioso = _pacote({"../../fora.txt": "x"})
    with pytest.raises(atualizacao.FalhaNaAtualizacao, match="fora da pasta"):
        atualizacao.instalar({"url": "https://x/app.zip", "sha256": captura.sha256(malicioso)}, tmp_path, baixar=lambda url: malicioso)
    assert not list(tmp_path.glob("**/*.py")) and not (tmp_path.parent / "fora.txt").exists()


def test_limite_negativo_pega_as_ultimas_obras(rel):
    coleta = pipeline.Coleta("X")
    assert [t for _, t in pipeline._titulos(rel, 2, coleta, "X")] == ["LIGUE O RADIO", "COISA FEITA"]
    assert [t for _, t in pipeline._titulos(rel, -2, coleta, "X")] == ["A DANCA DA PANELA", "DANCA DA PANELLA"]
    assert len(pipeline._titulos(rel, None, coleta, "X")) == 5 and "últimas obras de 5" in coleta.avisos[-1]


# --- andamento da coleta: barra de progresso e parada ---------------------------

from creditos import andamento as _andamento


def test_andamento_avanca_por_etapa_e_nunca_volta():
    relogio = [0.0]
    a = _andamento.Andamento(10, ("deezer", "spotify"), agora=lambda: relogio[0])
    assert a.etapas == ["deezer", "spotify", "prints"] and a.retrato()["fracao"] == 0
    vistas = []
    for texto in ["Deezer: discografia de ZECA LIMA", "Deezer: busca por título 1/10", "Deezer: busca por título 10/10",
                  "Deezer: lendo álbum 1/4", "Deezer: lendo álbum 4/4", "Spotify: buscando 1/10", "Spotify: lendo faixa 1/20",
                  "Deezer: lendo álbum 1/4",  # aviso atrasado de etapa que já passou: não faz a barra voltar
                  "Spotify: lendo faixa 20/20", "Deezer: print 1/2", "Deezer: print 2/2"]:
        a(texto)
        vistas.append(a.retrato()["fracao"])
    assert vistas == sorted(vistas) and vistas[0] > 0 and 0.9 < vistas[-1] <= 0.99  # só chega a 100% quando termina
    retrato = a.retrato()
    assert (retrato["etapa"], retrato["etapas"], retrato["nome"]) == (3, 3, "Prints das provas")
    assert retrato["detalhe"] == "Deezer: print 2 de 2"
    assert retrato["situacoes"] == [("Deezer", "concluída"), ("Spotify", "concluída"), ("Prints das provas", "em andamento")]


def test_andamento_estima_o_tempo_que_falta_pelo_ritmo_real():
    relogio = [0.0]
    a = _andamento.Andamento(10, ("deezer",), agora=lambda: relogio[0])
    assert a.texto_do_restante().startswith("falta")  # no começo, vale a estimativa do tamanho do relatório
    a("Deezer: lendo álbum 3/4")
    feito = a.fracao
    relogio[0] = 600.0
    assert a.segundos_restantes() == pytest.approx(600 * (1 - feito) / feito)
    relogio[0] = 16.0
    a.fracao = 0.99
    assert a.texto_do_restante() == "falta menos de 1 minuto"
    a("Abrindo o navegador")  # aviso sem plataforma: só muda o texto
    assert a.retrato()["detalhe"] == "Abrindo o navegador" and a.fracao == 0.99


def test_pedido_de_parada_interrompe_no_proximo_aviso_e_nao_vira_erro():
    def coleta(avisar):
        for i in range(1, 1000):
            avisar(f"Deezer: lendo álbum {i}/1000")
            try:
                time.sleep(0.005)
            except Exception:  # um `except Exception` da coleta não pode engolir a parada
                pass
        return "terminou"

    import time
    tarefa = _andamento.Tarefa(coleta, _andamento.Andamento(5, ("deezer",)))
    assert tarefa.viva
    tarefa.pedir_parada()
    tarefa.linha.join(5)
    assert not tarefa.viva and tarefa.interrompida and tarefa.erro is None and tarefa.resultado is None
    assert not issubclass(_andamento.Interrompida, Exception)

    def quebra(avisar):
        raise RuntimeError("navegador não abriu")
    com_erro = _andamento.Tarefa(quebra, _andamento.Andamento(5, ("deezer",)))
    com_erro.linha.join(5)
    assert isinstance(com_erro.erro, RuntimeError) and not com_erro.interrompida


def test_so_uma_coleta_por_vez(monkeypatch):
    import threading
    monkeypatch.setattr(_andamento, "_ATUAL", None)
    solta = threading.Event()
    primeira = _andamento.iniciar(lambda avisar: solta.wait(5), _andamento.Andamento(1, ()))
    assert _andamento.iniciar(lambda avisar: "outra", _andamento.Andamento(1, ())) is primeira  # não abre uma segunda
    _andamento.encerrar()
    assert _andamento.atual() is primeira  # em curso: não é esquecida
    solta.set()
    primeira.linha.join(5)
    _andamento.encerrar()
    assert _andamento.atual() is None


# --- print de tela inteira ------------------------------------------------------

from creditos import tela as _tela


def _pagina_falsa(largura=800, altura=500):
    """Uma "captura de página": faixa branca de carimbo, a linha vermelha e um conteúdo com desenho próprio."""
    import numpy as np
    pagina = np.zeros((altura, largura, 3), dtype=np.uint8)
    pagina[:40] = 255
    pagina[40:42] = _tela.VERMELHO
    y, x = np.mgrid[42:altura, 0:largura]
    pagina[42:, :, 0], pagina[42:, :, 1], pagina[42:, :, 2] = (x // 7 * 13) % 256, (y // 5 * 29) % 256, ((x + y) // 11 * 17) % 256
    return pagina


def _monitor_com(pagina, x=300, y=180, largura=1920, altura=1080, escala=1):
    """Um monitor cinza com a janela (barra de endereço de 90 px + página) na posição dada."""
    import numpy as np
    from PIL import Image
    monitor = np.full((altura, largura, 3), 60, dtype=np.uint8)
    if escala != 1:
        pagina = np.asarray(Image.fromarray(pagina).resize((round(pagina.shape[1] * escala), round(pagina.shape[0] * escala)), Image.NEAREST))
    alto, largo = min(pagina.shape[0], altura - y), min(pagina.shape[1], largura - x)
    monitor[max(0, y - 90):y, x:x + largo] = 230  # a moldura do navegador
    monitor[y:y + alto, x:x + largo] = pagina[:alto, :largo]
    return monitor


def _png(matriz) -> bytes:
    from PIL import Image
    saida = io.BytesIO()
    Image.fromarray(matriz).save(saida, "PNG")
    return saida.getvalue()


def test_tela_inteira_acha_a_janela_e_recusa_coberta_cortada_ou_ausente():
    pagina = _pagina_falsa()
    assert _tela.achar_faixa(pagina) == (40, 0, 800)
    assert _tela.conferir(_monitor_com(pagina), pagina) == (True, "")
    assert _tela.conferir(_monitor_com(pagina, escala=2, largura=3840, altura=2160), pagina) == (True, "")  # monitor de alta densidade
    # Monitor sem a janela (a pessoa está trabalhando nele, ou a janela foi minimizada).
    vazio = _monitor_com(pagina)[:, :250]
    assert _tela.conferir(vazio, pagina) == (False, "a janela do navegador não está à vista neste monitor")
    # Outra janela por cima de boa parte da página.
    coberta = _monitor_com(pagina)
    coberta[260:600, 300:1100] = 255
    assert _tela.conferir(coberta, pagina) == (False, "a janela do navegador está coberta por outra janela")
    # Janela arrastada para fora da borda: metade da página não aparece.
    cortada = _monitor_com(pagina, x=1500)
    assert _tela.conferir(cortada, pagina) == (False, "a janela do navegador não cabe inteira no monitor")
    # Um detalhe pequeno que se mexe (a barra de reprodução, o cursor) não reprova a foto.
    com_cursor = _monitor_com(pagina)
    com_cursor[400:412, 500:512] = 255
    assert _tela.conferir(com_cursor, pagina)[0]


def test_tela_inteira_fotografa_so_o_monitor_onde_a_pagina_esta(monkeypatch):
    pagina = _pagina_falsa()
    de_trabalho, livre = _monitor_com(pagina)[:, :250], _monitor_com(pagina)
    monitores = lambda: iter([(1, de_trabalho, b"png do monitor de trabalho"), (2, livre, b"png do monitor livre")])
    assert _tela.fotografar(_png(pagina), monitores) == (b"png do monitor livre", {"monitor": 2})
    monkeypatch.setattr(_tela.sys, "platform", "win32")
    foto, dados = _tela.fotografar(_png(pagina), lambda: iter([(1, de_trabalho, b"x")]))
    assert foto is None and dados == {"erro": "a janela do navegador não está à vista neste monitor"}


class _PaginaDeCaptura:
    """O mínimo de uma página do navegador para a captura: devolve uma imagem com a faixa de carimbo."""
    url = "https://exemplo/faixa/1"

    def __init__(self):
        self.trazida_para_frente = 0

    def evaluate(self, script, *args):
        return "" if "favicon" in script else False if "translated" in script else None

    def content(self):
        return "<html></html>"

    def screenshot(self, type="png"):
        return _png(_pagina_falsa())

    def wait_for_timeout(self, ms):
        pass

    def bring_to_front(self):
        self.trazida_para_frente += 1


def test_captura_guarda_a_tela_inteira_com_hash_e_mostra_no_pdf(tmp_path, monkeypatch, rel):
    monkeypatch.setattr(captura, "TELA_INTEIRA", True)
    foto = _png(_monitor_com(_pagina_falsa()))
    monkeypatch.setattr(_tela, "fotografar", lambda imagem: (foto, {"monitor": 2}))
    registro = captura.capturar(_PaginaDeCaptura(), tmp_path, "LIGUE O RADIO", "Banda do Baile", "spotify", {}, "creditos")
    assert registro["tela_inteira"] == {"monitor": 2}
    assert (tmp_path / registro["arquivos"]["tela"]).read_bytes() == foto and registro["sha256"]["tela"] == captura.sha256(foto)
    assert registro["arquivos"]["tela"].endswith("_creditos_tela.png")
    # No PDF de provas, vale a foto de tela inteira, com o hash dela; se for alterada, deixa de valer como prova.
    guardado = provas.registro_da_captura(tmp_path, registro["captura"])
    assert guardado["imagem"] == "tela" and guardado["integra"]
    (tmp_path / registro["arquivos"]["tela"]).write_bytes(foto + b"alterada")
    assert not provas.registro_da_captura(tmp_path, registro["captura"])["integra"]


def test_captura_sem_tela_inteira_continua_valendo_e_registra_o_motivo(tmp_path, monkeypatch):
    monkeypatch.setattr(captura, "TELA_INTEIRA", True)
    monkeypatch.setattr(_tela, "fotografar", lambda imagem: (None, {"erro": "a janela do navegador está coberta por outra janela"}))
    pagina = _PaginaDeCaptura()
    registro = captura.capturar(pagina, tmp_path, "LIGUE O RADIO", "Banda do Baile", "spotify", {}, "creditos")
    assert pagina.trazida_para_frente == 1  # tentou uma vez trazer a janela para a frente
    assert registro["tela_inteira"] == {"erro": "a janela do navegador está coberta por outra janela"}
    assert "tela" not in registro["arquivos"] and (tmp_path / registro["arquivos"]["png"]).is_file()
    assert provas.registro_da_captura(tmp_path, registro["captura"])["imagem"] == "png"
    # Desligado (o padrão), nada de tela inteira nem de chave nova na ficha.
    monkeypatch.setattr(captura, "TELA_INTEIRA", False)
    monkeypatch.setattr(_tela, "fotografar", lambda imagem: pytest.fail("não podia fotografar a tela"))
    assert "tela_inteira" not in captura.capturar(_PaginaDeCaptura(), tmp_path / "outra", "LIGUE O RADIO", "Banda do Baile", "spotify", {})


def test_executar_liga_a_tela_inteira_so_durante_a_coleta_e_deixa_o_navegador_visivel(monkeypatch, rel, tmp_path):
    vistos = []
    monkeypatch.setattr(pipeline, "_executar", lambda *a, **k: vistos.append((captura.TELA_INTEIRA, a[-1])) or [])
    pipeline.executar(rel, Config(), tmp_path, mostrar_navegador=False, tela_inteira=True)
    pipeline.executar(rel, Config(), tmp_path, mostrar_navegador=False)
    assert vistos == [(True, True), (False, False)] and captura.TELA_INTEIRA is False


def test_captura_reabre_a_janela_minimizada_antes_de_fotografar(tmp_path):
    comandos = []

    class Sessao:
        def __init__(self, estado): self.estado = estado
        def send(self, comando, parametros=None):
            comandos.append((comando, parametros))
            return {"windowId": 7, "bounds": {"windowState": self.estado}}
        def detach(self): pass

    def pagina_com_janela(estado):
        pagina = _PaginaDeCaptura()
        pagina.context = type("Contexto", (), {"new_cdp_session": lambda self, p: Sessao(estado)})()
        return pagina

    captura.capturar(pagina_com_janela("minimized"), tmp_path, "LIGUE O RADIO", "Banda do Baile", "spotify", {})
    assert ("Browser.setWindowBounds", {"windowId": 7, "bounds": {"windowState": "normal"}}) in comandos
    comandos.clear()
    captura.capturar(pagina_com_janela("normal"), tmp_path / "b", "LIGUE O RADIO", "Banda do Baile", "spotify", {})
    assert [c for c, _ in comandos] == ["Browser.getWindowForTarget"]  # janela aberta: não mexe nela


# --- YouTube Music: vídeos, títulos escritos de outro jeito e medleys ------------------

def test_youtube_busca_videos_de_todo_interprete_conhecido_e_tambem_so_pelo_titulo(rel):
    pedidas = []

    class Anota(YouTubeFalso):
        def buscar(self, consulta, abas):
            pedidas.append((consulta, abas))
            return super().buscar(consulta, abas)

    ytm = Anota(buscas={"COISA FEITA Cantora Original": [("v1", "Cantora Original - Coisa Feita (Clipe Oficial)", "Cantora Original", "9 mil visualizações", "Vídeo")],
                        "COISA FEITA": [("v2", "Coisa Feita", "Outro Cantor", "Álbum", "Música")]},
                leituras={"v1": _leitura("v1", "Cantora Original - Coisa Feita (Clipe Oficial)", "Cantora Original", MENU_SEM)})
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Banda do Baile"],
                                      pares={"COISA FEITA": ["Cantora Original"]}, limite_de_obras=2)
    da_obra = [p for p in pedidas if p[0].startswith("COISA FEITA")]
    assert da_obra == [("COISA FEITA Banda do Baile", ("Músicas", "Vídeos")), ("COISA FEITA ZECA LIMA", ("Músicas", "Vídeos")),
                       ("COISA FEITA Cantora Original", ("Músicas", "Vídeos")), ("COISA FEITA", ("Músicas", "Vídeos"))]
    clipe = next(g for g in coleta.gravacoes if g.link.endswith("v1"))
    assert (clipe.classificacao.status, clipe.classificacao.revisar, clipe.interprete) == (c.SEM_CREDITOS, False, "Cantora Original")
    assert [g.classificacao.status for g in coleta.gravacoes if g.link.endswith("v2")] == [c.HOMONIMA]


def test_youtube_reconhece_o_titulo_com_o_nome_do_artista_grudado():
    rel = relatorio_de_teste()
    nomes = ["Banda do Baile", "ZECA LIMA"]
    for exibido, esperado in [
        ("Banda do Baile- Coisa Feita", "Coisa Feita"),
        ("Grupo banda do baile coisa feita", "coisa feita"),
        ("Banda do Baile - Coisa Feita [Álbum As 10 Mais]", "Coisa Feita"),
        ("Coisa Feita - Banda do Baile Dvd 10 anos 2008", "Coisa Feita"),
        ("Zeca Lima: Coisa Feita", "Coisa Feita"),
        ("Banda do Baile - Coisa Feita ( Vídeo Clipe Oficial ) Part. Outra Banda", "Coisa Feita"),
        ("Banda do Baile - Coisa Feita (Part. Outra Banda)", "Coisa Feita (Part. Outra Banda)"),
        ("Banda do Baile - Coisa Feita ( Part.Esp. Outra Banda ) 2016/2017", "Coisa Feita ( Part.Esp. Outra Banda )"),
    ]:
        assert pipeline._titulo_do_video(exibido, rel, nomes) == esperado, exibido
    assert pipeline._titulo_do_video("Uma Música Completamente Diferente", rel, nomes) == "Uma Música Completamente Diferente"
    # Dupla citada no título do vídeo de um fã: o "e" não pode parti-la ao meio.
    assert pipeline._nome_dentro("Fulana e Beltrano", "Fulana & Beltrano - Coisa Feita") and pipeline._nome_dentro("Banda do Baile", "banda do baile ao vivo")
    assert not pipeline._nome_dentro("Fulano", "Beltrano Fulano - Ligue o Rádio")  # nome de uma palavra só nunca casa assim
    assert not pipeline._nome_dentro("Banda do Baile", "Banda Baile")


def test_youtube_le_titulo_parecido_e_medley_de_interprete_conhecido_mas_deixa_a_revisar(rel):
    ytm = YouTubeFalso(
        buscas={"COISA FEITA Banda do Baile": [
            ("v1", "Coisa Feitta", "Banda do Baile", "Álbum", "Música"),                      # título com letra a mais
            ("v2", "Ligue o Rádio / Coisa Feita (Ao Vivo)", "Banda do Baile", "Álbum", "Música"),  # medley com a obra
            ("v3", "Coisa Feitta", "Canal Qualquer", "100 visualizações", "Vídeo"),              # parecido e sem ligação: fica de fora
        ]},
        leituras={"v1": _leitura("v1", "Coisa Feitta", "Banda do Baile", MENU_SEM),
                  "v2": _leitura("v2", "Ligue o Rádio / Coisa Feita (Ao Vivo)", "Banda do Baile", MENU_SEM)},
    )
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Banda do Baile"])
    por_video = {g.link[-2:]: g for g in coleta.gravacoes}
    assert sorted(por_video) == ["v1", "v2", "v3"] and sorted(ytm.lidos) == ["v1", "v2"]
    # O de canal desconhecido não é aberto: fica nas pendências, dizendo que o título é só parecido.
    assert por_video["v3"].classificacao.status == c.HOMONIMA and "título parecido" in por_video["v3"].classificacao.fundamento
    # Lidos e listados, mas nunca como resultado firme: título só parecido e medley pedem conferência.
    assert (por_video["v1"].classificacao.status, por_video["v1"].classificacao.revisar) == (c.TITULO_APROXIMADO, True)
    assert por_video["v1"].classificacao.sugestao == c.SEM_CREDITOS
    assert (por_video["v2"].classificacao.status, por_video["v2"].classificacao.revisar) == (c.MEDLEY, True)
    assert provas.itens_da_peticao([coleta], "YOUTUBE") == ([], [])


def test_youtube_dupla_com_par_provado_sai_firme_mesmo_em_video_de_fa(rel):
    # "Fulana e Beltrano": a divisão do nome em partes não pode desfazer a dupla na hora de conferir o par provado.
    ytm = YouTubeFalso(
        buscas={"COISA FEITA Fulana & Beltrano": [
            ("v1", "Fulana e Beltrano - Coisa Feita", "Canal de Fã", "418 visualizações", "Vídeo"),
            ("v2", "Coisa Feita - Fulana & Beltrano Dvd 10 anos", "Fulana e Beltrano", "695 visualizações", "Vídeo"),
        ]},
        leituras={"v1": _leitura("v1", "Fulana e Beltrano - Coisa Feita", "Canal de Fã", MENU_SEM),
                  "v2": _leitura("v2", "Coisa Feita - Fulana & Beltrano Dvd 10 anos", "Fulana e Beltrano", MENU_SEM)},
    )
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Banda do Baile"],
                                      pares={"COISA FEITA": ["Fulana & Beltrano"]})
    assert sorted((g.classificacao.status, g.classificacao.revisar, g.interprete) for g in coleta.gravacoes) == [
        (c.SEM_CREDITOS, False, "Fulana & Beltrano"), (c.SEM_CREDITOS, False, "Fulana e Beltrano")]
    assert pipeline._par_conhecido({"X": ["Fulana & Beltrano"]}, "X", "Fulana e Beltrano remixed by Outro")
    assert not pipeline._par_conhecido({"X": ["Fulana & Beltrano"]}, "X", "Fulana")


def test_youtube_nome_da_banda_parecido_com_titulo_de_obra_nao_vira_a_obra():
    rel = relatorio_de_teste()   # tem a obra "COISA FEITA"; a banda conhecida se chama "Coisa Feitta"
    nomes = ["Coisa Feitta"]
    assert c.casar_titulo("Coisa Feitta", rel)[0] == "aproximado"
    # "Banda - Outra Música": o nome da banda não pode ser lido como o título parecido de uma obra.
    assert not pipeline._e_da_obra(pipeline._titulo_do_video("Coisa Feitta - Outra Música Qualquer", rel, nomes), "COISA FEITA", rel)
    assert not pipeline._e_da_obra(pipeline._titulo_do_video("Currículo - Coisa Feitta", rel, nomes), "COISA FEITA", rel)
    # Mas a banda cantando a obra continua valendo, pelo outro trecho do título.
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Coisa Feitta - Ligue o Rádio", rel, nomes), "LIGUE O RADIO", rel) == "exato"
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Coisa Feita", rel, nomes), "COISA FEITA", rel) == "exato"


def test_youtube_titulo_comum_nao_enche_as_pendencias(rel, monkeypatch):
    monkeypatch.setattr(pipeline, "PENDENCIAS_POR_OBRA", 2)
    ytm = YouTubeFalso(buscas={
        "COISA FEITA": [(f"a{i}", "Coisa Feita", f"Canal {i}", "", "Vídeo") for i in range(3)],
        "LIGUE O RADIO": [(f"b{i}", "Ligue o Rádio", f"Canal {i}", "", "Vídeo") for i in range(2)],
    }, leituras={})
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, recorrentes=["Banda do Baile"], limite_de_obras=2)
    assert sorted(g.link[-2:] for g in coleta.gravacoes) == ["b0", "b1"]  # título raro: lista; título comum: só avisa
    assert coleta.sem_vinculo_nao_abertos == 5 and any('"COISA FEITA" é um título comum' in a and "3 vídeos" in a for a in coleta.avisos)


def test_youtube_abre_alguns_videos_de_cada_interprete_e_os_melhores_primeiro(rel):
    famoso = [(f"f{i}", "Famoso - Coisa Feita (Cover)" if i == 0 else "Coisa Feita", "Canal de Fã" if i % 2 else "Famoso",
               f"{i} mil visualizações", "Vídeo") for i in range(30)]
    famoso = [(v, t if "Famoso" in t or canal == "Famoso" else "Famoso - Coisa Feita", canal, d, tipo) for v, t, canal, d, tipo in famoso]
    ytm = YouTubeFalso(
        buscas={"COISA FEITA Famoso": famoso,
                "COISA FEITA Cantora Original": [("c1", "Coisa Feita", "Cantora Original", "Álbum • 3:00", "Música")]},
        leituras=collections.defaultdict(lambda: _leitura("x", "Coisa Feita", "Famoso", MENU_SEM)),
    )
    ytm.leituras["c1"] = _leitura("c1", "Coisa Feita", "Cantora Original", MENU_SEM)
    coleta = pipeline.coletar_youtube(rel, Config(nomes_confirmados=["ZECA LIMA"]), ytm, [], {"COISA FEITA": ["Famoso", "Cantora Original"]},
                                      so_os_pares=True, por_interprete=3, limite_de_obras=2)
    # O intérprete com 30 vídeos não toma o lugar da outra: três dele, um dela.
    assert ytm.lidos[:2] == ["f28", "c1"] and len(ytm.lidos) == 4 and "c1" in ytm.lidos
    # Dele, primeiro o canal próprio com mais visualizações; o cover fica para o fim e nem entra.
    assert ytm.lidos == ["f28", "c1", "f26", "f24"] and "f0" not in ytm.lidos


def test_youtube_le_o_numero_de_visualizacoes():
    assert pipeline._visualizacoes("1,1 mi de visualizações • 7 mil marcações") == 1_100_000
    assert pipeline._visualizacoes("58 mil visualizações") == 58_000 and pipeline._visualizacoes("369 visualizações • 9 marcações") == 369
    assert pipeline._visualizacoes("Álbum Tal • 3:00") == 0 and pipeline._visualizacoes("") == 0


def test_youtube_acha_o_titulo_da_obra_no_meio_das_sobras_do_titulo_do_video():
    rel = relatorio_de_teste()
    nomes = ["Fulano"]
    # Três palavras ou mais: basta o título inteiro aparecer.
    assert pipeline._e_da_obra(pipeline._titulo_do_video("BANDA TAL - A DANCA DA PANELA [NOVA] setembro 2016", rel, []), "A DANCA DA PANELA", rel) == "exato"
    # Duas palavras: só quando o vídeo também cita um nome conhecido.
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Fulano Coisa Feita Video Music", rel, nomes), "COISA FEITA", rel) == "exato"
    assert not pipeline._e_da_obra(pipeline._titulo_do_video("Que coisa feita de qualquer jeito", rel, nomes), "COISA FEITA", rel)
    # Medley continua sendo medley, e não vira uma obra só.
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Ligue o Rádio / Coisa Feita (Ao Vivo)", rel, nomes), "COISA FEITA", rel) == "medley"


def test_youtube_numero_em_algarismo_casa_com_o_titulo_por_extenso():
    rel = relatorio_de_teste()
    rel.obras[0].titulo = "DEPOIS DAS TRES"
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Depois das 3", rel, []), "DEPOIS DAS TRES", rel) == "exato"
    assert pipeline._e_da_obra(pipeline._titulo_do_video("Dupla Tal - Depois das 3 (Ao Vivo)", rel, []), "DEPOIS DAS TRES", rel) == "exato"


def test_com_a_planilha_conferida_so_entra_quem_o_compositor_conferiu(rel):
    def g(obra, interprete, status, revisar=False, creditos=(), autores=()):
        return pipeline.Gravacao(plataforma="DEEZER", link=f"https://www.deezer.com/track/{obra[:2]}{interprete[:3]}", titulo=obra.title(),
                                 interprete=interprete, obra=obra, creditos=list(creditos),
                                 classificacao=c.Classificacao(status, "motivo", revisar=revisar, autores_reconhecidos=list(autores)))
    coleta = pipeline.Coleta("DEEZER", interpretes_inferidos=[{"nome": "Presumido", "obras": 3, "com_autor": 3}], gravacoes=[
        g("LIGUE O RADIO", "Banda do Baile", c.HOMONIMA),                      # conferido, sem crédito: vira firme
        g("LIGUE O RADIO", "Presumido", c.SEM_CREDITOS, revisar=True),         # presumido pelo app, fora da planilha: sai da conta
        g("COISA FEITA", "Outro Cantor", c.VIOLACAO, creditos=["Cida Dias"], autores=["CIDA DIAS"]),  # provado pelo crédito, fora da planilha
        g("COISA FEITA", "Cantora Original", c.OK, creditos=["Zeca Lima"]),    # conferido e creditado: fica como está
        g("LIGUE O RADIO", "Banda do Baile", c.VIOLACAO, revisar=True, creditos=["Fulano de Tal", "Beltrano"]),  # conferido, crédito de terceiros
    ])
    pipeline._so_os_conferidos(coleta, rel, Config(nomes_confirmados=["ZECA LIMA"]),
                               {"LIGUE O RADIO": ["Banda do Baile"], "COISA FEITA": ["Cantora Original"]})
    firme, presumido, fora, ok, errado = [x.classificacao for x in coleta.gravacoes]
    # crédito só com outros nomes, de intérprete que o compositor conferiu: violação firme, sem "a confirmar"
    assert (errado.status, errado.revisar) == (c.VIOLACAO, False) and "o compositor conferiu" in errado.fundamento
    assert (firme.status, firme.revisar) == (c.SEM_CREDITOS, False) and coleta.gravacoes[0].vinculo == "intérprete conferido pelo compositor"
    assert presumido.status == c.HOMONIMA and "não está na planilha conferida" in presumido.fundamento
    assert fora.status == c.HOMONIMA and "vale perguntar ao compositor" in fora.fundamento  # não some: fica nas pendências, com o motivo
    assert ok.status == c.OK and coleta.interpretes_inferidos == []
    assert provas.itens_da_peticao([coleta], "DEEZER")[0][0]["interprete"] == "Banda do Baile" and len(provas.itens_da_peticao([coleta], "DEEZER")[0]) == 2


# --- Amazon Music na web -------------------------------------------------------------

from creditos import amazon as _amazon


def test_amazon_le_resultados_da_busca_e_o_menu_da_faixa():
    itens = [{"titulo": "Coisa Feita", "interprete": "Cantora Original", "endereco": "/albums/B0AAA11111?trackAsin=B0TTT11111"},
             {"titulo": "Coisa Feita", "interprete": "Cantora Original", "endereco": "/albums/B0AAA22222?trackAsin=B0TTT11111"},  # repetida
             {"titulo": "", "interprete": "X", "endereco": "/albums/B0AAA33333?trackAsin=B0TTT33333"},                          # sem título
             {"titulo": "Um Álbum", "interprete": "Y", "endereco": "/albums/B0AAA44444"}]                                        # não é faixa
    [cand] = _amazon.interpretar_resultados(itens)
    assert (cand.album, cand.faixa, cand.titulo, cand.interprete) == ("B0AAA11111", "B0TTT11111", "Coisa Feita", "Cantora Original")
    assert cand.link == "https://music.amazon.com.br/albums/B0AAA11111?trackAsin=B0TTT11111"
    web = ["Ver álbum", "Ver artista", "Compartilhar esta música", "Reproduzir músicas semelhantes"]
    assert _amazon.interpretar_menu(web) == "sem_creditos"
    assert _amazon.interpretar_menu(["Adicionar à Playlist", *web]) == "sem_creditos"          # logado: mais itens, e ainda sem créditos
    assert _amazon.interpretar_menu([*web, "Créditos"]) == "com_creditos"                       # se o site passar a mostrar
    assert _amazon.interpretar_menu(["Biblioteca", "Baixadas"]) == "invalido" and _amazon.interpretar_menu([]) == "invalido"


class AmazonFalsa:
    """Busca e leitura com respostas fixas, no lugar do navegador."""
    nav = type("Nav", (), {"navegacoes": 0})()

    def __init__(self, buscas, leituras=None):
        self.buscas, self.leituras, self.lidos = buscas, leituras or {}, []

    def buscar(self, consulta):
        return [_amazon.Candidato(*c_) for c_ in self.buscas.get(consulta, [])]

    def ler(self, cand, obra):
        self.lidos.append(cand.faixa)
        return self.leituras.get(cand.faixa) or _amazon.Leitura(faixa=cand.faixa, link=cand.link, titulo=cand.titulo, interprete=cand.interprete,
                                                                 album="Álbum Tal", menu=["Ver álbum", "Ver artista"], provas=[f"prova-{cand.faixa}"])


def test_amazon_site_um_print_por_musica_de_cada_interprete_conhecido(rel):
    web = AmazonFalsa(buscas={"COISA FEITA": [
        ("A1", "T1", "Coisa Feita", "Cantora Original"), ("A2", "T2", "Coisa Feita (Ao Vivo)", "Cantora Original"),
        ("A3", "T3", "Coisa Feita", "Cantora Original"),      # a terceira da mesma intérprete não entra
        ("A4", "T4", "Coisa Feita", "Fulana & Beltrano"),     # outra intérprete conhecida: entra
        ("A5", "T5", "Coisa Feita", "Cantor Famoso"),         # homônimo de quem não tem ligação: conta, não abre
        ("A6", "T6", "Outra Música", "Cantora Original"),
    ]})
    coleta = pipeline.coletar_amazon(rel, Config(nomes_confirmados=["ZECA LIMA"]), web, [],
                                     {"COISA FEITA": ["Cantora Original", "Fulana e Beltrano"]}, limite_de_obras=2)
    assert web.lidos == ["T1", "T2", "T4"] and coleta.sem_vinculo_nao_abertos == 1 and coleta.plataforma == "AMAZON - SITE"
    primeira = coleta.gravacoes[0]
    # O site não mostra compositor: com o intérprete ligado ao titular, é "sem créditos" firme, com o print do menu.
    assert (primeira.classificacao.status, primeira.classificacao.revisar, primeira.provas) == (c.SEM_CREDITOS, False, ["prova-T1"])
    assert "não existe tela de créditos" in primeira.classificacao.fundamento
    assert primeira.link == "https://music.amazon.com.br/albums/A1?trackAsin=T1" and primeira.album == "Álbum Tal"
    firmes, _ = provas.itens_da_peticao([coleta], "AMAZON - SITE")
    assert len(firmes) == 3 and provas.NOMES["AMAZON - SITE"] == "Amazon Music (site)"


def test_amazon_site_leitura_que_falha_nunca_vira_sem_creditos(rel):
    falha = _amazon.Leitura(faixa="T1", link="https://music.amazon.com.br/albums/A1?trackAsin=T1", titulo="Coisa Feita",
                            coleta="erro", erro="o menu de ações da faixa não abriu (ou abriu outro menu)")
    web = AmazonFalsa(buscas={"COISA FEITA": [("A1", "T1", "Coisa Feita", "Cantora Original")]}, leituras={"T1": falha})
    coleta = pipeline.coletar_amazon(rel, Config(nomes_confirmados=["ZECA LIMA"]), web, [], {"COISA FEITA": ["Cantora Original"]}, limite_de_obras=2)
    assert [g.classificacao.status for g in coleta.gravacoes] == [c.ERRO_TECNICO]


def test_planilha_tem_as_duas_colunas_da_amazon(rel):
    assert saida.PLATAFORMAS[saida.PLATAFORMAS.index("AMAZON") + 1] == "AMAZON - SITE"
    assert saida._nao_coletada("AMAZON")[0] == "não coletado: aplicativo de desktop (print manual)"
    assert "amazon" in pipeline.TODAS and "AMAZON - SITE" in pipeline.ORDEM


# --- Amazon Music, aplicativo de desktop -------------------------------------------------

from creditos import amazon_app as _amazon_app


def test_amazon_app_le_o_menu_e_a_janela_de_creditos():
    menu = ["Reproduzir a próxima", "Adicionar à fila", "Adicionar à playlist", "Compartilhar música", "Créditos",
            "Reproduzir músicas semelhantes", "Fazer download"]
    assert _amazon_app.interpretar_menu(menu) == "com_creditos"
    assert _amazon_app.interpretar_menu([i for i in menu if i != "Créditos"]) == "sem_creditos"
    assert _amazon_app.interpretar_menu(["Biblioteca", "Baixadas"]) == "invalido"  # não é o menu de uma faixa
    # a lista de playlists que o aplicativo guarda escondida na página não é o menu da faixa: nunca vira "sem créditos"
    escondida = ["Adicionar à playlist", "Criar uma playlist", "Minhas curtidas", "Reproduzir a próxima", "Adicionar à fila", "Fazer download"]
    assert _amazon_app.interpretar_menu(escondida) == "invalido" and _amazon_app.interpretar_menu(escondida + menu) == "invalido"
    blocos = [{"rotulo": "Compositores", "bloco": "Créditos\nCompositores\nZeca Lima, Cida Dias"}]
    secoes = _amazon_app.interpretar_creditos(blocos)
    assert secoes == {"Compositores": ["Zeca Lima", "Cida Dias"]} and _amazon_app.autores(secoes) == ["Zeca Lima", "Cida Dias"]
    dois = [{"rotulo": "Compositores", "bloco": "Compositores\nZeca Lima\nLetristas\nCida Dias"},
            {"rotulo": "Letristas", "bloco": "Compositores\nZeca Lima\nLetristas\nCida Dias"}]
    assert _amazon_app.interpretar_creditos(dois) == {"Compositores": ["Zeca Lima"], "Letristas": ["Cida Dias"]}
    assert _amazon_app.interpretar_creditos([{"rotulo": "Compositores", "bloco": "Compositores"}]) == {}  # janela sem nome nenhum
    assert _amazon_app.interpretar_creditos([]) == {} and _amazon_app.autores({"Produtores": ["Alguém"]}) == []


class AppFalso:
    navegacoes = 0

    def __init__(self, leituras, falha=None):
        self.leituras, self.falha, self.lidos = leituras, falha, []

    def ler(self, cand, obra):
        if self.falha:
            raise self.falha
        self.lidos.append(cand.faixa)
        return self.leituras[cand.faixa]


def test_amazon_app_classifica_pelos_compositores_e_guarda_os_dois_prints(rel):
    web = AmazonFalsa(buscas={"COISA FEITA": [("A1", "T1", "Coisa Feita", "Cantora Original"), ("A2", "T2", "Coisa Feita", "Fulana & Beltrano")]})
    def lida(faixa, creditos):
        return _amazon_app.Leitura(faixa=faixa, titulo="Coisa Feita", creditos=creditos, tem_item_de_creditos=True,
                                   provas=[f"menu-{faixa}", f"creditos-{faixa}"])
    app = AppFalso({"T1": lida("T1", ["Zeca Lima", "Cida Dias"]), "T2": lida("T2", ["Outro Autor"])})
    coleta = pipeline.coletar_amazon_app(rel, Config(nomes_confirmados=["ZECA LIMA"]), web, app, [],
                                         {"COISA FEITA": ["Cantora Original", "Fulana e Beltrano"]}, limite_de_obras=2)
    creditada, errada = coleta.gravacoes
    assert coleta.plataforma == "AMAZON" and app.lidos == ["T1", "T2"]
    assert creditada.classificacao.status == c.OK
    assert (errada.classificacao.status, errada.provas) == (c.VIOLACAO, ["menu-T2", "creditos-T2"])  # dois prints por música
    assert errada.link == "https://music.amazon.com.br/albums/A2?trackAsin=T2" and "aplicativo de desktop" in errada.classificacao.fundamento


def test_amazon_app_sem_o_aplicativo_a_coleta_para_e_avisa_sem_inventar_nada(rel, tmp_path, monkeypatch):
    web = AmazonFalsa(buscas={"COISA FEITA": [("A1", "T1", "Coisa Feita", "Cantora Original")]})
    app = AppFalso({}, falha=_amazon_app.AplicativoIndisponivel("o aplicativo Amazon Music não foi encontrado neste computador"))
    coleta = pipeline.coletar_amazon_app(rel, Config(nomes_confirmados=["ZECA LIMA"]), web, app, [], {"COISA FEITA": ["Cantora Original"]}, limite_de_obras=2)
    assert coleta.gravacoes == [] and "não foi encontrado" in coleta.interrompida and any("PAROU antes do fim" in a for a in coleta.avisos)
    # Fora do Windows, ou sem o programa instalado, o próprio AmazonApp diz por quê, em vez de quebrar.
    monkeypatch.setattr(_amazon_app.AmazonApp, "_responde", lambda self: False)
    monkeypatch.setattr(_amazon_app.sys, "platform", "darwin")
    with pytest.raises(_amazon_app.AplicativoIndisponivel, match="só existe no Windows"):
        _amazon_app.AmazonApp(tmp_path).pagina
    monkeypatch.setattr(_amazon_app.sys, "platform", "win32")
    monkeypatch.setattr(_amazon_app, "caminho_do_aplicativo", lambda: None)
    with pytest.raises(_amazon_app.AplicativoIndisponivel, match="não foi encontrado"):
        _amazon_app.AmazonApp(tmp_path).pagina
    assert provas.NOMES["AMAZON"] == "Amazon Music (aplicativo)" and "amazon_app" not in pipeline.TODAS  # só entra quando a pessoa marca


def test_amazon_app_acha_o_programa_perguntando_ao_windows(tmp_path, monkeypatch):
    pasta = tmp_path / "Outro Lugar" / "Amazon Music"
    pasta.mkdir(parents=True)
    (pasta / "Amazon Music.exe").write_text("programa")
    (pasta / "unins000.exe").write_text("desinstalador")
    for variavel in ("QUEMCANTA_AMAZON_EXE", "LOCALAPPDATA", "PROGRAMFILES", "PROGRAMFILES(X86)"):
        monkeypatch.delenv(variavel, raising=False)
    # Nada nas pastas de costume: vale o que o Windows responder, nesta ordem: aberto agora, registro, loja.
    respostas = {_amazon_app._ABERTO: [], _amazon_app._REGISTRO: [str(pasta / "unins000.exe") + ",0", str(pasta)], _amazon_app._DA_LOJA: []}
    monkeypatch.setattr(_amazon_app, "_powershell", lambda comando: respostas[comando])
    assert _amazon_app.caminho_do_aplicativo() == pasta / "Amazon Music.exe"  # a pasta de instalação serve; o desinstalador, não
    respostas[_amazon_app._REGISTRO] = []
    assert _amazon_app.caminho_do_aplicativo() is None
    (pasta / "Amazon Music Helper.exe").write_text("ajudante")
    respostas[_amazon_app._ABERTO] = [str(pasta / "Amazon Music Helper.exe")]
    assert _amazon_app.caminho_do_aplicativo() == pasta / "Amazon Music.exe"  # o ajudante aberto leva ao programa principal
    so_ajudante = tmp_path / "So Ajudante"
    so_ajudante.mkdir()
    (so_ajudante / "Amazon Music Helper.exe").write_text("ajudante")
    assert _amazon_app._executavel(so_ajudante / "Amazon Music Helper.exe") is None and _amazon_app._executavel(so_ajudante) is None
    monkeypatch.setenv("QUEMCANTA_AMAZON_EXE", str(pasta / "Amazon Music.exe"))
    assert _amazon_app.caminho_do_aplicativo() == pasta / "Amazon Music.exe"
    assert _amazon_app.e_da_loja(r"C:\Program Files\WindowsApps\AmazonMobileLLC.AmazonMusic_9.5\Amazon Music.exe")
    assert not _amazon_app.e_da_loja(pasta / "Amazon Music.exe")
    monkeypatch.undo()
    assert sys.platform == "win32" or _amazon_app._powershell("Get-Date") == []  # fora do Windows, não pergunta nada


def test_amazon_app_da_loja_e_aberto_pelo_comando_do_windows(tmp_path, monkeypatch):
    loja = tmp_path / "WindowsApps" / "AmazonMobileLLC.AmazonMusic_9.5" 
    loja.mkdir(parents=True)
    (loja / "Amazon Music.exe").write_text("programa")
    comando = _amazon_app.comando_da_loja(loja / "Amazon Music.exe", 9333)
    assert "Invoke-CommandInDesktopPackage" in comando and "--remote-debugging-port=9333" in comando and "Get-AppxPackageManifest" in comando
    # Com a versão da loja, o app não tenta abrir o programa direto (o Windows recusa): usa o comando, e diz o que ele respondeu.
    pedidos, abertos = [], []
    monkeypatch.setattr(_amazon_app.sys, "platform", "win32")
    monkeypatch.setattr(_amazon_app, "caminho_do_aplicativo", lambda: loja / "Amazon Music.exe")
    monkeypatch.setattr(_amazon_app, "_powershell", lambda cmd, com_erros=False: pedidos.append(cmd) or ["pacote: X | aplicativo: Y"])
    monkeypatch.setattr(_amazon_app.subprocess, "run", lambda *a, **k: None)
    monkeypatch.setattr(_amazon_app.subprocess, "Popen", lambda *a, **k: abertos.append(a))
    monkeypatch.setattr(_amazon_app.time, "sleep", lambda s: None)
    relogio = iter(range(0, 100000, 30))
    monkeypatch.setattr(_amazon_app.time, "monotonic", lambda: next(relogio))
    monkeypatch.setattr(_amazon_app.AmazonApp, "_responde", lambda self: False)
    with pytest.raises(_amazon_app.AplicativoIndisponivel, match="aberto pela Microsoft Store. O que o Windows respondeu: pacote: X"):
        _amazon_app.AmazonApp(tmp_path)._abrir_o_aplicativo()
    assert len(pedidos) == 1 and "Invoke-CommandInDesktopPackage" in pedidos[0] and abertos == []


def test_cdp_trata_funcao_e_expressao_como_o_playwright():
    from creditos import cdp
    enviados = []

    class Falsa(cdp.PaginaCDP):
        def __init__(self):
            self.keyboard = cdp._Teclado(self)
        def comando(self, metodo, **parametros):
            enviados.append((metodo, parametros))
            return {"result": {"value": "valor"}} if metodo == "Runtime.evaluate" else {"data": "aW1hZ2Vt"}

    pagina = Falsa()
    assert pagina.evaluate("(texto) => texto.length", "é isso") == "valor"
    assert enviados[-1][1]["expression"] == '((texto) => texto.length)("é isso")' and enviados[-1][1]["awaitPromise"] is True
    pagina.evaluate("async () => { return 1; }")
    assert enviados[-1][1]["expression"] == "(async () => { return 1; })()"
    pagina.evaluate("document.body.innerText.slice(0, 1500)")  # expressão: vai como está
    assert enviados[-1][1]["expression"] == "document.body.innerText.slice(0, 1500)"
    assert pagina.screenshot(type="png") == b"imagem"
    pagina.keyboard.press("Escape")
    assert [p["type"] for m, p in enviados[-2:]] == ["keyDown", "keyUp"] and enviados[-1][1]["key"] == "Escape"
    pagina.clicar(10, 20)
    assert [p["type"] for m, p in enviados[-3:]] == ["mouseMoved", "mousePressed", "mouseReleased"]


def test_amazon_app_navega_pelo_trecho_depois_do_cerquilha(tmp_path):
    assert _amazon_app.e_de_pagina_unica("https://www.amazon.com.br/morpho/webapp/index.html#/")
    assert _amazon_app.e_de_pagina_unica("https://www.amazon.com.br/morpho/webapp/index.html")
    assert not _amazon_app.e_de_pagina_unica("https://music.amazon.com.br/albums/B0AAA11111")
    feitos = []

    class Pagina:
        def __init__(self, url): self.url = url
        def evaluate(self, script, *args): feitos.append(("script", args))
        def goto(self, url, **_): feitos.append(("goto", url))

    app = _amazon_app.AmazonApp(tmp_path)
    app._pagina = Pagina("https://www.amazon.com.br/morpho/webapp/index.html#/")
    app._ir_para("/albums/B0AAA11111?trackAsin=B0TTT11111")      # no aplicativo: troca o trecho, não carrega outro endereço
    app._pagina = Pagina("https://music.amazon.com.br/")
    app._ir_para("/albums/B0AAA11111?trackAsin=B0TTT11111")      # no site: o caminho vai no endereço
    assert feitos == [("script", ("/albums/B0AAA11111?trackAsin=B0TTT11111",)),
                      ("goto", "https://music.amazon.com.br/albums/B0AAA11111?trackAsin=B0TTT11111")]
    app._pagina = Pagina("file:///C:/app/index.html")
    with pytest.raises(_amazon_app.AplicativoIndisponivel, match="endereço interno"):
        app._ir_para("/albums/X")
    app._pagina = None


def test_scripts_usados_no_aplicativo_da_amazon_servem_num_navegador_antigo():
    """O navegador embutido no aplicativo é antigo: "?." e "??" derrubam o script inteiro (foi o que impediu o print)."""
    import inspect
    from creditos import captura as _captura
    for modulo in (_captura, _amazon_app):
        fonte = inspect.getsource(modulo)
        assert "?." not in fonte and "??" not in fonte and "replaceAll" not in fonte


def test_pdf_de_provas_do_aplicativo_da_amazon_leva_os_dois_prints_de_cada_musica(rel, tmp_path):
    """No aplicativo a prova são dois prints (o menu da faixa e a janela de créditos): os dois vão para o PDF."""
    import pdfplumber

    coleta = _coleta_para_peticao(tmp_path)
    coleta.plataforma = "AMAZON"
    for g in coleta.gravacoes:
        g.plataforma = "AMAZON"
    coleta.gravacoes[0].provas = ["prova-1", "prova-3"]  # LIGUE O RADIO: menu e créditos
    pdf, avisos = provas.pdf_de_provas(rel, [coleta], "AMAZON", tmp_path, "Amazon Music (aplicativo)")
    with pdfplumber.open(io.BytesIO(pdf)) as lido:
        paginas = [p.extract_text() for p in lido.pages]
    assert len(paginas) == 4 and "3 (1 de 2)" in paginas[0] and "3 (2 de 2)" in paginas[0]
    assert "Ligue o Radio" in paginas[2] and "(1 de 2)" in paginas[2] and "prova-1.png" in paginas[2]
    assert "Ligue o Radio" in paginas[3] and "(2 de 2)" in paginas[3] and "prova-3.png" in paginas[3]
    # nas outras plataformas continua um print por música
    coleta.plataforma = "YOUTUBE"
    for g in coleta.gravacoes:
        g.plataforma = "YOUTUBE"
    with pdfplumber.open(io.BytesIO(provas.pdf_de_provas(rel, [coleta], "YOUTUBE", tmp_path)[0])) as lido:
        assert len(lido.pages) == 3


def test_ajustes_do_teste_com_o_gabarito_da_amazon_no_aplicativo():
    import pandas as pd
    from creditos import ubc
    # letra dobrada não faz outro artista; nome diferente continua diferente
    assert pipeline._mesma_grafia("Israel e Rodolfo", "Israel & Rodolffo") and pipeline._mesma_grafia("Jefferson e Emerson", "Jeffersson e Emerson")
    assert not pipeline._mesma_grafia("Ana", "Anna Maria") and not pipeline._mesma_grafia("", "")
    assert pipeline._par_conhecido({"COISA FEITA": ["Israel e Rodolfo"]}, "COISA FEITA", "Israel & Rodolffo")
    # o aviso "Créditos indisponíveis" da janela não é nome de autor: a faixa fica sem crédito, e não com crédito errado
    secoes = _amazon_app.interpretar_creditos([{"rotulo": "Compositores", "bloco": "Compositores\n\nCréditos indisponíveis"}])
    assert secoes == {} and _amazon_app.autores(secoes) == []
    # relatório em planilha, sem coluna de código: cada obra tem o seu identificador, e não um em branco igual para todas
    rel = ubc.ler_planilha(pd.DataFrame([{"Título": "UMA", "Compositor": "ZECA LIMA"}, {"Título": "OUTRA", "Compositor": "ZECA LIMA"}]))
    assert len({o.codigo for o in rel.obras}) == 2 and all(o.codigo for o in rel.obras)


def test_claro_le_a_lista_da_busca_e_o_cartao_de_informacoes(rel):
    from types import SimpleNamespace
    from creditos import claro as _claro

    itens = [{"faixa": "50360898", "album": "6455627", "titulo": "Coisa Feita (Ao Vivo)", "artistas": ["Banda do Baile"], "nome_do_album": "Ao Vivo", "duracao": "02:44"},
             {"faixa": "50360898", "album": "6455627", "titulo": "Coisa Feita (Ao Vivo)", "artistas": ["Banda do Baile"]},  # repetida
             {"faixa": "", "album": "1", "titulo": "Sem faixa"}]
    (cand,) = _claro.interpretar_resultados(itens)
    assert (cand.album, cand.faixa, cand.interprete, cand.duracao) == ("6455627", "50360898", "Banda do Baile", "02:44")
    assert cand.link == "https://www.claromusica.com/album/6455627/BR#faixa-50360898"
    # o cartão: as duas primeiras linhas são o cabeçalho; rótulo sem valor fica vazio
    cartao = ["Coisa Feita (Ao Vivo)", "Banda do Baile", "Música", "Coisa Feita (Ao Vivo)", "Duração", "02:44", "Artista", "Banda do Baile",
              "Autores", "Zeca Lima, Cida Dias", "Álbum", "Ao Vivo", "Gravadora", "Selo Tal", "Ano", "2018"]
    lido = _claro.interpretar_cartao(cartao)
    assert _claro.autores(lido) == ["Zeca Lima", "Cida Dias"] and _claro.campo(lido, "gravadora") == "Selo Tal" and _claro.campo(lido, "musica") == "Coisa Feita (Ao Vivo)"
    sem = _claro.interpretar_cartao(["Música", "Coisa Feita", "Autores", "Álbum", "Ao Vivo"])
    assert _claro.autores(sem) == [] and _claro.campo(sem, "album") == "Ao Vivo"

    class ClaroFalsa:
        nav = SimpleNamespace(navegacoes=3)
        def __init__(self, creditos, falha=None):
            self.creditos, self.falha, self.consultas = creditos, falha, []
        def buscar(self, consulta):
            self.consultas.append(consulta)
            return [cand] if consulta == "COISA FEITA" else []
        def ler(self, c_, obra):
            if self.falha:
                raise self.falha
            return _claro.Leitura(faixa=c_.faixa, titulo=c_.titulo, creditos=self.creditos, album="Ao Vivo", gravadora="Selo Tal", ano="2018",
                                  provas=["coisa-feita_claro_informacoes"])

    pares, cfg = {"COISA FEITA": ["Banda do Baile"]}, Config(nomes_confirmados=["ZECA LIMA"])
    falsa = ClaroFalsa(["Fulano de Tal"])
    coleta = pipeline.coletar_claro(rel, cfg, falsa, pares=pares)
    (g,) = coleta.gravacoes
    assert coleta.plataforma == "CLARO" and g.classificacao.status == c.VIOLACAO and g.provas == ["coisa-feita_claro_informacoes"]
    assert all(" " not in q.replace("COISA FEITA", "").strip() or q in [o.titulo for o in rel.obras] for q in falsa.consultas)  # só por título
    assert pipeline.coletar_claro(rel, cfg, ClaroFalsa(["Zeca Lima"]), pares=pares).gravacoes[0].classificacao.status == c.OK
    assert pipeline.coletar_claro(rel, cfg, ClaroFalsa([]), pares=pares).gravacoes[0].classificacao.status == c.SEM_CREDITOS
    # sem a sessão logada, a coleta para e avisa, sem inventar resultado
    parada = pipeline.coletar_claro(rel, cfg, ClaroFalsa([], falha=_claro.SemLogin("a Claro Música não está logada")), pares=pares)
    assert parada.interrompida and not parada.gravacoes and "não está logada" in parada.avisos[-1]


def test_claro_procura_pela_pagina_do_artista_quando_a_busca_nao_traz_a_faixa(rel):
    from types import SimpleNamespace
    from creditos import claro as _claro

    class ClaroFalsa:
        nav = SimpleNamespace(navegacoes=0)
        def __init__(self):
            self.pedidos = []
        def buscar(self, consulta):  # título comum demais: a busca não devolve a faixa do intérprete
            return [_claro.Candidato("1", "9", "Coisa Feita", "Outro Cantor")]
        def buscar_pelo_artista(self, nome, mesmo_artista, e_da_obra, no_maximo=2):
            self.pedidos.append(nome)
            assert mesmo_artista("Banda do Baile") and not mesmo_artista("Banda do Bairro") and e_da_obra("Coisa Feita (Ao Vivo)") and not e_da_obra("Outra Coisa")
            return [_claro.Candidato("77", "77n3", "Coisa Feita (Ao Vivo)", "Banda do Baile")]
        def ler(self, cand, obra):
            return _claro.Leitura(faixa=cand.faixa, titulo=cand.titulo, informacoes={"Autores": "Fulano de Tal / Beltrano"}, provas=["print"])

    falsa = ClaroFalsa()
    coleta = pipeline.coletar_claro(rel, Config(nomes_confirmados=["ZECA LIMA"]), falsa, pares={"COISA FEITA": ["Banda do Baile"]})
    (g,) = coleta.gravacoes
    assert falsa.pedidos == ["Banda do Baile"] and g.interprete == "Banda do Baile" and g.link.endswith("/album/77/BR#faixa-77n3")
    assert g.creditos == ["Fulano de Tal", "Beltrano"] and g.classificacao.status == c.VIOLACAO  # a barra também separa os nomes


def test_limpar_o_guardado_apaga_so_as_pastas_das_plataformas_escolhidas(tmp_path):
    caso = tmp_path / "casos" / "fulano"
    for nome in ("spotify/leituras", "claro", "deezer", "amazon-site", "apple-music"):
        (caso / nome).mkdir(parents=True)
        (caso / nome / "a.json").write_text("{}")
    (tmp_path / "perfil-claro").mkdir()  # a sessão logada fica fora do caso
    assert sorted(pipeline.limpar_o_guardado(caso, ["spotify", "claro", "tidal"])) == ["claro", "spotify"]
    assert not (caso / "spotify").exists() and not (caso / "claro").exists()
    assert (caso / "deezer").exists() and (caso / "amazon-site").exists() and (tmp_path / "perfil-claro").exists()
    assert sorted(pipeline.limpar_o_guardado(caso, ["apple", "amazon"])) == ["amazon-site", "apple-music"]
