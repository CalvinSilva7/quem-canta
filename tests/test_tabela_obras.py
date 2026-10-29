"""Tabela de obras: recortes do relatório analítico e do ISWC, os prints do Ecadnet e o Word (dados inventados)."""

import io
import sys

import pytest

from creditos import tabela_obras as t

AZUL, CINZA, TAREFAS = (0, 94, 255), (230, 230, 230), (238, 238, 238)


def _relatorio() -> bytes:
    """Duas páginas no jeito do relatório analítico do titular 368368: a 3ª obra começa no pé da 1ª e os nomes
    continuam na 2ª, onde está o cliente."""
    from reportlab.pdfgen import canvas

    saida = io.BytesIO()
    folha = canvas.Canvas(saida, pagesize=(595, 842))

    def cabecalho():
        folha.setFont("Helvetica-Bold", 7)
        folha.drawString(43, 842 - 105, "TITULAR:")
        folha.drawString(77, 842 - 105, "368368")
        folha.drawString(118, 842 - 105, "FULANO DE TAL")
        for x, rotulo in [(30, "CÓD."), (77, "ISWC"), (139, "TÍTULO"), (305, "ASS."), (80, "NOME"), (226, "PSEUDÔNIMO")]:
            folha.drawString(x, 842 - (130 if rotulo not in ("NOME", "PSEUDÔNIMO") else 146), rotulo)

    def obra(y, codigo, iswc, titulo):
        folha.setFont("Courier-Bold", 7)
        folha.drawString(30, 842 - y, codigo)
        folha.drawString(77, 842 - y, iswc)
        folha.drawString(142, 842 - y, titulo)
        folha.drawString(305, 842 - y, "UBC")

    def nome(y, texto, codigo="12345"):
        folha.setFont("Courier", 7)
        folha.drawString(30, 842 - y, codigo)
        folha.drawString(226, 842 - y, "APELIDO")
        folha.setFont("Courier-Oblique", 7)
        folha.drawString(80, 842 - y, texto)

    cabecalho()
    obra(170, "111", "T-123.456.789-0", "OBRA UM")
    nome(181, "BELTRANO DA SILVA")
    nome(192, "FULANO DE TAL", "368368")
    nome(203, "WARNER CHAPPELL EDICOES MUSICAIS LTDA")  # depois do cliente: não entra
    obra(220, "222", "-  .  .  -", "OBRA SEM ISWC")
    nome(231, "CICLANO")
    obra(780, "333", "T-000.000.001-0", "OBRA TRES")
    nome(791, "AAA PRIMEIRO")
    folha.setFont("Helvetica", 7)
    folha.drawString(30, 842 - 815, "27/04/2026 - 15:30   Página 1")
    folha.showPage()
    cabecalho()
    nome(170, "FULANO DE TAL", "368368")
    nome(181, "ZZZ DEPOIS DO CLIENTE")
    obra(200, "444", "T-999.999.999-9", "OBRA QUATRO")
    nome(201, "JOSE")
    folha.save()
    return saida.getvalue()


def _iswc(*obras) -> bytes:
    """Uma página por obra, no jeito do resultado do ISWC Network: o quadro com as linhas separadas por traços."""
    from reportlab.pdfgen import canvas

    saida = io.BytesIO()
    folha = canvas.Canvas(saida, pagesize=(842, 1191))
    for iswc, titulo in obras:
        folha.setFont("Helvetica", 9)
        folha.drawString(52, 1191 - 180, "ISWC:")  # o do campo de busca: o valor fica em outra linha
        folha.drawString(56, 1191 - 205, iswc)
        folha.rect(52, 1191 - 655, 738, 288)
        for y in (412, 458, 503, 563, 608):
            folha.rect(53, 1191 - y, 737, 0.5, fill=1)
        for y, rotulo, valor in [(392, "ISWC:", iswc), (438, "Original Title:", titulo), (483, "Rank:", "1"),
                                 (535, "Creator Name(s):", "TAL FULANO DE"), (588, "IP Name Number:", "00012345678")]:
            folha.drawString(61, 1191 - y, rotulo)
            folha.drawString(277, 1191 - y, valor)
        folha.showPage()
    folha.save()
    return saida.getvalue()


def test_relatorio_acha_cada_obra_com_os_titulares_mesmo_na_pagina_seguinte():
    obras = t.obras_do_relatorio(_relatorio())
    assert [(o.codigo, o.iswc, o.titulo) for o in obras] == [
        ("111", "T-123.456.789-0", "OBRA UM"), ("222", "", "OBRA SEM ISWC"),
        ("333", "T-000.000.001-0", "OBRA TRES"), ("444", "T-999.999.999-9", "OBRA QUATRO")]
    um, _, tres, _ = obras
    # a linha da obra e os nomes até o do cliente (pelo código dele, 368368); a Warner, que vem depois, fica de fora
    assert len(um.trechos) == 1 and 192 < um.trechos[0].bottom < 203
    assert [tr.pagina for tr in tres.trechos] == [0, 1]  # o cliente só aparece na página 2: ela entra
    assert 170 < tres.trechos[1].bottom < 181  # e para no cliente
    # o recorte começa na coluna do ISWC (sem o código da obra) e não pega o pseudônimo, que não é itálico
    assert um.trechos[0].x0 < 77 and um.trechos[0].x1 < 226
    imagens = t.recortes_do_relatorio(_relatorio(), tres)
    assert len(imagens) == 2 and all(im.width > 100 and im.height > 20 for im in imagens)


def test_iswc_recorta_o_quadro_ate_os_criadores_e_acha_pelo_iswc():
    recortes = t.recortes_do_iswc(_iswc(("T-123.456.789-0", "OBRA UM"), ("T-000.000.001-0", "OBRA TRES")))
    assert set(recortes) == {"1234567890", "0000000010"}
    imagem = recortes["1234567890"]
    # da borda de cima do quadro até a divisória antes do "IP Name Number" (563 - 367 pontos), na escala 3
    assert abs(imagem.height - (563 - 367 + 2) * t.ESCALA_DO_ISWC) <= 6 and abs(imagem.width - 740 * t.ESCALA_DO_ISWC) <= 9


def test_linhas_sao_as_obras_do_iswc_na_ordem_do_relatorio_com_o_titulo_do_relatorio():
    linhas, obras, avisos = t.preparar(_relatorio(), [_iswc(("T-000.000.001-0", "X"), ("T-123.456.789-0", "Y"), ("T-555.555.555-5", "Z"))])
    assert [(linha.titulo, linha.iswc, len(linha.relatorio)) for linha in linhas] == [
        ("Obra Um", "T-123.456.789-0", 1), ("Obra Tres", "T-000.000.001-0", 2)]
    assert len(avisos) == 1 and "T-555.555.555-5" in avisos[0]  # está no ISWC e não no relatório


def _print(barra_em=150, rolado=False, linhas_da_tabela=(300, 400, 500), largura=1920, altura=1080):
    """Um print de tela do Ecadnet: fundo cinza, o quadro branco no meio, a barra azul, as divisórias e a barra de
    tarefas do Windows embaixo (cinza, com a caixa de busca branca)."""
    from PIL import Image, ImageDraw

    imagem = Image.new("RGB", (largura, altura), CINZA)
    desenho = ImageDraw.Draw(imagem)
    desenho.rectangle((460, 87, 1444, altura - 49), fill="white")
    if not rolado:
        desenho.rectangle((501, barra_em, 1403, barra_em + 37), fill=AZUL)
    for y in linhas_da_tabela:
        desenho.line((501, y, 1403, y), fill=CINZA)
    desenho.text((501, linhas_da_tabela[-1] + 40), "*As informações disponibilizadas no Ecadnet...", fill="gray")
    desenho.rectangle((0, altura - 48, largura, altura), fill=TAREFAS)
    desenho.rectangle((740, altura - 40, 960, altura - 10), fill="white")
    return imagem


def test_ecadnet_acha_a_barra_azul_e_recorta_ate_a_ultima_tabela():
    p = t.PrintDoEcadnet(imagem=_print())
    p.barra = t.barra_azul(p.imagem)
    assert p.barra == (501, 150, 1403, 187)
    recorte = t.recorte_do_ecadnet(p, (501, 1403))
    # da barra até a última divisória: sem o rodapé do site e sem a barra de tarefas
    assert recorte.size[1] == 500 + 2 - 148 and 900 <= recorte.size[0] <= 906
    rolado = t.PrintDoEcadnet(imagem=_print(rolado=True), topo_da_pagina=103)
    assert t.barra_azul(rolado.imagem) is None
    assert t.recorte_do_ecadnet(rolado, (501, 1403)).size[1] == 500 + 2 - 103


def test_o_que_o_leitor_leu_vira_codigo_iswc_e_titulo():
    p = t.PrintDoEcadnet(imagem=None, barra=(501, 155, 1403, 192))
    linhas = [{"texto": "2; ecadnet.ecad.org.br/consulta/detalhes/obras-musicais/18437950", "y0": 171, "y1": 213},
              {"texto": "u.• A Farra Perdeu Pro Amor", "y0": 498, "y1": 546},
              {"texto": "T-040.z79.666-0", "y0": 818, "y1": 843}]  # fmt: skip
    t.interpretar_print(p, linhas, 3)
    assert (p.codigo, p.iswc, p.titulo, p.topo_da_pagina) == ("18437950", "0404796660", "A Farra Perdeu Pro Amor", 87)


def test_cada_print_vai_para_a_obra_certa():
    obras = [t.ObraDoRelatorio("111", "T-123.456.789-0", "OBRA UM"), t.ObraDoRelatorio("333", "T-000.000.001-0", "CARDÁPIO DE CAMA"),
             t.ObraDoRelatorio("555", "", "FORA DO ISWC")]  # fmt: skip
    barra = (501, 155, 1403, 192)
    prints = [t.PrintDoEcadnet(None, barra, codigo="111"),  # pelo código do endereço
              t.PrintDoEcadnet(None, None, codigo="111"),  # a mesma página, rolada
              t.PrintDoEcadnet(None, barra, iswc="0000000011"),  # ISWC lido com um dígito errado
              t.PrintDoEcadnet(None, None),  # rolado, sem endereço lido: continuação do anterior
              t.PrintDoEcadnet(None, barra, titulo="Cardapio De Carna"),  # só o título, lido com erro
              t.PrintDoEcadnet(None, barra, codigo="555"),  # obra do relatório que não está no ISWC
              t.PrintDoEcadnet(None, None)]  # continuação da obra de fora: fica de fora também
    t.ligar_prints(prints, obras, {"1234567890", "0000000010"})
    assert [(p.obra, p.como) for p in prints] == [
        ("1234567890", "código"), ("1234567890", "continuação"), ("0000000010", "ISWC"), ("0000000010", "continuação"),
        ("0000000010", "título"), ("", "fora da tabela"), ("", "fora da tabela")]


def test_word_no_modelo_do_escritorio():
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from PIL import Image

    imagem = t.png(Image.new("RGB", (400, 200), "white"))
    linhas = [t.LinhaDaTabela("Obra Um", "T-123.456.789-0", [imagem, imagem], imagem, [imagem]),
              t.LinhaDaTabela("Obra Tres", "T-000.000.001-0", [imagem], imagem, [])]
    doc = Document(io.BytesIO(t.montar_docx(linhas)))
    assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE
    tabela = doc.tables[0]
    assert [c.text.split("\n")[0] for c in tabela.rows[0].cells] == ["OBRA MUSICAL", "RELATÓRIO ANALÍTICO", "ISWC", "ECADNET"]
    assert [r.cells[0].text for r in tabela.rows[1:]] == ["Obra Um", "Obra Tres"]
    figuras = [[len(c._tc.xpath(".//a:blip")) for c in r.cells] for r in tabela.rows[1:]]
    assert figuras == [[0, 2, 1, 1], [0, 1, 1, 0]]  # sem Ecadnet: a célula fica vazia
    # cada recorte na largura útil da coluna (a coluna menos as margens da célula)
    larguras = {int(e.get("cx")) for e in tabela.rows[1].cells[1]._tc.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")}
    assert larguras == {(3355 - 216) * 635}
    assert tabela.rows[0].cells[0].paragraphs[0].runs[0].font.name == "Bookman Old Style"


def test_arquivos_enviados_viram_pdfs_e_prints():
    import zipfile

    from docx import Document
    from PIL import Image

    pacote = io.BytesIO()
    with zipfile.ZipFile(pacote, "w") as z:
        z.writestr("02 - B.pdf", b"%PDF-b")
        z.writestr("01 - A.pdf", b"%PDF-a")
    assert t.pdfs_de("ISWC.zip", pacote.getvalue()) == [b"%PDF-a", b"%PDF-b"]  # na ordem do nome (01, 02…)
    doc = Document()
    for cor in ("red", "blue"):
        doc.add_picture(io.BytesIO(t.png(Image.new("RGB", (50, 30), cor))))
    saida = io.BytesIO()
    doc.save(saida)
    prints = t.imagens_de("Ecadnet.docx", saida.getvalue())
    assert [Image.open(io.BytesIO(p)).getpixel((0, 0)) for p in prints] == [(255, 0, 0), (0, 0, 255)]


@pytest.mark.skipif(sys.platform != "win32", reason="o leitor de texto é o do Windows")
def test_leitor_do_windows_le_o_endereco():
    from PIL import Image, ImageDraw, ImageFont

    imagem = Image.new("RGB", (1400, 120), "white")
    ImageDraw.Draw(imagem).text((20, 30), "ecadnet.ecad.org.br/consulta/detalhes/obras-musicais/18437950", fill="black",
                                font=ImageFont.truetype("arial.ttf", 36))
    try:
        linhas = t.ler_textos([imagem])
    except t.LeitorIndisponivel as e:
        pytest.skip(f"leitor do Windows indisponível aqui: {e}")
    assert "18437950" in " ".join(linha["texto"] for linha in linhas[0]).replace(" ", "")


def test_prints_do_ecadnet_em_pdf(tmp_path):
    """O PDF que a Réplica (ECAD) junta (uma imagem por página): cada página vira um print, na resolução original."""
    from PIL import Image

    from creditos import replica

    caminhos = []
    for i, cor in enumerate(("red", "blue")):
        Image.new("RGB", (1920, 1080), cor).save(tmp_path / f"{i}.png")
        caminhos.append(tmp_path / f"{i}.png")
    prints = t.imagens_de("Ecadnet.pdf", replica.pdf_dos_prints(caminhos))
    imagens = [Image.open(io.BytesIO(p)) for p in prints]
    assert [im.size for im in imagens] == [(1920, 1080), (1920, 1080)]
    assert [im.convert("RGB").getpixel((5, 5)) for im in imagens] == [(255, 0, 0), (0, 0, 255)]
